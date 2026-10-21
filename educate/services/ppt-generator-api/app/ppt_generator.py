"""
PowerPoint Presentation Generation Engine for Acadion AI PPT Generator.
Converts structured PresentationData into a polished, editable, widescreen .pptx presentation.
"""

import math
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches, Pt

from app.diagrams import DiagramRenderer, RectBounds
from app.layouts import SlideLayoutEngine, ThemeColors
from app.models import (
    BulletPoint,
    PresentationData,
    PPTGeneratorError,
    Section,
    SectionType,
    SlideVisual,
    VisualType,
    VisualRenderError,
)
from app.readability_validator import ReadabilityValidator


class PPTGenerator:
    """Engine responsible for building editable, professionally styled PowerPoint presentations."""

    MAX_BULLETS_PER_SLIDE = 4  # Prevent overflowing slides by splitting large sections

    def __init__(
        self,
        layout_engine: Optional[SlideLayoutEngine] = None,
        diagram_renderer: Optional[DiagramRenderer] = None,
        theme: Optional[str | ThemeColors] = None,
        layout_style: Optional[str] = "smart_mixed",
    ):
        if layout_engine:
            self.layout = layout_engine
        elif theme:
            self.layout = SlideLayoutEngine(theme=theme)
        else:
            self.layout = SlideLayoutEngine()
        self.colors = self.layout.colors
        self.diagram_renderer = diagram_renderer or DiagramRenderer(self.layout)
        self.layout_style = (layout_style or "smart_mixed").lower().strip()


    def generate(
        self,
        data: PresentationData,
        output_path: str | Path,
        layout_style: Optional[str] = None,
        max_slides: Optional[int] = None,
    ) -> Path:
        """Generate PowerPoint presentation from structured data and save to output_path.
        
        Returns:
            Path: The resolved path to the created presentation file.
        """
        out_file = Path(output_path)
        out_file.parent.mkdir(parents=True, exist_ok=True)

        effective_style = (layout_style or self.layout_style or "smart_mixed").lower().strip()

        prs = Presentation()
        self.layout.set_presentation_dimensions(prs)

        # Plan all slides first to know the total slide count for accurate "Slide X of Y" footers
        slide_plans = self._plan_slides(data, layout_style=effective_style, max_slides=max_slides)
        total_slides = len(slide_plans)

        for plan in slide_plans:
            slide_type = plan["type"]
            slide = None
            if slide_type == "title":
                slide = self._render_title_slide(prs, data, plan)
            elif slide_type == "overview":
                slide = self._render_overview_slide(prs, data, plan, total_slides)
            elif slide_type == "section_divider":
                slide = self._render_section_divider_slide(prs, data, plan, total_slides)
            elif slide_type in ("pull_quote", "quote"):
                slide = self._render_pull_quote_slide(prs, data, plan, total_slides)
            elif slide_type in ("split_image", "split_image_insight"):
                slide = self._render_split_image_slide(prs, data, plan, total_slides)
            elif slide_type in ("quadrant_matrix", "matrix"):
                slide = self._render_quadrant_matrix_slide(prs, data, plan, total_slides)
            elif slide_type in ("process_chevrons", "chevrons"):
                slide = self._render_process_chevrons_slide(prs, data, plan, total_slides)
            elif slide_type == "editorial":
                slide = self._render_editorial_content_slide(prs, data, plan, total_slides)
            elif slide_type == "asymmetric":
                slide = self._render_asymmetric_slide(prs, data, plan, total_slides)
            elif slide_type == "case_study":
                slide = self._render_case_study_slide(prs, data, plan, total_slides)
            elif slide_type == "multi_column":
                slide = self._render_multi_column_slide(prs, data, plan, total_slides)
            elif slide_type == "discussion":
                slide = self._render_discussion_slide(prs, data, plan, total_slides)
            elif slide_type == "cards_grid":
                slide = self._render_cards_grid_slide(prs, data, plan, total_slides)
            elif slide_type == "two_column":
                slide = self._render_two_column_slide(prs, data, plan, total_slides)
            elif slide_type in ("stats_metrics", "kpi_metric_trio", "kpi_metrics"):
                slide = self._render_metrics_slide(prs, data, plan, total_slides)
            elif slide_type == "split":
                slide = self._render_split_slide(prs, data, plan, total_slides)
            elif slide_type == "visual":
                slide = self._render_visual_slide(prs, data, plan, total_slides)
            elif slide_type == "summary":
                slide = self._render_summary_slide(prs, data, plan, total_slides)
            else:
                slide = self._render_content_slide(prs, data, plan, total_slides)

            # Attach speaker notes if present in plan
            notes = plan.get("speaker_notes")
            if notes and slide:
                slide.notes_slide.notes_text_frame.text = notes

        try:
            prs.save(str(out_file))
            return out_file.resolve()
        except PermissionError:
            # If the destination file is open in Microsoft PowerPoint, write to a versioned path
            fallback_name = f"{out_file.stem}_generated{out_file.suffix}"
            fallback_path = out_file.parent / fallback_name
            prs.save(str(fallback_path))
            print(f"\n[NOTE] '{out_file.name}' is currently locked (open in PowerPoint). Saved to '{fallback_name}' instead.")
            return fallback_path.resolve()

    def _disperse_content_units(
        self,
        title: str,
        bullets: List[BulletPoint],
    ) -> Tuple[str, Optional[List[BulletPoint]], Optional[str]]:
        """
        Disperse multi-step titles into discrete logical bullet units before layout assignment.
        Ensures headings remain concise descriptions of the slide content, sequential steps
        become individual process units, and appropriate layouts are selected.
        """
        import re
        step_pattern = re.compile(
            r"(?:^|(?<=\s))(?:(?:Step|Phase|Stage)\s*(\d+)[:\.]?|(\d+)[\.\)])\s*([^0-9\n]+?)(?=(?:\s+(?:(?:Step|Phase|Stage)\s*\d+[:\.]?|\d+[\.\)]))|\Z)",
            re.IGNORECASE
        )
        matches = list(step_pattern.finditer(title))
        if len(matches) >= 2:
            first_match_start = matches[0].start()
            prefix_text = title[:first_match_start].strip().rstrip(":-|–—")

            dispersed: List[BulletPoint] = []
            for m in matches:
                num = m.group(1) or m.group(2)
                desc = m.group(3).strip().rstrip(".,;")
                if ":" in desc:
                    pfx, txt = desc.split(":", 1)
                    dispersed.append(BulletPoint(prefix=f"Step {num}: {pfx.strip()}", text=txt.strip()))
                else:
                    dispersed.append(BulletPoint(prefix=f"Step {num}", text=desc))

            # Preserve any existing bullets without duplicating
            for b in bullets:
                if not any(b.text.lower() in d.text.lower() for d in dispersed):
                    dispersed.append(b)

            if prefix_text and len(prefix_text) >= 4:
                clean_title = prefix_text
            else:
                first_step_text = matches[0].group(3).split(":")[0].strip()
                last_step_text = matches[-1].group(3).split(":")[0].strip()
                clean_title = f"{first_step_text} to {last_step_text} Execution Pipeline"
                if len(clean_title) > 60:
                    clean_title = "System Execution & Process Workflow"

            return clean_title, dispersed, "process_chevrons"

        return title, None, None

    def _plan_slides(
        self,
        data: PresentationData,
        layout_style: Optional[str] = None,
        max_slides: Optional[int] = None,
    ) -> List[Dict[str, Any]]:
        """Compute slide breakdown, layout family assignment, and pagination."""
        import re
        style = (layout_style or getattr(self, "layout_style", "smart_mixed") or "smart_mixed").lower().strip()
        plans: List[Dict[str, Any]] = []

        # 1. Title Slide (Slide 1)
        title_style = getattr(data, "title_composition", "editorial") or "editorial"
        plans.append({
            "type": "title",
            "title_style": title_style,
            "slide_num": 1,
            "speaker_notes": f"Introduction and lecture overview: {data.title}."
        })
        current_num = 2
        content_sections = data.content_sections

        # 2. Overview / Agenda Slide (Include if overview present or synthesize for standalone presentations)
        overview_sec = data.overview_section
        if not overview_sec and not data.is_part_of_series and (max_slides is None or max_slides >= 4):
            overview_bullets = [
                BulletPoint(text=sec.title[:80], prefix=f"Module {idx+1:02d}")
                for idx, sec in enumerate(content_sections[:5])
            ]
            overview_sec = Section(
                title="Curriculum Overview & Agenda",
                section_type=SectionType.INTRO_OVERVIEW,
                bullets=overview_bullets,
            )

        if overview_sec and (max_slides is None or max_slides >= 4):
            overview_style = getattr(data, "overview_composition", "agenda_flow") or "agenda_flow"
            overview_notes = overview_sec.speaker_notes if overview_sec else ""
            plans.append({
                "type": "overview",
                "overview_style": overview_style,
                "slide_num": current_num,
                "section": overview_sec,
                "speaker_notes": overview_notes or f"Overview of topics covered in {data.title}.",
            })
            current_num += 1

        # 3. Content Slides (with layout family intelligence & pagination)
        content_sections = data.content_sections
        for sec_idx, sec in enumerate(content_sections):
            bullets = list(sec.bullets)
            paragraphs = sec.paragraphs
            speaker_notes = sec.speaker_notes
            visual = sec.visual
            comp_type = (sec.composition_type or "").lower().strip()

            # Content dispersion before layout selection
            clean_title, dispersed_bullets, suggested_type = self._disperse_content_units(sec.title, bullets)
            sec_title = clean_title
            if dispersed_bullets is not None:
                bullets = dispersed_bullets
                if not comp_type:
                    comp_type = suggested_type

            # If section has valid native diagram visual, use visual diagram when in adaptive/flow layout
            if style in ("smart_mixed", "auto", "visual_flow") and visual and visual.visual_type != VisualType.NONE and self._is_visual_valid(visual):
                plans.append({
                    "type": "visual",
                    "slide_num": current_num,
                    "title": sec_title,
                    "category": self._get_visual_category(visual.visual_type),
                    "visual": visual,
                    "paragraphs": paragraphs,
                    "bullets": bullets,
                    "teaching_objective": sec.teaching_objective,
                    "key_message": sec.key_message,
                    "speaker_notes": speaker_notes or f"Visual analysis and architecture for {sec_title}.",
                })
                current_num += 1
                continue

            # Determine layout type based on requested style and section semantics
            chosen_type = "editorial"
            lower_title = sec_title.lower()

            if style in ("cards_grid", "two_column", "stats_metrics", "classic_bullets", "visual_flow") and style not in ("smart_mixed", "auto"):
                if style == "cards_grid":
                    chosen_type = "cards_grid"
                elif style == "two_column":
                    chosen_type = "two_column"
                elif style == "stats_metrics":
                    chosen_type = "stats_metrics"
                elif style == "classic_bullets":
                    chosen_type = "editorial"
                elif style == "visual_flow":
                    chosen_type = "process_chevrons"
            elif comp_type and comp_type in (
                "editorial", "asymmetric", "case_study", "discussion",
                "two_column", "multi_column", "cards_grid", "kpi_metrics", "split",
                "pull_quote", "quote", "split_image", "split_image_insight",
                "quadrant_matrix", "matrix", "process_chevrons", "chevrons",
                "stats_metrics", "kpi_metric_trio"
            ):
                if comp_type in ("kpi_metrics", "kpi_metric_trio"):
                    chosen_type = "stats_metrics"
                elif comp_type in ("pull_quote", "quote"):
                    chosen_type = "pull_quote"
                elif comp_type in ("split_image", "split_image_insight"):
                    chosen_type = "split_image"
                elif comp_type in ("quadrant_matrix", "matrix"):
                    chosen_type = "quadrant_matrix"
                elif comp_type in ("process_chevrons", "chevrons"):
                    chosen_type = "process_chevrons"
                else:
                    chosen_type = comp_type
            else:
                # Content-aware semantic analysis
                if any(w in lower_title for w in ("quote", "axiom", "philosophy", "maxim", "principle")):
                    chosen_type = "pull_quote"
                elif any(w in lower_title for w in ("matrix", "quadrant", "tradeoff", "trade-off", "swot")):
                    chosen_type = "quadrant_matrix"
                elif any(w in lower_title for w in ("pipeline", "step-by-step", "lifecycle", "workflow", "stages", "chevrons", "sequence", "phases")) or (len(bullets) >= 2 and all(re.match(r"^(?:Step|Phase|Stage|\d+)", b.prefix or "") for b in bullets)):
                    chosen_type = "process_chevrons"
                elif any(w in lower_title for w in ("diagram", "sensor fusion", "telemetry", "visual explainer")):
                    chosen_type = "split_image"
                elif any(w in lower_title for w in ("discussion", "question", "review", "exercise", "debate", "socratic")):
                    chosen_type = "discussion"
                elif any(w in lower_title for w in ("case study", "empirical study", "breakthrough", "real-world deployment")) and (len(bullets) >= 3 or (len(paragraphs) >= 1 and len(bullets) >= 2)):
                    chosen_type = "case_study"
                elif any(w in lower_title for w in (" vs ", " versus ", "comparison", "classical vs", "deep learning vs")):
                    chosen_type = "two_column"
                elif any(w in lower_title for w in ("architecture", "transformer", "neural", "mechanic", "foundational", "formula")):
                    chosen_type = "asymmetric"
                elif len(bullets) == 1:
                    chosen_type = "asymmetric" if len(paragraphs) > 0 else "editorial"
                elif len(bullets) == 2:
                    chosen_type = "two_column"
                elif len(bullets) == 3:
                    chosen_type = "multi_column"
                elif any(re.search(r"(\d+[%xXkKMmB]?|\d{4})", b.prefix or b.text) for b in bullets) and len(bullets) <= 4:
                    chosen_type = "stats_metrics"
                elif len(bullets) == 4 and sec_idx % 3 == 0:
                    chosen_type = "cards_grid"
                else:
                    chosen_type = "editorial"

            # Evaluate content readability against layout and apply 4-stage fallback if needed
            validated_type, _ = ReadabilityValidator.evaluate_content_for_layout(
                layout_type=chosen_type,
                title=sec_title,
                bullets=bullets,
                paragraphs=paragraphs,
            )
            chosen_type = validated_type

            category_labels = {
                "editorial": "CURRICULUM TOPIC",
                "asymmetric": "CORE ARCHITECTURE",
                "case_study": "CASE STUDY & BREAKTHROUGH",
                "discussion": "SEMINAR & DISCUSSION",
                "two_column": "PARADIGM COMPARISON",
                "multi_column": "FOUNDATIONAL PILLARS",
                "cards_grid": "MODULAR CONCEPTS",
                "stats_metrics": "KEY METRICS & HIGHLIGHTS",
                "split": "SYSTEM ARCHITECTURE",
                "pull_quote": "CORE PEDAGOGICAL AXIOM",
                "split_image": "VISUAL SYSTEM ARCHITECTURE",
                "quadrant_matrix": "DECISION MATRIX & TRADEOFFS",
                "process_chevrons": "SYSTEM EXECUTION PIPELINE",
                "content": "MODULE CONTENT",
                "content_bullets": "MODULE CONTENT",
            }
            category_text = category_labels.get(chosen_type, "MODULE CONTENT")

            # Extract image path if visual image is present
            img_path = None
            img_caption = None
            if sec.visual and hasattr(sec.visual, "image") and sec.visual.image:
                img_path = getattr(sec.visual.image, "file_path", None)
                img_caption = getattr(sec.visual.image, "caption", None)

            # Filter valid non-empty bullets
            valid_bullets = [b for b in bullets if (b.text and b.text.strip()) or (b.prefix and b.prefix.strip())]

            # Pagination check: split large bullet lists to prevent overcrowding
            # Note: For multi_column, 3 items fit cleanly. For editorial or cards_grid, 4 items max.
            max_items = 3 if chosen_type == "multi_column" else self.MAX_BULLETS_PER_SLIDE
            if len(valid_bullets) > max_items:
                chunks = [
                    valid_bullets[i:i + max_items]
                    for i in range(0, len(valid_bullets), max_items)
                ]
                # If splitting would breach max_slides, reflow onto 1 slide with safe margins & compact spacing
                if max_slides is not None and len(plans) + len(chunks) > max_slides:
                    plans.append({
                        "type": "content_bullets" if len(valid_bullets) > 4 else chosen_type,
                        "slide_num": current_num,
                        "title": sec_title,
                        "category": category_text,
                        "paragraphs": paragraphs,
                        "bullets": valid_bullets,
                        "teaching_objective": sec.teaching_objective,
                        "key_message": sec.key_message,
                        "speaker_notes": speaker_notes,
                        "image_path": img_path,
                        "caption": img_caption,
                    })
                    current_num += 1
                else:
                    total_parts = len(chunks)
                    for part_idx, chunk in enumerate(chunks, 1):
                        # Ensure chunk has non-empty text
                        if not any((b.text and b.text.strip()) or (b.prefix and b.prefix.strip()) for b in chunk):
                            continue
                        plans.append({
                            "type": chosen_type,
                            "slide_num": current_num,
                            "title": f"{sec_title} (Part {part_idx}/{total_parts})",
                            "category": category_text,
                            "paragraphs": paragraphs if part_idx == 1 else [],
                            "bullets": chunk,
                            "teaching_objective": sec.teaching_objective,
                            "key_message": sec.key_message if part_idx == 1 else (chunk[0].text if chunk else ""),
                            "speaker_notes": speaker_notes if part_idx == 1 else f"Continuation of {sec_title}.",
                            "image_path": img_path if part_idx == 1 else None,
                            "caption": img_caption if part_idx == 1 else None,
                        })
                        current_num += 1
            else:
                plans.append({
                    "type": chosen_type,
                    "slide_num": current_num,
                    "title": sec_title,
                    "category": category_text,
                    "paragraphs": paragraphs,
                    "bullets": valid_bullets,
                    "teaching_objective": sec.teaching_objective,
                    "key_message": sec.key_message,
                    "speaker_notes": speaker_notes,
                    "image_path": img_path,
                    "caption": img_caption,
                })
                current_num += 1

        # 4. Summary / Conclusion Slide
        conclusion_style = getattr(data, "conclusion_composition", "synthesis_takeaways") or "synthesis_takeaways"
        can_add_summary = (max_slides is None or len(plans) < max_slides)
        is_series_part = getattr(data, "is_part_of_series", False)

        if data.conclusion_section and can_add_summary:
            plans.append({
                "type": "summary",
                "summary_style": conclusion_style,
                "slide_num": current_num,
                "section": data.conclusion_section,
                "speaker_notes": data.conclusion_section.speaker_notes,
            })
            current_num += 1
        elif not is_series_part and can_add_summary:
            # Generate default summary for standalone presentation only
            plans.append({
                "type": "summary",
                "summary_style": conclusion_style,
                "slide_num": current_num,
                "section": Section(
                    title="Summary & Key Takeaways",
                    section_type=SectionType.SUMMARY_CONCLUSION,
                    bullets=[
                        BulletPoint(text="Review the core conceptual paradigms discussed in this lecture.", prefix="Foundational Mastery"),
                        BulletPoint(text="Analyze real-world applications and ethical considerations.", prefix="Practical Application"),
                        BulletPoint(text="Consult supplementary curriculum readings and problem sets.", prefix="Next Steps"),
                    ],
                ),
                "speaker_notes": "Conclude presentation with review of key takeaways and student assignments.",
            })
            current_num += 1

        # Enforce absolute hard maximum if specified
        if max_slides is not None and len(plans) > max_slides:
            plans = plans[:max_slides]

        # Re-number slide_num from 1 to len(plans)
        for idx, p in enumerate(plans, 1):
            p["slide_num"] = idx

        return plans

    # =========================================================================
    # Slide Renderers
    # =========================================================================

    def _render_title_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Optional[Dict[str, Any]] = None
    ) -> Slide:
        """Render Title slide with content-aware presentation style."""
        title_style = (plan.get("title_style") if plan else None) or getattr(data, "title_composition", "editorial") or "editorial"
        title_style = title_style.lower().strip()

        if title_style == "cinematic":
            return self._render_cinematic_title_slide(prs, data, plan)
        elif title_style == "split":
            return self._render_split_title_slide(prs, data, plan)
        else:
            return self._render_editorial_title_slide(prs, data, plan)

    def _render_editorial_title_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Optional[Dict[str, Any]] = None
    ) -> Slide:
        """Render modern, clean editorial title slide with generous whitespace and high typographic hierarchy."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        # Top Category / Series Pill
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.9), Inches(1.1), Inches(4.5), Inches(0.42)
        )
        pill.fill.solid()
        pill.fill.fore_color.rgb = self.colors.card_tint
        pill.line.color.rgb = self.colors.secondary
        pill.line.width = Pt(1)
        tf_pill = pill.text_frame
        p_pill = tf_pill.paragraphs[0]
        p_pill.text = f"{data.institution.upper()} • CURRICULUM SERIES"
        p_pill.font.name = self.layout.FONT_HEADING
        p_pill.font.size = Pt(9.5)
        p_pill.font.bold = True
        p_pill.font.color.rgb = self.colors.secondary
        p_pill.alignment = PP_ALIGN.CENTER

        # Giant High-Impact Presentation Title
        title_box = slide.shapes.add_textbox(
            Inches(0.9), Inches(1.8), Inches(11.5), Inches(2.3)
        )
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = data.title
        p_title.font.name = self.layout.FONT_HEADING
        p_title.font.size = Pt(36)
        p_title.font.bold = True
        p_title.font.color.rgb = self.colors.text_primary if self.colors.is_dark else self.colors.primary
        p_title.line_spacing = 1.1

        # Subtitle
        sub_box = slide.shapes.add_textbox(
            Inches(0.9), Inches(4.2), Inches(11.5), Inches(0.85)
        )
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = data.subtitle or "Comprehensive Curriculum & Faculty Lecture Notes"
        p_sub.font.name = self.layout.FONT_BODY
        p_sub.font.size = Pt(16.5)
        p_sub.font.color.rgb = self.colors.text_secondary

        # Decorative dual-accent rules
        bar1 = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.9), Inches(5.2), Inches(2.2), Inches(0.06)
        )
        bar1.fill.solid()
        bar1.fill.fore_color.rgb = self.colors.accent
        bar1.line.fill.background()

        bar2 = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(3.2), Inches(5.22), Inches(9.2), Inches(0.02)
        )
        bar2.fill.solid()
        bar2.fill.fore_color.rgb = self.colors.card_border
        bar2.line.fill.background()

        # Metadata Card at Bottom
        meta_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.9), Inches(5.45), Inches(11.5), Inches(1.1)
        )
        meta_card.fill.solid()
        meta_card.fill.fore_color.rgb = self.colors.card_bg
        meta_card.line.color.rgb = self.colors.card_border
        meta_card.line.width = Pt(1)

        tf_meta = meta_card.text_frame
        tf_meta.word_wrap = True
        tf_meta.margin_left = Inches(0.3)
        tf_meta.margin_top = Inches(0.18)

        p_m1 = tf_meta.paragraphs[0]
        p_m1.text = f"Academic Institution: {data.institution}"
        p_m1.font.name = self.layout.FONT_BODY
        p_m1.font.size = Pt(11)
        p_m1.font.bold = True
        p_m1.font.color.rgb = self.colors.accent if self.colors.is_dark else self.colors.primary

        p_m2 = tf_meta.add_paragraph()
        date_display = data.date_str or "Academic Session"
        p_m2.text = f"Delivered: {date_display}   •   Format: 16:9 Widescreen Courseware   •   Standard Educational Deck"
        p_m2.font.name = self.layout.FONT_BODY
        p_m2.font.size = Pt(10)
        p_m2.font.color.rgb = self.colors.text_muted

        return slide

    def _render_cinematic_title_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Optional[Dict[str, Any]] = None
    ) -> Slide:
        """Render high-impact cinematic hero title slide fully respecting active theme colors."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)

        # Apply canvas background from the active theme
        self.layout.apply_canvas_background(slide)

        is_dark = self.colors.is_dark

        # Outer Container Frame: card_bg surface (pure white in light themes, dark card surface in dark themes)
        frame = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(1.0), Inches(0.8), Inches(11.333), Inches(5.9)
        )
        frame.name = "HeroFrame"
        frame.fill.solid()
        frame.fill.fore_color.rgb = self.colors.card_bg
        frame.line.color.rgb = self.colors.accent if is_dark else self.colors.card_border
        frame.line.width = Pt(1.5)

        # Top Category / Accent Pill
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(1.5), Inches(1.3), Inches(4.2), Inches(0.4)
        )
        pill.fill.solid()
        pill.fill.fore_color.rgb = self.colors.card_tint
        pill.line.color.rgb = self.colors.accent if is_dark else self.colors.secondary
        pill.line.width = Pt(1)
        tf_pill = pill.text_frame
        p_pill = tf_pill.paragraphs[0]
        p_pill.text = "ADVANCED LECTURE & RESEARCH SERIES"
        p_pill.font.name = self.layout.FONT_HEADING
        p_pill.font.size = Pt(9.5)
        p_pill.font.bold = True
        p_pill.font.color.rgb = self.colors.accent if is_dark else self.colors.secondary
        p_pill.alignment = PP_ALIGN.CENTER

        # Title
        title_box = slide.shapes.add_textbox(
            Inches(1.5), Inches(2.0), Inches(10.3), Inches(2.3)
        )
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = data.title
        p_title.font.name = self.layout.FONT_HEADING
        p_title.font.size = Pt(36)
        p_title.font.bold = True
        p_title.font.color.rgb = self.colors.text_primary if is_dark else self.colors.primary
        p_title.line_spacing = 1.1

        # Subtitle
        sub_box = slide.shapes.add_textbox(
            Inches(1.5), Inches(4.3), Inches(10.3), Inches(0.8)
        )
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = data.subtitle or "Comprehensive Curriculum & Faculty Lecture Notes"
        p_sub.font.name = self.layout.FONT_BODY
        p_sub.font.size = Pt(16.5)
        p_sub.font.color.rgb = self.colors.text_secondary

        # Bottom Accent Rule
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(1.5), Inches(5.2), Inches(10.3), Inches(0.04)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = self.colors.accent if is_dark else self.colors.secondary
        line.line.fill.background()

        tb_meta = slide.shapes.add_textbox(
            Inches(1.5), Inches(5.35), Inches(10.3), Inches(0.6)
        )
        tf_m = tb_meta.text_frame
        p_m = tf_m.paragraphs[0]
        p_m.text = f"{data.institution}   •   {data.date_str or 'Academic Session'}   •   16:9 Widescreen Master Deck"
        p_m.font.name = self.layout.FONT_BODY
        p_m.font.size = Pt(10.5)
        p_m.font.color.rgb = self.colors.text_muted

        return slide

    def _render_split_title_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Optional[Dict[str, Any]] = None
    ) -> Slide:
        """Render Title slide with left hero accent column and right content area."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)

        # Base canvas matching the active theme
        self.layout.apply_canvas_background(slide)

        is_dark = self.colors.is_dark

        # Left Accent Hero Column
        col_width = Inches(3.8)
        col = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), col_width, self.layout.SLIDE_HEIGHT
        )
        col.name = "HeroColumn"
        col.fill.solid()
        col_fill = self.colors.card_bg if is_dark else self.colors.card_tint
        col.fill.fore_color.rgb = col_fill
        col.line.fill.background()

        # Vertical accent line along the column border
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            col_width, Inches(0), Inches(0.08), self.layout.SLIDE_HEIGHT
        )
        line.fill.solid()
        line.fill.fore_color.rgb = self.colors.accent if is_dark else self.colors.secondary
        line.line.fill.background()

        # Left Column decorative badge & institution branding
        badge_box = slide.shapes.add_textbox(
            Inches(0.6), Inches(1.0), Inches(2.8), Inches(4.5)
        )
        tf_badge = badge_box.text_frame
        tf_badge.word_wrap = True

        p1 = tf_badge.paragraphs[0]
        p1.text = "ACADION"
        p1.font.name = self.layout.FONT_HEADING
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = self.colors.text_primary if is_dark else self.colors.primary

        p2 = tf_badge.add_paragraph()
        p2.text = "CURRICULUM SERIES"
        p2.font.name = self.layout.FONT_HEADING
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = self.colors.accent if is_dark else self.colors.secondary
        p2.space_after = Pt(24)

        p3 = tf_badge.add_paragraph()
        p3.text = "Standard Educational Courseware & Presentation Engine"
        p3.font.name = self.layout.FONT_BODY
        p3.font.size = Pt(11)
        p3.font.color.rgb = self.colors.text_muted if is_dark else self.colors.text_secondary

        # Main Title & Subtitle Card on the Right (Inches 4.5 to 12.5)
        right_start = Inches(4.5)
        right_width = Inches(8.0)

        # Tag pill
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            right_start, Inches(1.6), Inches(2.6), Inches(0.38)
        )
        pill.fill.solid()
        pill.fill.fore_color.rgb = self.colors.card_tint
        pill.line.color.rgb = self.colors.secondary
        pill.line.width = Pt(1)
        tf_pill = pill.text_frame
        p_pill = tf_pill.paragraphs[0]
        p_pill.text = "ACADEMIC LECTURE NOTES"
        p_pill.font.name = self.layout.FONT_HEADING
        p_pill.font.size = Pt(9.5)
        p_pill.font.bold = True
        p_pill.font.color.rgb = self.colors.secondary
        p_pill.alignment = PP_ALIGN.CENTER

        # Main Presentation Title Box
        title_box = slide.shapes.add_textbox(
            right_start, Inches(2.2), right_width, Inches(2.2)
        )
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = data.title
        p_title.font.name = self.layout.FONT_HEADING
        p_title.font.size = Pt(34)
        p_title.font.bold = True
        p_title.font.color.rgb = self.colors.text_primary if self.colors.is_dark else self.colors.primary
        p_title.line_spacing = 1.1

        # Subtitle Box
        sub_box = slide.shapes.add_textbox(
            right_start, Inches(4.3), right_width, Inches(1.0)
        )
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = data.subtitle or "Comprehensive Curriculum & Faculty Lecture Notes"
        p_sub.font.name = self.layout.FONT_BODY
        p_sub.font.size = Pt(16)
        p_sub.font.color.rgb = self.colors.text_secondary

        # Metadata Card (Author, Institution, Date)
        meta_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            right_start, Inches(5.4), right_width, Inches(1.1)
        )
        meta_card.fill.solid()
        meta_card.fill.fore_color.rgb = self.colors.card_bg
        meta_card.line.color.rgb = self.colors.card_border
        meta_card.line.width = Pt(1)

        tf_meta = meta_card.text_frame
        tf_meta.word_wrap = True
        tf_meta.margin_left = Inches(0.3)
        tf_meta.margin_top = Inches(0.18)

        p_meta1 = tf_meta.paragraphs[0]
        p_meta1.text = f"Institution: {data.institution}"
        p_meta1.font.name = self.layout.FONT_BODY
        p_meta1.font.size = Pt(11)
        p_meta1.font.bold = True
        p_meta1.font.color.rgb = self.colors.accent if self.colors.is_dark else self.colors.primary

        p_meta2 = tf_meta.add_paragraph()
        p_meta2.text = f"Delivered: {data.date_str or 'Academic Session'}   •   Format: 16:9 Widescreen Courseware"
        p_meta2.font.name = self.layout.FONT_BODY
        p_meta2.font.size = Pt(10)
        p_meta2.font.color.rgb = self.colors.text_muted

        return slide

    def _render_overview_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int
    ) -> Slide:
        """Render Agenda / Overview slide with varied content-aware compositions."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        sec: Optional[Section] = plan.get("section")
        title = sec.title if sec else "Curriculum Overview & Agenda"
        overview_style = (plan.get("overview_style") or getattr(data, "overview_composition", "agenda_flow") or "agenda_flow").lower().strip()

        self.layout.add_header(slide, title=title, category="OVERVIEW & AGENDA")
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        # If user/notes provided narrative text explicitly
        if sec and sec.paragraphs and overview_style != "agenda_flow":
            narrative_text = " ".join(sec.paragraphs)
            self.layout.create_card_container(
                slide,
                left=Inches(0.8),
                top=Inches(1.75),
                width=Inches(11.733),
                height=Inches(1.3),
                accent_color=self.colors.accent,
                bg_color=self.colors.card_tint,
            )
            box = slide.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(11.2), Inches(1.1))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = narrative_text
            p.font.name = self.layout.FONT_BODY
            p.font.size = Pt(13.5)
            p.font.color.rgb = self.colors.text_primary
            p.line_spacing = 1.25

            card_top = Inches(3.2)
            card_height = Inches(3.3)
            self.layout.create_card_container(
                slide,
                left=Inches(0.8),
                top=card_top,
                width=Inches(11.733),
                height=card_height,
                accent_color=self.colors.secondary,
            )
            self._render_bullets_in_box(
                slide,
                left=Inches(1.2),
                top=card_top + Inches(0.25),
                width=Inches(11.0),
                height=card_height - Inches(0.5),
                bullets=sec.bullets[:4],
            )
        elif overview_style == "editorial":
            self._render_editorial_roadmap(slide, data.content_sections)
        elif overview_style == "cards":
            self._render_agenda_grid(slide, data.content_sections)
        else:
            # Default modern agenda flow: progressive roadmap cards with connecting track
            self._render_agenda_flow(slide, data.content_sections, data.overview_section)

        return slide

    def _render_agenda_flow(
        self,
        slide: Slide,
        sections: List[Section],
        overview_sec: Optional[Section] = None
    ) -> None:
        """Render a sequential roadmap flow with stage cards and connection track."""
        roadmap_items: List[Tuple[str, str, str]] = []  # (stage_label, title, desc)

        if overview_sec and overview_sec.bullets:
            for idx, b in enumerate(overview_sec.bullets[:5]):
                stage_lbl = b.prefix or f"STAGE {idx+1:02d}"
                title = b.text
                # Find matching section description if possible
                matching_sec = next((s for s in sections if s.title == b.text or b.text in s.title), None)
                desc = ""
                if matching_sec and matching_sec.bullets:
                    desc = matching_sec.bullets[0].text
                elif matching_sec and matching_sec.paragraphs:
                    desc = matching_sec.paragraphs[0]
                else:
                    desc = "Core curriculum topic and foundational principles."
                roadmap_items.append((stage_lbl, title, desc))
        else:
            count = min(len(sections), 5)
            for idx in range(count):
                sec = sections[idx]
                stage_lbl = f"STAGE {idx+1:02d}"
                desc = sec.bullets[0].text if sec.bullets else (sec.paragraphs[0] if sec.paragraphs else "Core foundational study.")
                roadmap_items.append((stage_lbl, sec.title, desc))

        count = len(roadmap_items)
        if count == 0:
            return

        # Connecting track line across the middle of the cards
        track_y = Inches(3.7)
        track = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(1.2), track_y, Inches(10.9), Inches(0.04)
        )
        track.fill.solid()
        track.fill.fore_color.rgb = self.colors.secondary
        track.line.fill.background()

        gap = Inches(0.3)
        col_w = (Inches(11.733) - (gap * (count - 1))) / count
        card_h = Inches(4.5)
        top_y = Inches(1.85)

        for i, (stage_lbl, title, desc) in enumerate(roadmap_items):
            c_left = Inches(0.8) + i * (col_w + gap)

            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                c_left, top_y, col_w, card_h
            )
            card.fill.solid()
            card.fill.fore_color.rgb = self.colors.card_bg
            card.line.color.rgb = self.colors.card_border
            card.line.width = Pt(1)

            # Node Badge
            badge_size = Inches(0.55)
            badge_x = c_left + (col_w - badge_size) / 2
            badge = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                badge_x, top_y + Inches(0.3), badge_size, badge_size
            )
            badge.fill.solid()
            badge.fill.fore_color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent
            badge.line.color.rgb = self.colors.accent
            badge.line.width = Pt(1.5)
            tf_b = badge.text_frame
            p_b = tf_b.paragraphs[0]
            p_b.text = f"{i+1:02d}"
            p_b.font.name = self.layout.FONT_HEADING
            p_b.font.size = Pt(11)
            p_b.font.bold = True
            p_b.font.color.rgb = self.colors.text_white
            p_b.alignment = PP_ALIGN.CENTER

            # Stage Label
            tb_lbl = slide.shapes.add_textbox(c_left + Inches(0.1), top_y + Inches(1.05), col_w - Inches(0.2), Inches(0.35))
            tf_lbl = tb_lbl.text_frame
            tf_lbl.word_wrap = True
            p_lbl = tf_lbl.paragraphs[0]
            p_lbl.text = stage_lbl.upper()
            p_lbl.font.name = self.layout.FONT_HEADING
            p_lbl.font.size = Pt(9.5)
            p_lbl.font.bold = True
            p_lbl.font.color.rgb = self.colors.accent if self.colors.is_dark else self.colors.secondary
            p_lbl.alignment = PP_ALIGN.CENTER

            # Module Title Box
            tb_title = slide.shapes.add_textbox(c_left + Inches(0.15), top_y + Inches(1.45), col_w - Inches(0.3), Inches(1.3))
            tf_t = tb_title.text_frame
            tf_t.word_wrap = True
            p_t = tf_t.paragraphs[0]
            clean_title = title if len(title) <= 45 else title[:42].rsplit(" ", 1)[0] + "..."
            p_t.text = clean_title
            p_t.font.name = self.layout.FONT_HEADING
            p_t.font.size = Pt(12)
            p_t.font.bold = True
            p_t.font.color.rgb = self.colors.text_primary
            p_t.alignment = PP_ALIGN.CENTER

            # Divider inside card
            div = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, c_left + Inches(0.2), top_y + Inches(2.85), col_w - Inches(0.4), Inches(0.02)
            )
            div.fill.solid()
            div.fill.fore_color.rgb = self.colors.card_border
            div.line.fill.background()

            # Bullets or description preview with strict boundary fit
            tb_body = slide.shapes.add_textbox(c_left + Inches(0.15), top_y + Inches(2.95), col_w - Inches(0.3), Inches(1.35))
            tf_body = tb_body.text_frame
            tf_body.word_wrap = True
            p_desc = tf_body.paragraphs[0]
            clean_desc = desc if len(desc) <= 75 else desc[:72].rsplit(" ", 1)[0] + "..."
            p_desc.text = clean_desc
            p_desc.font.name = self.layout.FONT_BODY
            p_desc.font.size = Pt(10)
            p_desc.font.color.rgb = self.colors.text_secondary
            p_desc.line_spacing = 1.15

    def _render_editorial_roadmap(self, slide: Slide, sections: List[Section]) -> None:
        """Render a clean, high-legibility two-column editorial curriculum agenda."""
        col_w = Inches(5.6)
        row_h = Inches(1.35)
        lefts = [Inches(0.8), Inches(6.9)]
        start_y = Inches(1.85)

        for i, sec in enumerate(sections[:6]):
            col_idx = i % 2
            row_idx = i // 2
            c_left = lefts[col_idx]
            c_top = start_y + (row_idx * (row_h + Inches(0.25)))

            # Transparent card with left accent rule
            bar = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, c_left, c_top, Inches(0.08), row_h
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = self.colors.accent if i % 2 == 0 else self.colors.secondary
            bar.line.fill.background()

            tb = slide.shapes.add_textbox(c_left + Inches(0.25), c_top, col_w - Inches(0.3), row_h)
            tf = tb.text_frame
            tf.word_wrap = True

            p1 = tf.paragraphs[0]
            p1.text = f"MODULE {i+1:02d}  •  CURRICULUM UNIT"
            p1.font.name = self.layout.FONT_HEADING
            p1.font.size = Pt(9.5)
            p1.font.bold = True
            p1.font.color.rgb = self.colors.secondary

            p2 = tf.add_paragraph()
            p2.text = sec.title
            p2.font.name = self.layout.FONT_HEADING
            p2.font.size = Pt(13)
            p2.font.bold = True
            p2.font.color.rgb = self.colors.text_primary
            p2.space_before = Pt(3)

            p3 = tf.add_paragraph()
            sub = sec.bullets[0].text if sec.bullets else "Core curriculum concepts and analysis."
            p3.text = sub
            p3.font.name = self.layout.FONT_BODY
            p3.font.size = Pt(10.5)
            p3.font.color.rgb = self.colors.text_secondary
            p3.space_before = Pt(2)

    def _render_agenda_grid(self, slide: Slide, sections: List[Section]) -> None:
        """Render a clean grid of lecture modules when no explicit overview text is provided."""
        col_width = Inches(5.6)
        row_height = Inches(1.3)
        lefts = [Inches(0.8), Inches(6.9)]
        start_top = Inches(1.8)

        for i, sec in enumerate(sections[:6]):
            col_idx = i % 2
            row_idx = i // 2
            c_left = lefts[col_idx]
            c_top = start_top + (row_idx * (row_height + Inches(0.25)))

            # Card shape
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                c_left, c_top, col_width, row_height
            )
            card.fill.solid()
            card.fill.fore_color.rgb = self.colors.card_bg
            card.line.color.rgb = self.colors.card_border
            card.line.width = Pt(1)

            # Number badge
            badge = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                c_left + Inches(0.25), c_top + Inches(0.25), Inches(0.7), Inches(0.4)
            )
            badge.fill.solid()
            badge.fill.fore_color.rgb = self.colors.card_tint
            badge.line.color.rgb = self.colors.secondary
            badge.line.width = Pt(1)
            tf_b = badge.text_frame
            p_b = tf_b.paragraphs[0]
            p_b.text = f"{i+1:02d}"
            p_b.font.name = self.layout.FONT_HEADING
            p_b.font.size = Pt(12)
            p_b.font.bold = True
            p_b.font.color.rgb = self.colors.secondary
            p_b.alignment = PP_ALIGN.CENTER

            # Module Title
            box = slide.shapes.add_textbox(
                c_left + Inches(1.1), c_top + Inches(0.2), Inches(4.3), Inches(0.85)
            )
            tf_m = box.text_frame
            tf_m.word_wrap = True
            p_m = tf_m.paragraphs[0]
            p_m.text = sec.title
            p_m.font.name = self.layout.FONT_HEADING
            p_m.font.size = Pt(13)
            p_m.font.bold = True
            p_m.font.color.rgb = self.colors.primary

    def _render_content_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int
    ) -> Slide:
        """Render a modular content slide with card containers and formatted bullets."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Lecture Topic")
        category = plan.get("category", "MODULE CONTENT")
        paragraphs = plan.get("paragraphs", [])
        bullets: List[BulletPoint] = plan.get("bullets", [])

        rule_top = self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        current_top = max(Inches(1.7), (rule_top if rule_top is not None else Inches(1.45)) + Inches(0.22))
        available_height = max(Inches(3.0), Inches(6.65) - current_top)

        # If there are narrative paragraphs, render a concise top note card
        if paragraphs:
            para_text = " ".join(paragraphs)
            para_card_height = Inches(1.1)
            self.layout.create_card_container(
                slide,
                left=Inches(0.8),
                top=current_top,
                width=Inches(11.733),
                height=para_card_height,
                accent_color=self.colors.accent,
                bg_color=self.colors.card_tint,
            )
            p_box = slide.shapes.add_textbox(
                Inches(1.1), current_top + Inches(0.12), Inches(11.2), Inches(0.9)
            )
            tf_p = p_box.text_frame
            tf_p.word_wrap = True
            p_elem = tf_p.paragraphs[0]
            p_elem.text = para_text
            p_elem.font.name = self.layout.FONT_BODY
            p_elem.font.size = Pt(13.5)
            p_elem.font.color.rgb = self.colors.text_primary

            current_top += para_card_height + Inches(0.2)
            available_height -= (para_card_height + Inches(0.2))

        # Main bullets card container
        self.layout.create_card_container(
            slide,
            left=Inches(0.8),
            top=current_top,
            width=Inches(11.733),
            height=available_height,
            accent_color=self.colors.secondary,
        )

        self._render_bullets_in_box(
            slide,
            left=Inches(1.2),
            top=current_top + Inches(0.25),
            width=Inches(11.0),
            height=available_height - Inches(0.5),
            bullets=bullets,
        )

        return slide

    def _render_editorial_content_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render modern editorial layout with breathable typography and prominent key takeaways."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Lecture Topic")
        category = plan.get("category", "CURRICULUM TOPIC")
        bullets: List[BulletPoint] = plan.get("bullets", [])
        paragraphs = plan.get("paragraphs", [])
        key_msg = plan.get("key_message") or (paragraphs[0] if paragraphs else None)

        self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        current_top = Inches(1.75)

        # Prominent Key Message / Pull-quote Callout Banner
        if key_msg:
            pull_card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.8), current_top, Inches(11.733), Inches(1.15)
            )
            pull_card.fill.solid()
            pull_card.fill.fore_color.rgb = self.colors.card_tint
            pull_card.line.color.rgb = self.colors.secondary
            pull_card.line.width = Pt(1.5)

            # Left accent highlight strip on callout
            strip = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0.8), current_top, Inches(0.12), Inches(1.15)
            )
            strip.fill.solid()
            strip.fill.fore_color.rgb = self.colors.accent
            strip.line.fill.background()

            tb_msg = slide.shapes.add_textbox(
                Inches(1.1), current_top + Inches(0.12), Inches(11.2), Inches(0.9)
            )
            tf_m = tb_msg.text_frame
            tf_m.word_wrap = True
            p_lbl = tf_m.paragraphs[0]
            p_lbl.text = "CORE TAKEAWAY / KEY MESSAGE"
            p_lbl.font.name = self.layout.FONT_HEADING
            p_lbl.font.size = Pt(9.5)
            p_lbl.font.bold = True
            p_lbl.font.color.rgb = self.colors.accent if self.colors.is_dark else self.colors.primary

            p_txt = tf_m.add_paragraph()
            p_txt.text = key_msg
            p_txt.font.name = self.layout.FONT_BODY
            p_txt.font.size = Pt(13)
            p_txt.font.color.rgb = self.colors.text_primary
            p_txt.space_before = Pt(3)

            current_top += Inches(1.35)

        # Editorial Content Card: Clean, high-legibility container with open spacing
        remaining_h = Inches(6.65) - current_top
        ed_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), current_top, Inches(11.733), remaining_h
        )
        ed_card.fill.solid()
        ed_card.fill.fore_color.rgb = self.colors.card_bg
        ed_card.line.color.rgb = self.colors.card_border
        ed_card.line.width = Pt(1)

        # Render bullet points with refined typographic spacing
        self._render_bullets_in_box(
            slide,
            left=Inches(1.2),
            top=current_top + Inches(0.25),
            width=Inches(11.0),
            height=remaining_h - Inches(0.5),
            bullets=bullets,
        )

        return slide

    def _render_asymmetric_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render asymmetric focal layout (40% focal panel + 60% structured breakdown)."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Core Architecture")
        category = plan.get("category", "CORE ARCHITECTURE")
        bullets: List[BulletPoint] = plan.get("bullets", [])
        obj = plan.get("teaching_objective") or "Foundational Mastery"
        key_msg = plan.get("key_message") or (bullets[0].text if bullets else "Core architectural principles.")

        self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        top_y = Inches(1.8)
        content_h = Inches(4.75)

        # Left Focal Panel (4.4" wide)
        left_w = Inches(4.4)
        focal_panel = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), top_y, left_w, content_h
        )
        focal_panel.fill.solid()
        panel_color = self.colors.primary if self.colors.is_dark else self.colors.card_tint
        focal_panel.fill.fore_color.rgb = panel_color
        focal_panel.line.color.rgb = self.colors.secondary
        focal_panel.line.width = Pt(1.5)

        # Top Accent Ribbon on Focal Panel
        ribbon = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), top_y, left_w, Inches(0.12)
        )
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = self.colors.accent
        ribbon.line.fill.background()

        # Focal Text Box
        tb_focal = slide.shapes.add_textbox(
            Inches(1.05), top_y + Inches(0.35), left_w - Inches(0.5), content_h - Inches(0.7)
        )
        tf_focal = tb_focal.text_frame
        tf_focal.word_wrap = True

        p_obj_lbl = tf_focal.paragraphs[0]
        p_obj_lbl.text = "TEACHING OBJECTIVE"
        p_obj_lbl.font.name = self.layout.FONT_HEADING
        p_obj_lbl.font.size = Pt(9.5)
        p_obj_lbl.font.bold = True
        p_obj_lbl.font.color.rgb = self.colors.accent if self.colors.is_dark else self.colors.primary

        p_obj = tf_focal.add_paragraph()
        p_obj.text = obj
        p_obj.font.name = self.layout.FONT_HEADING
        p_obj.font.size = Pt(15)
        p_obj.font.bold = True
        p_obj.font.color.rgb = self.colors.text_white if self.colors.is_dark else self.colors.primary
        p_obj.space_before = Pt(4)
        p_obj.space_after = Pt(18)

        p_msg_lbl = tf_focal.add_paragraph()
        p_msg_lbl.text = "FOUNDATIONAL PRINCIPLE"
        p_msg_lbl.font.name = self.layout.FONT_HEADING
        p_msg_lbl.font.size = Pt(9.5)
        p_msg_lbl.font.bold = True
        p_msg_lbl.font.color.rgb = self.colors.secondary

        p_msg = tf_focal.add_paragraph()
        p_msg.text = key_msg
        p_msg.font.name = self.layout.FONT_BODY
        p_msg.font.size = Pt(12.5)
        p_msg.font.color.rgb = self.colors.text_primary if not self.colors.is_dark else self.colors.text_white
        p_msg.space_before = Pt(4)
        p_msg.line_spacing = 1.25

        # Right Breakdown Panel (6.9" wide)
        right_x = Inches(5.6)
        right_w = Inches(6.933)
        self.layout.create_card_container(
            slide,
            left=right_x,
            top=top_y,
            width=right_w,
            height=content_h,
            accent_color=self.colors.secondary,
        )

        # Bullets in Right Panel (skip bullet 0 if it was identical to key_msg and we have >2 bullets)
        render_bullets = bullets[1:] if (len(bullets) > 2 and bullets[0].text == key_msg) else bullets
        self._render_bullets_in_box(
            slide,
            left=right_x + Inches(0.4),
            top=top_y + Inches(0.3),
            width=right_w - Inches(0.8),
            height=content_h - Inches(0.6),
            bullets=render_bullets,
        )

        return slide

    def _render_case_study_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render empirical case study, application, or breakthrough slide with context & impact cards."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Case Study & Real-World Application")
        category = plan.get("category", "CASE STUDY & BREAKTHROUGH")
        bullets: List[BulletPoint] = plan.get("bullets", [])
        paragraphs = plan.get("paragraphs", [])

        self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        # 1. Top Context Card
        top_context_h = Inches(1.2)
        top_context_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), Inches(1.75), Inches(11.733), top_context_h
        )
        top_context_card.fill.solid()
        top_context_card.fill.fore_color.rgb = self.colors.card_tint
        top_context_card.line.color.rgb = self.colors.secondary
        top_context_card.line.width = Pt(1)

        # Badge & Text in Top Context
        tb_ctx = slide.shapes.add_textbox(
            Inches(1.05), Inches(1.85), Inches(11.2), Inches(0.95)
        )
        tf_ctx = tb_ctx.text_frame
        tf_ctx.word_wrap = True

        p_b = tf_ctx.paragraphs[0]
        p_b.text = "EMPIRICAL CONTEXT & SCOPE"
        p_b.font.name = self.layout.FONT_HEADING
        p_b.font.size = Pt(9.5)
        p_b.font.bold = True
        p_b.font.color.rgb = self.colors.accent if self.colors.is_dark else self.colors.primary

        p_t = tf_ctx.add_paragraph()
        # Partition bullets strictly without cross-box duplication
        avail_bullets = list(bullets)
        if not paragraphs and avail_bullets:
            ctx_text = avail_bullets.pop(0).text
        else:
            ctx_text = " ".join(paragraphs) if paragraphs else "Case analysis context."

        p_t.text = ctx_text
        p_t.font.name = self.layout.FONT_BODY
        p_t.font.size = Pt(12.5)
        p_t.font.color.rgb = self.colors.text_primary
        p_t.space_before = Pt(3)

        # 2. Middle Split: Challenge / Architecture vs Empirical Outcome / Impact
        mid_y = Inches(3.1)
        mid_h = Inches(2.4)
        col_w = Inches(5.65)
        gap = Inches(0.43)

        # Left: Challenge / System Implementation
        left_x = Inches(0.8)
        left_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left_x, mid_y, col_w, mid_h
        )
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = self.colors.card_bg
        left_card.line.color.rgb = self.colors.card_border
        left_card.line.width = Pt(1)

        l_strip = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left_x, mid_y, col_w, Inches(0.1)
        )
        l_strip.fill.solid()
        l_strip.fill.fore_color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent
        l_strip.line.fill.background()

        tb_l = slide.shapes.add_textbox(left_x + Inches(0.25), mid_y + Inches(0.2), col_w - Inches(0.5), mid_h - Inches(0.35))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True
        p_lh = tf_l.paragraphs[0]
        p_lh.text = "SYSTEM ARCHITECTURE & CHALLENGE"
        p_lh.font.name = self.layout.FONT_HEADING
        p_lh.font.size = Pt(11)
        p_lh.font.bold = True
        p_lh.font.color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent

        left_pts = [avail_bullets.pop(0)] if avail_bullets else []
        if left_pts:
            for b in left_pts:
                pref = b.prefix.strip().rstrip(":") if b.prefix else ""
                txt = b.text.strip() if b.text else ""
                if pref and txt.lower().startswith((pref + ":").lower()):
                    txt = txt[len(pref) + 1:].strip()
                p_item = tf_l.add_paragraph()
                p_item.text = f"▪ {pref + ': ' if pref else ''}{txt}"
                p_item.font.name = self.layout.FONT_BODY
                p_item.font.size = Pt(11)
                p_item.font.color.rgb = self.colors.text_secondary
                p_item.space_before = Pt(6)
        else:
            p_item = tf_l.add_paragraph()
            p_item.text = "▪ System architecture and operational challenges."
            p_item.font.name = self.layout.FONT_BODY
            p_item.font.size = Pt(11)
            p_item.font.color.rgb = self.colors.text_secondary
            p_item.space_before = Pt(6)

        # Right: Empirical Outcome & Benchmark
        right_x = left_x + col_w + gap
        right_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, right_x, mid_y, col_w, mid_h
        )
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = self.colors.card_bg
        right_card.line.color.rgb = self.colors.card_border
        right_card.line.width = Pt(1)

        r_strip = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, right_x, mid_y, col_w, Inches(0.1)
        )
        r_strip.fill.solid()
        r_strip.fill.fore_color.rgb = self.colors.secondary
        r_strip.line.fill.background()

        tb_r = slide.shapes.add_textbox(right_x + Inches(0.25), mid_y + Inches(0.2), col_w - Inches(0.5), mid_h - Inches(0.35))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True
        p_rh = tf_r.paragraphs[0]
        p_rh.text = "MEASURED RESULTS & BENCHMARKS"
        p_rh.font.name = self.layout.FONT_HEADING
        p_rh.font.size = Pt(11)
        p_rh.font.bold = True
        p_rh.font.color.rgb = self.colors.secondary

        right_pts = [avail_bullets.pop(0)] if avail_bullets else []
        if right_pts:
            for b in right_pts:
                pref = b.prefix.strip().rstrip(":") if b.prefix else ""
                txt = b.text.strip() if b.text else ""
                if pref and txt.lower().startswith((pref + ":").lower()):
                    txt = txt[len(pref) + 1:].strip()
                p_item = tf_r.add_paragraph()
                p_item.text = f"▪ {pref + ': ' if pref else ''}{txt}"
                p_item.font.name = self.layout.FONT_BODY
                p_item.font.size = Pt(11)
                p_item.font.color.rgb = self.colors.text_secondary
                p_item.space_before = Pt(6)
        else:
            p_item = tf_r.add_paragraph()
            p_item.text = "▪ Evaluated against performance benchmarks and requirements."
            p_item.font.name = self.layout.FONT_BODY
            p_item.font.size = Pt(11)
            p_item.font.color.rgb = self.colors.text_secondary
            p_item.space_before = Pt(6)

        # 3. Bottom Takeaway Banner
        bot_y = Inches(5.65)
        bot_banner = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), bot_y, Inches(11.733), Inches(0.9)
        )
        bot_banner.fill.solid()
        bot_banner.fill.fore_color.rgb = self.colors.card_tint
        bot_banner.line.color.rgb = self.colors.accent
        bot_banner.line.width = Pt(1)

        tb_bot = slide.shapes.add_textbox(Inches(1.05), bot_y + Inches(0.1), Inches(11.2), Inches(0.7))
        tf_bot = tb_bot.text_frame
        tf_bot.word_wrap = True
        p_t1 = tf_bot.paragraphs[0]
        p_t1.text = "PRACTICAL INDUSTRY IMPLICATION:"
        p_t1.font.name = self.layout.FONT_HEADING
        p_t1.font.size = Pt(9.5)
        p_t1.font.bold = True
        p_t1.font.color.rgb = self.colors.accent if self.colors.is_dark else self.colors.primary

        p_t2 = tf_bot.add_paragraph()
        takeaway = avail_bullets.pop(0).text if avail_bullets else "Deploying architectures requires continuous validation against production metrics."
        p_t2.text = takeaway
        p_t2.font.name = self.layout.FONT_BODY
        p_t2.font.size = Pt(11.5)
        p_t2.font.color.rgb = self.colors.text_primary
        p_t2.space_before = Pt(2)

        return slide

    def _render_multi_column_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render 3 distinct foundational pillar columns with accent headers."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Foundational Pillars")
        category = plan.get("category", "FOUNDATIONAL PILLARS")
        bullets: List[BulletPoint] = plan.get("bullets", [])

        self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        count = min(len(bullets), 3)
        col_w = Inches(3.64)
        gap = Inches(0.4)
        start_x = Inches(0.8)
        top_y = Inches(1.8)
        col_h = Inches(4.75)

        accent_colors = [
            self.colors.primary if not self.colors.is_dark else self.colors.accent,
            self.colors.secondary,
            self.colors.accent
        ]

        for idx in range(count):
            b = bullets[idx]
            card_x = start_x + idx * (col_w + gap)

            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, top_y, col_w, col_h)
            card.fill.solid()
            card.fill.fore_color.rgb = self.colors.card_bg
            card.line.color.rgb = self.colors.card_border
            card.line.width = Pt(1)

            # Top Accent Header Strip
            bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, top_y, col_w, Inches(0.12))
            bar.fill.solid()
            bar.fill.fore_color.rgb = accent_colors[idx % len(accent_colors)]
            bar.line.fill.background()

            # Pillar Badge (e.g. PILLAR 01)
            tb_badge = slide.shapes.add_textbox(card_x + Inches(0.2), top_y + Inches(0.25), col_w - Inches(0.4), Inches(0.4))
            tf_badge = tb_badge.text_frame
            tf_badge.word_wrap = True
            p_b = tf_badge.paragraphs[0]
            p_b.text = f"PILLAR {idx + 1:02d}"
            p_b.font.name = self.layout.FONT_HEADING
            p_b.font.size = Pt(10)
            p_b.font.bold = True
            p_b.font.color.rgb = accent_colors[idx % len(accent_colors)]

            # Pillar Title
            tb_title = slide.shapes.add_textbox(card_x + Inches(0.2), top_y + Inches(0.65), col_w - Inches(0.4), Inches(0.8))
            tf_title = tb_title.text_frame
            tf_title.word_wrap = True
            p_t = tf_title.paragraphs[0]
            p_t.text = b.prefix or f"Component {idx + 1}"
            p_t.font.name = self.layout.FONT_HEADING
            p_t.font.size = Pt(14)
            p_t.font.bold = True
            p_t.font.color.rgb = self.colors.text_primary

            # Divider line inside card
            div = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, card_x + Inches(0.2), top_y + Inches(1.5), col_w - Inches(0.4), Inches(0.02)
            )
            div.fill.solid()
            div.fill.fore_color.rgb = self.colors.card_border
            div.line.fill.background()

            # Body description
            tb_body = slide.shapes.add_textbox(card_x + Inches(0.2), top_y + Inches(1.65), col_w - Inches(0.4), col_h - Inches(1.8))
            tf_body = tb_body.text_frame
            tf_body.word_wrap = True
            p_desc = tf_body.paragraphs[0]
            p_desc.text = b.text
            p_desc.font.name = self.layout.FONT_BODY
            p_desc.font.size = Pt(12)
            p_desc.font.color.rgb = self.colors.text_secondary
            p_desc.line_spacing = 1.25

        return slide

    def _render_discussion_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render seminar discussion & Socratic review slide with inquiry cards and reflection prompt."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Seminar Discussion & Ethical Review")
        category = plan.get("category", "SEMINAR & DISCUSSION")
        bullets: List[BulletPoint] = plan.get("bullets", [])

        self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        # Top Seminar Discussion Banner
        top_banner = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), Inches(1.75), Inches(11.733), Inches(1.15)
        )
        top_banner.fill.solid()
        top_banner.fill.fore_color.rgb = self.colors.card_tint
        top_banner.line.color.rgb = self.colors.secondary
        top_banner.line.width = Pt(1.5)

        tb_tb = slide.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(11.2), Inches(0.95))
        tf_tb = tb_tb.text_frame
        tf_tb.word_wrap = True
        p_bh = tf_tb.paragraphs[0]
        p_bh.text = "SOCRATIC SEMINAR & CRITICAL INQUIRY"
        p_bh.font.name = self.layout.FONT_HEADING
        p_bh.font.size = Pt(9.5)
        p_bh.font.bold = True
        p_bh.font.color.rgb = self.colors.accent if self.colors.is_dark else self.colors.primary

        p_bt = tf_tb.add_paragraph()
        p_bt.text = "Evaluate architectural tradeoffs, societal implications, and rigorous engineering responsibilities."
        p_bt.font.name = self.layout.FONT_BODY
        p_bt.font.size = Pt(13)
        p_bt.font.color.rgb = self.colors.text_primary
        p_bt.space_before = Pt(3)

        # Discussion Cards: 2 or 3 question cards side-by-side
        q_count = min(len(bullets), 3) if bullets else 2
        card_w = (Inches(11.733) - (Inches(0.4) * (q_count - 1))) / q_count
        card_y = Inches(3.05)
        card_h = Inches(2.55)

        for q_idx in range(q_count):
            q_x = Inches(0.8) + q_idx * (card_w + Inches(0.4))
            q_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, q_x, card_y, card_w, card_h)
            q_card.fill.solid()
            q_card.fill.fore_color.rgb = self.colors.card_bg
            q_card.line.color.rgb = self.colors.card_border
            q_card.line.width = Pt(1)

            # Left vertical indicator
            q_ind = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, q_x, card_y, Inches(0.1), card_h)
            q_ind.fill.solid()
            q_ind.fill.fore_color.rgb = self.colors.accent if q_idx % 2 == 0 else self.colors.secondary
            q_ind.line.fill.background()

            tb_q = slide.shapes.add_textbox(q_x + Inches(0.25), card_y + Inches(0.15), card_w - Inches(0.4), card_h - Inches(0.3))
            tf_q = tb_q.text_frame
            tf_q.word_wrap = True

            p_qh = tf_q.paragraphs[0]
            p_qh.text = f"DISCUSSION TOPIC {q_idx + 1:02d}"
            p_qh.font.name = self.layout.FONT_HEADING
            p_qh.font.size = Pt(10)
            p_qh.font.bold = True
            p_qh.font.color.rgb = self.colors.accent if q_idx % 2 == 0 else self.colors.secondary

            b_item = bullets[q_idx] if q_idx < len(bullets) else None
            p_qtitle = tf_q.add_paragraph()
            p_qtitle.text = (b_item.prefix if b_item and b_item.prefix else f"Question {q_idx + 1}")
            p_qtitle.font.name = self.layout.FONT_HEADING
            p_qtitle.font.size = Pt(13)
            p_qtitle.font.bold = True
            p_qtitle.font.color.rgb = self.colors.text_primary
            p_qtitle.space_before = Pt(3)

            p_qbody = tf_q.add_paragraph()
            p_qbody.text = b_item.text if b_item else "Synthesize arguments and propose concrete mitigations."
            p_qbody.font.name = self.layout.FONT_BODY
            p_qbody.font.size = Pt(11)
            p_qbody.font.color.rgb = self.colors.text_secondary
            p_qbody.space_before = Pt(5)
            p_qbody.line_spacing = 1.2

        # Bottom Cohort Reflection Callout
        cohort_y = Inches(5.75)
        cohort = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), cohort_y, Inches(11.733), Inches(0.8)
        )
        cohort.fill.solid()
        cohort.fill.fore_color.rgb = self.colors.card_tint
        cohort.line.color.rgb = self.colors.secondary
        cohort.line.width = Pt(1)

        tb_coh = slide.shapes.add_textbox(Inches(1.05), cohort_y + Inches(0.12), Inches(11.2), Inches(0.6))
        tf_coh = tb_coh.text_frame
        tf_coh.word_wrap = True
        p_coh = tf_coh.paragraphs[0]
        p_coh.text = "Cohort Exercise: Break into discussion pods and formulate actionable governance standards."
        p_coh.font.name = self.layout.FONT_BODY
        p_coh.font.size = Pt(11)
        p_coh.font.bold = True
        p_coh.font.color.rgb = self.colors.primary

        return slide

    def _render_cards_grid_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render content in a structured 2x2 or 3-card grid container with top accent bars."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Lecture Topic")
        category = plan.get("category", "CARDS & GRID")
        bullets: List[BulletPoint] = plan.get("bullets", [])

        rule_top = self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        count = len(bullets)
        if count == 0:
            return slide

        header_rule = rule_top if rule_top is not None else Inches(1.45)
        card_top = max(Inches(1.8), header_rule + Inches(0.25))
        card_h = max(Inches(3.8), Inches(6.65) - card_top)

        if count == 1:
            # Single Concept: render expansive full-width card container balancing canvas without empty void
            col_w = Inches(11.733)
            start_x = Inches(0.8)
            b = bullets[0]

            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, start_x, card_top, col_w, card_h)
            card.fill.solid()
            card.fill.fore_color.rgb = self.colors.card_bg
            card.line.color.rgb = self.colors.card_border
            card.line.width = Pt(1.5)

            # Top accent strip
            strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, start_x, card_top, col_w, Inches(0.14))
            strip.fill.solid()
            strip.fill.fore_color.rgb = self.colors.accent
            strip.line.fill.background()

            # Concept Header Box
            tb_h = slide.shapes.add_textbox(start_x + Inches(0.4), card_top + Inches(0.35), col_w - Inches(0.8), Inches(0.8))
            tf_h = tb_h.text_frame
            tf_h.word_wrap = True
            p_h = tf_h.paragraphs[0]
            p_h.text = b.prefix or "Core Architectural Principle"
            p_h.font.name = self.layout.FONT_HEADING
            p_h.font.size = Pt(20)
            p_h.font.bold = True
            p_h.font.color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent

            # Concept Body Text
            tb_b = slide.shapes.add_textbox(start_x + Inches(0.4), card_top + Inches(1.25), col_w - Inches(0.8), card_h - Inches(1.5))
            tf_b = tb_b.text_frame
            tf_b.word_wrap = True
            p_b = tf_b.paragraphs[0]
            p_b.text = b.text
            p_b.font.name = self.layout.FONT_BODY
            p_b.font.size = Pt(14)
            p_b.font.color.rgb = self.colors.text_primary
            p_b.line_spacing = 1.3

        elif count == 2:
            # 2 cards balanced across the entire slide width
            col_w = Inches(5.65)
            gap = Inches(0.433)
            start_x = Inches(0.8)

            for idx, b in enumerate(bullets[:2]):
                card_x = start_x + idx * (col_w + gap)
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_top, col_w, card_h)
                card.fill.solid()
                card.fill.fore_color.rgb = self.colors.card_bg
                card.line.color.rgb = self.colors.card_border
                card.line.width = Pt(1)

                # Top accent strip
                strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_top, col_w, Inches(0.12))
                strip.fill.solid()
                strip.fill.fore_color.rgb = self.colors.accent if idx == 0 else self.colors.secondary
                strip.line.fill.background()

                # Card Header Box
                tb_h = slide.shapes.add_textbox(card_x + Inches(0.3), card_top + Inches(0.35), col_w - Inches(0.6), Inches(0.75))
                tf_h = tb_h.text_frame
                tf_h.word_wrap = True
                p_h = tf_h.paragraphs[0]
                p_h.text = b.prefix or f"Concept {idx + 1:02d}"
                p_h.font.name = self.layout.FONT_HEADING
                p_h.font.size = Pt(16)
                p_h.font.bold = True
                p_h.font.color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent

                # Card Body Text
                tb_b = slide.shapes.add_textbox(card_x + Inches(0.3), card_top + Inches(1.2), col_w - Inches(0.6), card_h - Inches(1.5))
                tf_b = tb_b.text_frame
                tf_b.word_wrap = True
                p_b = tf_b.paragraphs[0]
                p_b.text = b.text
                p_b.font.name = self.layout.FONT_BODY
                p_b.font.size = Pt(13)
                p_b.font.color.rgb = self.colors.text_secondary
                p_b.line_spacing = 1.25

        elif count == 3:
            # 3 cards across (1 row)
            col_w = Inches(3.64)
            gap = Inches(0.40)
            start_x = Inches(0.8)

            for idx, b in enumerate(bullets[:3]):
                card_x = start_x + idx * (col_w + gap)
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_top, col_w, card_h)
                card.fill.solid()
                card.fill.fore_color.rgb = self.colors.card_bg
                card.line.color.rgb = self.colors.card_border
                card.line.width = Pt(1)

                # Top accent strip
                strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_top, col_w, Inches(0.12))
                strip.fill.solid()
                strip.fill.fore_color.rgb = self.colors.accent if idx == 0 else self.colors.secondary
                strip.line.fill.background()

                # Card Header Box (Prefix or Number)
                tb_h = slide.shapes.add_textbox(card_x + Inches(0.2), card_top + Inches(0.3), col_w - Inches(0.4), Inches(0.7))
                tf_h = tb_h.text_frame
                tf_h.word_wrap = True
                p_h = tf_h.paragraphs[0]
                p_h.text = b.prefix or f"Concept {idx + 1:02d}"
                p_h.font.name = self.layout.FONT_HEADING
                p_h.font.size = Pt(14)
                p_h.font.bold = True
                p_h.font.color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent

                # Card Body Text
                tb_b = slide.shapes.add_textbox(card_x + Inches(0.2), card_top + Inches(1.1), col_w - Inches(0.4), card_h - Inches(1.3))
                tf_b = tb_b.text_frame
                tf_b.word_wrap = True
                p_b = tf_b.paragraphs[0]
                p_b.text = b.text
                p_b.font.name = self.layout.FONT_BODY
                p_b.font.size = Pt(11.5)
                p_b.font.color.rgb = self.colors.text_secondary
                p_b.line_spacing = 1.2
        else:
            # 4 cards in 2x2 grid
            col_w = Inches(5.65)
            row_h = (card_h - Inches(0.3)) / 2
            gap_x = Inches(0.433)
            gap_y = Inches(0.3)
            start_x = Inches(0.8)
            start_y = card_top

            for idx, b in enumerate(bullets[:4]):
                col_idx = idx % 2
                row_idx = idx // 2
                card_x = start_x + col_idx * (col_w + gap_x)
                card_y = start_y + row_idx * (row_h + gap_y)

                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y, col_w, row_h)
                card.fill.solid()
                card.fill.fore_color.rgb = self.colors.card_bg
                card.line.color.rgb = self.colors.card_border
                card.line.width = Pt(1)

                # Left accent strip
                strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y, Inches(0.12), row_h)
                strip.fill.solid()
                strip.fill.fore_color.rgb = self.colors.secondary if col_idx == 0 else self.colors.accent
                strip.line.fill.background()

                # Card Text Frame
                tb = slide.shapes.add_textbox(card_x + Inches(0.28), card_y + Inches(0.15), col_w - Inches(0.45), row_h - Inches(0.3))
                tf = tb.text_frame
                tf.word_wrap = True

                p1 = tf.paragraphs[0]
                p1.text = b.prefix or f"Point {idx + 1:02d}"
                p1.font.name = self.layout.FONT_HEADING
                p1.font.size = Pt(13)
                p1.font.bold = True
                p1.font.color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent

                p2 = tf.add_paragraph()
                p2.text = b.text
                p2.font.name = self.layout.FONT_BODY
                card_font_pt = 10.5 if len(b.text) > 130 else 11.0
                p2.font.size = Pt(card_font_pt)
                p2.font.color.rgb = self.colors.text_secondary
                p2.space_before = Pt(3 if card_font_pt < 11.0 else 4)
                p2.line_spacing = 1.1 if card_font_pt < 11.0 else 1.15

        return slide

    def _render_two_column_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render content in balanced dual columns with distinct accent banners."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Comparative Analysis")
        category = plan.get("category", "TWO-COLUMN COMPARISON")
        bullets: List[BulletPoint] = plan.get("bullets", [])

        rule_top = self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        mid = (len(bullets) + 1) // 2
        left_bullets = bullets[:mid]
        right_bullets = bullets[mid:]

        col_w = Inches(5.6)
        gap = Inches(0.53)
        top_y = max(Inches(1.8), (rule_top if rule_top is not None else Inches(1.45)) + Inches(0.25))
        col_h = max(Inches(3.8), Inches(6.65) - top_y)

        # Left Column Card
        left_x = Inches(0.8)
        left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, top_y, col_w, col_h)
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = self.colors.card_bg
        left_card.line.color.rgb = self.colors.card_border
        left_card.line.width = Pt(1)

        # Left Top Banner
        l_banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, top_y, col_w, Inches(0.12))
        l_banner.fill.solid()
        l_banner.fill.fore_color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent
        l_banner.line.fill.background()

        self._render_bullets_in_box(slide, left_x + Inches(0.3), top_y + Inches(0.3), col_w - Inches(0.6), col_h - Inches(0.5), left_bullets)

        # Right Column Card
        right_x = left_x + col_w + gap
        right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, top_y, col_w, col_h)
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = self.colors.card_bg
        right_card.line.color.rgb = self.colors.card_border
        right_card.line.width = Pt(1)

        # Right Top Banner
        r_banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, top_y, col_w, Inches(0.12))
        r_banner.fill.solid()
        r_banner.fill.fore_color.rgb = self.colors.secondary
        r_banner.line.fill.background()

        self._render_bullets_in_box(slide, right_x + Inches(0.3), top_y + Inches(0.3), col_w - Inches(0.6), col_h - Inches(0.5), right_bullets)

        return slide

    def _render_metrics_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render high-impact metric statistics cards with large bold numbers."""
        import re
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Key Metrics & Insights")
        category = plan.get("category", "KEY METRICS & HIGHLIGHTS")
        bullets: List[BulletPoint] = plan.get("bullets", [])

        rule_top = self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        count = min(len(bullets), 4)
        if count == 0:
            return slide

        gap = Inches(0.35)
        col_w = (Inches(11.733) - (gap * (count - 1))) / count
        top_y = max(Inches(1.8), (rule_top if rule_top is not None else Inches(1.45)) + Inches(0.25))
        card_h = max(Inches(3.8), Inches(6.65) - top_y)

        for idx, b in enumerate(bullets[:count]):
            card_x = Inches(0.8) + idx * (col_w + gap)

            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, top_y, col_w, card_h)
            card.fill.solid()
            card.fill.fore_color.rgb = self.colors.card_bg
            card.line.color.rgb = self.colors.card_border
            card.line.width = Pt(1)

            # Top accent bar
            bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, top_y, col_w, Inches(0.1))
            bar.fill.solid()
            bar.fill.fore_color.rgb = self.colors.accent if idx % 2 == 0 else self.colors.secondary
            bar.line.fill.background()

            # Extract metric number from text or prefix (e.g. 99%, 1950, 10x, Phase 1, 50-year)
            num_match = re.search(r"(\d+[%xXkKMmB]?|\d+[\-–]\w+|\d{4})", b.prefix or b.text)
            metric_val = num_match.group(1) if num_match else f"{idx + 1:02d}"

            # Metric Number Box
            tb_num = slide.shapes.add_textbox(card_x + Inches(0.15), top_y + Inches(0.3), col_w - Inches(0.3), Inches(0.9))
            tf_num = tb_num.text_frame
            tf_num.word_wrap = True
            p_num = tf_num.paragraphs[0]
            p_num.text = metric_val
            p_num.font.name = self.layout.FONT_HEADING
            p_num.font.size = Pt(32)
            p_num.font.bold = True
            p_num.font.color.rgb = self.colors.accent if idx % 2 == 0 else (self.colors.primary if not self.colors.is_dark else self.colors.accent)

            # Metric Label
            tb_lbl = slide.shapes.add_textbox(card_x + Inches(0.18), top_y + Inches(1.25), col_w - Inches(0.36), Inches(0.7))
            tf_lbl = tb_lbl.text_frame
            tf_lbl.word_wrap = True
            p_lbl = tf_lbl.paragraphs[0]
            p_lbl.text = b.prefix or f"Metric {idx + 1}"
            p_lbl.font.name = self.layout.FONT_HEADING
            p_lbl.font.size = Pt(13)
            p_lbl.font.bold = True
            p_lbl.font.color.rgb = self.colors.text_primary

            # Description Box
            tb_desc = slide.shapes.add_textbox(card_x + Inches(0.18), top_y + Inches(2.0), col_w - Inches(0.36), card_h - Inches(2.15))
            tf_desc = tb_desc.text_frame
            tf_desc.word_wrap = True
            p_desc = tf_desc.paragraphs[0]
            p_desc.text = b.text
            p_desc.font.name = self.layout.FONT_BODY
            p_desc.font.size = Pt(11)
            p_desc.font.color.rgb = self.colors.text_secondary
            p_desc.line_spacing = 1.2

        return slide

    def _fit_text_box(
        self,
        text: str,
        max_pt: float = 14.0,
        min_pt: float = 10.5,
        target_chars: int = 150,
    ) -> float:
        """Dynamically compute font size based on text length to prevent clipping."""
        length = len(text.strip())
        if length <= target_chars:
            return max_pt
        ratio = min(1.0, (length - target_chars) / 250.0)
        return max(min_pt, max_pt - (ratio * (max_pt - min_pt)))

    def _render_pull_quote_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render a high-typography editorial pull-quote slide."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Foundational Principle")
        category = plan.get("category", "CORE PEDAGOGICAL AXIOM")
        self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        bullets: List[BulletPoint] = plan.get("bullets", [])
        quote_text = plan.get("key_message") or (bullets[0].text if bullets else title)
        attribution = (bullets[0].prefix if bullets and bullets[0].prefix else "Foundational Axiom")

        # Top Hero Quote Card
        self.layout.create_card_container(
            slide,
            left=Inches(0.8),
            top=Inches(1.8),
            width=Inches(11.733),
            height=Inches(3.2),
            accent_color=self.colors.accent,
            bg_color=self.colors.card_tint,
        )

        # Large quotation mark graphic
        q_mark = slide.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(1.2), Inches(1.2))
        tf_q = q_mark.text_frame
        tf_q.word_wrap = False
        p_q = tf_q.paragraphs[0]
        p_q.text = "\u201C"
        p_q.font.name = "Georgia"
        p_q.font.size = Pt(72)
        p_q.font.bold = True
        p_q.font.color.rgb = self.colors.accent

        # Quote Text
        q_box = slide.shapes.add_textbox(Inches(2.2), Inches(2.1), Inches(9.8), Inches(2.0))
        tf_text = q_box.text_frame
        tf_text.word_wrap = True
        p_body = tf_text.paragraphs[0]
        p_body.text = quote_text
        p_body.font.name = self.layout.FONT_HEADING
        font_sz = self._fit_text_box(quote_text, max_pt=22.0, min_pt=16.0, target_chars=110)
        p_body.font.size = Pt(font_sz)
        p_body.font.italic = True
        p_body.font.bold = True
        p_body.font.color.rgb = self.colors.primary

        # Attribution line
        p_attr = tf_text.add_paragraph()
        p_attr.text = f"\u2014 {attribution}"
        p_attr.font.name = self.layout.FONT_BODY
        p_attr.font.size = Pt(13)
        p_attr.font.bold = True
        p_attr.font.color.rgb = self.colors.text_secondary
        p_attr.space_before = Pt(12)

        # Bottom Reflection / Implications Card
        sub_bullets = bullets[1:] if len(bullets) > 1 else bullets
        if sub_bullets:
            self.layout.create_card_container(
                slide,
                left=Inches(0.8),
                top=Inches(5.2),
                width=Inches(11.733),
                height=Inches(1.5),
                accent_color=self.colors.secondary,
            )
            tb_sub = slide.shapes.add_textbox(Inches(1.1), Inches(5.3), Inches(11.2), Inches(1.3))
            tf_sub = tb_sub.text_frame
            tf_sub.word_wrap = True
            for i, b in enumerate(sub_bullets[:2]):
                p_s = tf_sub.paragraphs[0] if i == 0 else tf_sub.add_paragraph()
                p_s.text = f"\u2022 {b.prefix}: {b.text}" if b.prefix else f"\u2022 {b.text}"
                p_s.font.name = self.layout.FONT_BODY
                p_s.font.size = Pt(11.5)
                p_s.font.color.rgb = self.colors.text_primary
                p_s.space_before = Pt(4) if i > 0 else Pt(0)

        return slide

    def _render_split_image_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render a split layout with diagram/visual on left and takeaways on right."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "System Architecture & Visual Model")
        category = plan.get("category", "VISUAL SYSTEM ARCHITECTURE")
        bullets: List[BulletPoint] = plan.get("bullets", [])

        self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        # Left 45% Visual Canvas Box
        self.layout.create_card_container(
            slide,
            left=Inches(0.8),
            top=Inches(1.8),
            width=Inches(5.2),
            height=Inches(4.9),
            accent_color=self.colors.accent,
            bg_color=self.colors.card_tint,
        )

        has_local_img = False
        img_path_str = plan.get("image_path")
        local_img_candidate = Path(img_path_str) if img_path_str else None
        caption_text = plan.get("caption") or ""

        if local_img_candidate and local_img_candidate.exists():
            try:
                from PIL import Image as PILImage
                with PILImage.open(local_img_candidate) as p_img:
                    orig_w, orig_h = p_img.size

                has_caption = bool(caption_text and caption_text.strip())
                max_box_w = Inches(4.8)
                max_box_h = Inches(3.9) if has_caption else Inches(4.5)
                aspect = orig_w / float(max(1, orig_h))

                if aspect >= (max_box_w / max_box_h):
                    render_w = max_box_w
                    render_h = render_w / aspect
                else:
                    render_h = max_box_h
                    render_w = render_h * aspect

                center_x = Inches(1.0) + (max_box_w - render_w) / 2
                center_y = Inches(1.95) + (max_box_h - render_h) / 2

                slide.shapes.add_picture(
                    str(local_img_candidate),
                    center_x, center_y, width=render_w, height=render_h
                )
                has_local_img = True

                # Render caption underneath the image
                if has_caption:
                    cap_y = Inches(1.95) + max_box_h + Inches(0.08)
                    cap_tb = slide.shapes.add_textbox(Inches(0.9), cap_y, Inches(5.0), Inches(0.55))
                    tf_cap = cap_tb.text_frame
                    tf_cap.word_wrap = True
                    tf_cap.margin_left = Inches(0.05)
                    tf_cap.margin_right = Inches(0.05)
                    tf_cap.margin_top = Inches(0.02)
                    p_cap = tf_cap.paragraphs[0]
                    p_cap.text = caption_text[:110]
                    p_cap.font.name = self.layout.FONT_BODY
                    p_cap.font.size = Pt(10)
                    p_cap.font.italic = True
                    p_cap.font.color.rgb = self.colors.text_muted
                    p_cap.alignment = PP_ALIGN.CENTER
            except Exception:
                has_local_img = False

        if not has_local_img:
            # Honest Fallback: Informative Figure Reference Card (never dummy image or fabricated vector)
            fb_box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(1.2), Inches(2.3), Inches(4.4), Inches(3.8)
            )
            fb_box.fill.solid()
            fb_box.fill.fore_color.rgb = self.colors.primary
            fb_box.line.color.rgb = self.colors.accent
            fb_box.line.width = Pt(1.5)
            tf_fb = fb_box.text_frame
            tf_fb.word_wrap = True

            p_h = tf_fb.paragraphs[0]
            p_h.text = "DIAGRAM / MODEL REFERENCE"
            p_h.font.name = self.layout.FONT_HEADING
            p_h.font.size = Pt(12)
            p_h.font.bold = True
            p_h.font.color.rgb = self.colors.accent
            p_h.alignment = PP_ALIGN.CENTER

            p_ref = tf_fb.add_paragraph()
            p_ref.text = caption_text or title
            p_ref.font.name = self.layout.FONT_BODY
            p_ref.font.size = Pt(11)
            p_ref.font.bold = True
            p_ref.font.color.rgb = RGBColor(255, 255, 255)
            p_ref.space_before = Pt(8)
            p_ref.alignment = PP_ALIGN.CENTER

            p_note = tf_fb.add_paragraph()
            p_note.text = "Original source diagram could not be extracted as an image asset from document stream. Conceptual specifications and model elements are detailed on the right."
            p_note.font.name = self.layout.FONT_BODY
            p_note.font.size = Pt(9.5)
            p_note.font.color.rgb = RGBColor(226, 232, 240)
            p_note.space_before = Pt(12)
            p_note.alignment = PP_ALIGN.CENTER

        # Right 55% Structured Takeaways Box
        right_left = Inches(6.3)
        right_width = Inches(6.233)

        takeaway = plan.get("key_message") or (bullets[0].text if bullets else "Architectural analysis")
        self.layout.create_card_container(
            slide,
            left=right_left,
            top=Inches(1.8),
            width=right_width,
            height=Inches(1.2),
            accent_color=self.colors.secondary,
        )
        tb_tw = slide.shapes.add_textbox(right_left + Inches(0.2), Inches(1.9), right_width - Inches(0.4), Inches(1.0))
        tf_tw = tb_tw.text_frame
        tf_tw.word_wrap = True
        p_tw = tf_tw.paragraphs[0]
        p_tw.text = f"KEY INSIGHT: {takeaway}"
        p_tw.font.name = self.layout.FONT_HEADING
        p_tw.font.size = Pt(13)
        p_tw.font.bold = True
        p_tw.font.color.rgb = self.colors.primary

        # Detail Bullets
        self.layout.create_card_container(
            slide,
            left=right_left,
            top=Inches(3.2),
            width=right_width,
            height=Inches(3.5),
            accent_color=self.colors.accent,
        )
        tb_bullets = slide.shapes.add_textbox(right_left + Inches(0.25), Inches(3.35), right_width - Inches(0.5), Inches(3.2))
        tf_b = tb_bullets.text_frame
        tf_b.word_wrap = True

        display_bullets = bullets[:4]
        for i, b in enumerate(display_bullets):
            p = tf_b.paragraphs[0] if i == 0 else tf_b.add_paragraph()
            font_size = self._fit_text_box(b.text, max_pt=12.5, min_pt=10.5)
            if b.prefix:
                run_p = p.add_run()
                run_p.text = f"{b.prefix}: "
                run_p.font.name = self.layout.FONT_HEADING
                run_p.font.bold = True
                run_p.font.size = Pt(font_size)
                run_p.font.color.rgb = self.colors.primary

            run_t = p.add_run()
            run_t.text = b.text
            run_t.font.name = self.layout.FONT_BODY
            run_t.font.size = Pt(font_size)
            run_t.font.color.rgb = self.colors.text_primary
            p.space_before = Pt(8) if i > 0 else Pt(0)
            p.line_spacing = 1.15

        return slide

    def _render_quadrant_matrix_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render a 2x2 decision or tradeoff matrix slide."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Decision Matrix & Architecture Tradeoffs")
        category = plan.get("category", "DECISION MATRIX & TRADEOFFS")
        bullets: List[BulletPoint] = plan.get("bullets", [])

        self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        # 2x2 Grid Coordinates
        q_w = Inches(5.7)
        q_h = Inches(2.3)
        lefts = [Inches(0.8), Inches(6.833)]
        tops = [Inches(1.8), Inches(4.4)]

        q_labels = [
            ("Q1: HIGH IMPACT \u2022 IMMEDIATE VELOCITY", self.colors.accent),
            ("Q2: STRATEGIC \u2022 ARCHITECTURAL DEPTH", self.colors.secondary),
            ("Q3: FOUNDATIONAL \u2022 STABLE BASELINE", self.colors.primary),
            ("Q4: SPECIALIZED \u2022 BOUNDED DOMAIN", self.colors.text_secondary),
        ]

        quad_bullets = []
        for i in range(4):
            if i < len(bullets):
                quad_bullets.append(bullets[i])
            else:
                quad_bullets.append(BulletPoint(text="Structured balance criteria across execution constraints.", prefix=f"Pillar {i+1}"))

        coords = [
            (lefts[0], tops[0]),
            (lefts[1], tops[0]),
            (lefts[0], tops[1]),
            (lefts[1], tops[1]),
        ]

        for i, (q_x, q_y) in enumerate(coords):
            lbl, border_col = q_labels[i]
            b_item = quad_bullets[i]

            self.layout.create_card_container(
                slide,
                left=q_x,
                top=q_y,
                width=q_w,
                height=q_h,
                accent_color=border_col,
                bg_color=self.colors.card_tint,
            )

            # Quadrant Badge
            tb_hdr = slide.shapes.add_textbox(q_x + Inches(0.2), q_y + Inches(0.12), q_w - Inches(0.4), Inches(0.4))
            tf_h = tb_hdr.text_frame
            tf_h.word_wrap = True
            p_h = tf_h.paragraphs[0]
            p_h.text = lbl
            p_h.font.name = self.layout.FONT_HEADING
            p_h.font.size = Pt(11)
            p_h.font.bold = True
            p_h.font.color.rgb = border_col

            # Quadrant Body
            tb_body = slide.shapes.add_textbox(q_x + Inches(0.2), q_y + Inches(0.55), q_w - Inches(0.4), q_h - Inches(0.65))
            tf_b = tb_body.text_frame
            tf_b.word_wrap = True

            p_title = tf_b.paragraphs[0]
            p_title.text = b_item.prefix or f"Dimension {i+1}"
            p_title.font.name = self.layout.FONT_HEADING
            p_title.font.size = Pt(13)
            p_title.font.bold = True
            p_title.font.color.rgb = self.colors.text_primary

            p_desc = tf_b.add_paragraph()
            p_desc.text = b_item.text
            p_desc.font.name = self.layout.FONT_BODY
            p_desc.font.size = Pt(11.5)
            p_desc.font.color.rgb = self.colors.text_secondary
            p_desc.space_before = Pt(4)
            p_desc.line_spacing = 1.15

        return slide

    def _render_process_chevrons_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render a sequential horizontal process pipeline slide."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "End-to-End System Execution Pipeline")
        category = plan.get("category", "SYSTEM EXECUTION PIPELINE")
        bullets: List[BulletPoint] = plan.get("bullets", [])

        rule_top = self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        num_steps = max(2, min(4, len(bullets) if bullets else 3))
        display_bullets = bullets[:num_steps] if bullets else [
            BulletPoint(text="Ingest raw signals and perform normalization", prefix="Stage 1"),
            BulletPoint(text="Execute distributed matrix transformations", prefix="Stage 2"),
            BulletPoint(text="Evaluate convergence and return telemetry", prefix="Stage 3"),
        ]

        card_gap = Inches(0.3)
        total_width = Inches(11.733)
        card_w = (total_width - (card_gap * (num_steps - 1))) / num_steps
        card_y = max(Inches(1.85), (rule_top if rule_top is not None else Inches(1.45)) + Inches(0.25))
        card_h = max(Inches(3.8), Inches(6.65) - card_y)

        for idx, b in enumerate(display_bullets):
            card_x = Inches(0.8) + (idx * (card_w + card_gap))

            self.layout.create_card_container(
                slide,
                left=card_x,
                top=card_y,
                width=card_w,
                height=card_h,
                accent_color=self.colors.accent if idx == 0 else self.colors.secondary,
                bg_color=self.colors.card_tint,
            )

            # Step Circle Badge
            step_circle = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                card_x + Inches(0.25), card_y + Inches(0.25), Inches(0.7), Inches(0.7)
            )
            step_circle.fill.solid()
            step_circle.fill.fore_color.rgb = self.colors.primary
            step_circle.line.color.rgb = self.colors.accent
            step_circle.line.width = Pt(1.5)
            tf_c = step_circle.text_frame
            p_c = tf_c.paragraphs[0]
            p_c.text = f"0{idx+1}"
            p_c.font.name = self.layout.FONT_HEADING
            p_c.font.size = Pt(14)
            p_c.font.bold = True
            p_c.font.color.rgb = RGBColor(255, 255, 255)
            p_c.alignment = PP_ALIGN.CENTER

            # Step Header
            tb_hdr = slide.shapes.add_textbox(card_x + Inches(0.2), card_y + Inches(1.15), card_w - Inches(0.4), Inches(0.8))
            tf_h = tb_hdr.text_frame
            tf_h.word_wrap = True
            p_h = tf_h.paragraphs[0]
            p_h.text = b.prefix or f"Phase {idx+1}"
            p_h.font.name = self.layout.FONT_HEADING
            p_h.font.size = Pt(14)
            p_h.font.bold = True
            p_h.font.color.rgb = self.colors.primary

            # Step Description
            tb_body = slide.shapes.add_textbox(card_x + Inches(0.2), card_y + Inches(2.0), card_w - Inches(0.4), card_h - Inches(2.2))
            tf_b = tb_body.text_frame
            tf_b.word_wrap = True
            p_b = tf_b.paragraphs[0]
            p_b.text = b.text
            p_b.font.name = self.layout.FONT_BODY
            p_b.font.size = Pt(self._fit_text_box(b.text, max_pt=12.0, min_pt=10.0))
            p_b.font.color.rgb = self.colors.text_primary
            p_b.line_spacing = 1.2

            # Chevron Arrow between cards (if not last)
            if idx < num_steps - 1:
                arrow = slide.shapes.add_shape(
                    MSO_SHAPE.CHEVRON,
                    card_x + card_w + Inches(0.06), card_y + Inches(2.1), Inches(0.18), Inches(0.4)
                )
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = self.colors.accent
                arrow.line.fill.background()

        return slide

    def _render_section_divider_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render a high-impact chapter/module section divider slide matching active theme."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        is_dark = self.colors.is_dark

        # Elegant section hero card matching theme styling
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(1.2), Inches(1.8), Inches(10.933), Inches(3.8)
        )
        card.name = "DividerCard"
        card.fill.solid()
        card.fill.fore_color.rgb = self.colors.card_bg
        card.line.color.rgb = self.colors.card_border if not is_dark else self.colors.accent
        card.line.width = Pt(1.5)

        # Accent border banner across the middle
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(2.2), Inches(0.15), Inches(3.0))
        banner.fill.solid()
        banner.fill.fore_color.rgb = self.colors.accent
        banner.line.fill.background()

        # Text Frame for Module & Title
        tb = slide.shapes.add_textbox(Inches(1.9), Inches(2.1), Inches(9.8), Inches(3.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p_badge = tf.paragraphs[0]
        p_badge.text = plan.get("category", "CURRICULUM MODULE").upper()
        p_badge.font.name = self.layout.FONT_HEADING
        p_badge.font.size = Pt(13)
        p_badge.font.bold = True
        p_badge.font.color.rgb = self.colors.accent

        p_title = tf.add_paragraph()
        p_title.text = plan.get("title", "Curriculum Section")
        p_title.font.name = self.layout.FONT_HEADING
        p_title.font.size = Pt(32)
        p_title.font.bold = True
        p_title.font.color.rgb = self.colors.primary if not is_dark else self.colors.text_primary
        p_title.space_before = Pt(8)

        if plan.get("subtitle"):
            p_sub = tf.add_paragraph()
            p_sub.text = plan.get("subtitle")
            p_sub.font.name = self.layout.FONT_BODY
            p_sub.font.size = Pt(15)
            p_sub.font.color.rgb = self.colors.text_secondary
            p_sub.space_before = Pt(10)

        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)
        return slide

    def _render_split_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render image/diagram and text split slide."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Applied Systems & Analysis")
        category = plan.get("category", "IMAGE & TEXT SPLIT")
        bullets: List[BulletPoint] = plan.get("bullets", [])

        self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        left_w = Inches(4.8)
        right_w = Inches(6.5)
        gap = Inches(0.43)
        top_y = Inches(1.8)
        col_h = Inches(4.7)

        # Left Column Card (Visual / Diagram Container)
        left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_y, left_w, col_h)
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = self.colors.card_tint
        left_card.line.color.rgb = self.colors.card_border
        left_card.line.width = Pt(1)

        # Left Visual Pillar Elements
        tb_left = slide.shapes.add_textbox(Inches(1.1), top_y + Inches(0.4), left_w - Inches(0.6), Inches(0.8))
        tf_l = tb_left.text_frame
        tf_l.word_wrap = True
        p_l = tf_l.paragraphs[0]
        p_l.text = "CORE ARCHITECTURE"
        p_l.font.name = self.layout.FONT_HEADING
        p_l.font.size = Pt(13)
        p_l.font.bold = True
        p_l.font.color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent

        # 3 decorative conceptual pillar blocks
        block_h = Inches(0.8)
        block_y = top_y + Inches(1.4)
        for b_idx in range(min(3, max(1, len(bullets)))):
            blk = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), block_y + b_idx * Inches(0.95), left_w - Inches(0.6), block_h)
            blk.fill.solid()
            blk.fill.fore_color.rgb = self.colors.card_bg
            blk.line.color.rgb = self.colors.secondary
            blk.line.width = Pt(1)
            tf_blk = blk.text_frame
            tf_blk.word_wrap = True
            p_blk = tf_blk.paragraphs[0]
            label = bullets[b_idx].prefix if b_idx < len(bullets) and bullets[b_idx].prefix else f"Pillar {b_idx + 1:02d}"
            p_blk.text = label
            p_blk.font.name = self.layout.FONT_HEADING
            p_blk.font.size = Pt(11.5)
            p_blk.font.bold = True
            p_blk.font.color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent
            p_blk.alignment = PP_ALIGN.CENTER

        # Right Column Card (Text & Takeaways)
        right_x = Inches(0.8) + left_w + gap
        self.layout.create_card_container(
            slide,
            left=right_x,
            top=top_y,
            width=right_w,
            height=col_h,
            accent_color=self.colors.secondary,
        )
        self._render_bullets_in_box(
            slide,
            left=right_x + Inches(0.35),
            top=top_y + Inches(0.3),
            width=right_w - Inches(0.6),
            height=col_h - Inches(0.6),
            bullets=bullets,
        )

        return slide

    def _render_summary_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int
    ) -> Slide:
        """Render Conclusion & Key Takeaways slide with varied content-aware compositions."""
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        sec: Section = plan["section"]
        title = sec.title if sec else "Summary & Future Outlook"
        summary_style = (plan.get("summary_style") or getattr(data, "conclusion_composition", "synthesis_takeaways") or "synthesis_takeaways").lower().strip()

        self.layout.add_header(slide, title=title, category="COURSE CONCLUSION")
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        bullets = sec.bullets if sec else []

        if summary_style == "discussion_roadmap":
            # Dual-column synthesis: Core Takeaways vs Future Horizons
            mid = (len(bullets) + 1) // 2
            left_pts = bullets[:mid]
            right_pts = bullets[mid:] if mid < len(bullets) else [
                BulletPoint(text="Review supplementary readings and complete analytical coursework.", prefix="Immediate Actions"),
                BulletPoint(text="Formulate research inquiries around unresolved theoretical boundaries.", prefix="Future Research"),
            ]

            col_w = Inches(5.65)
            gap = Inches(0.43)
            col_h = Inches(4.75)
            top_y = Inches(1.8)

            # Left Panel: Core Retrospective
            left_x = Inches(0.8)
            left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, top_y, col_w, col_h)
            left_card.fill.solid()
            left_card.fill.fore_color.rgb = self.colors.card_bg
            left_card.line.color.rgb = self.colors.card_border
            left_card.line.width = Pt(1)

            l_strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, top_y, col_w, Inches(0.12))
            l_strip.fill.solid()
            l_strip.fill.fore_color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent
            l_strip.line.fill.background()

            tb_lh = slide.shapes.add_textbox(left_x + Inches(0.25), top_y + Inches(0.25), col_w - Inches(0.5), Inches(0.5))
            tf_lh = tb_lh.text_frame
            p_lh = tf_lh.paragraphs[0]
            p_lh.text = "CURRICULUM SYNTHESIS"
            p_lh.font.name = self.layout.FONT_HEADING
            p_lh.font.size = Pt(11)
            p_lh.font.bold = True
            p_lh.font.color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent

            self._render_bullets_in_box(
                slide,
                left=left_x + Inches(0.25),
                top=top_y + Inches(0.75),
                width=col_w - Inches(0.5),
                height=col_h - Inches(0.9),
                bullets=left_pts,
            )

            # Right Panel: Future Horizons & Actions
            right_x = left_x + col_w + gap
            right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, top_y, col_w, col_h)
            right_card.fill.solid()
            right_card.fill.fore_color.rgb = self.colors.card_bg
            right_card.line.color.rgb = self.colors.card_border
            right_card.line.width = Pt(1)

            r_strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, top_y, col_w, Inches(0.12))
            r_strip.fill.solid()
            r_strip.fill.fore_color.rgb = self.colors.secondary
            r_strip.line.fill.background()

            tb_rh = slide.shapes.add_textbox(right_x + Inches(0.25), top_y + Inches(0.25), col_w - Inches(0.5), Inches(0.5))
            tf_rh = tb_rh.text_frame
            p_rh = tf_rh.paragraphs[0]
            p_rh.text = "FUTURE HORIZONS & ACTIONS"
            p_rh.font.name = self.layout.FONT_HEADING
            p_rh.font.size = Pt(11)
            p_rh.font.bold = True
            p_rh.font.color.rgb = self.colors.secondary

            self._render_bullets_in_box(
                slide,
                left=right_x + Inches(0.25),
                top=top_y + Inches(0.75),
                width=col_w - Inches(0.5),
                height=col_h - Inches(0.9),
                bullets=right_pts,
            )

        elif summary_style == "synthesis_takeaways" and len(bullets) >= 3:
            # High-impact 3-pillar takeaway layout
            # Top synthesis narrative
            top_h = Inches(1.15)
            top_card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(11.733), top_h
            )
            top_card.fill.solid()
            top_card.fill.fore_color.rgb = self.colors.card_tint
            top_card.line.color.rgb = self.colors.secondary
            top_card.line.width = Pt(1)

            tb_top = slide.shapes.add_textbox(Inches(1.05), Inches(1.85), Inches(11.2), Inches(0.95))
            tf_top = tb_top.text_frame
            tf_top.word_wrap = True
            p_th = tf_top.paragraphs[0]
            p_th.text = "EXECUTIVE CURRICULUM SYNTHESIS"
            p_th.font.name = self.layout.FONT_HEADING
            p_th.font.size = Pt(9.5)
            p_th.font.bold = True
            p_th.font.color.rgb = self.colors.accent if self.colors.is_dark else self.colors.primary

            p_tt = tf_top.add_paragraph()
            p_tt.text = "Mastery of foundational theory, rigorous architectural evaluation, and proactive governance."
            p_tt.font.name = self.layout.FONT_BODY
            p_tt.font.size = Pt(12.5)
            p_tt.font.color.rgb = self.colors.text_primary
            p_tt.space_before = Pt(3)

            # 3 Takeaway Cards
            card_w = Inches(3.64)
            gap = Inches(0.4)
            top_y = Inches(3.1)
            card_h = Inches(2.6)
            for idx in range(3):
                b = bullets[idx]
                card_x = Inches(0.8) + idx * (card_w + gap)

                c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, top_y, card_w, card_h)
                c.fill.solid()
                c.fill.fore_color.rgb = self.colors.card_bg
                c.line.color.rgb = self.colors.card_border
                c.line.width = Pt(1)

                bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, top_y, card_w, Inches(0.1))
                bar.fill.solid()
                bar.fill.fore_color.rgb = self.colors.accent if idx % 2 == 0 else self.colors.secondary
                bar.line.fill.background()

                tb_c = slide.shapes.add_textbox(card_x + Inches(0.2), top_y + Inches(0.2), card_w - Inches(0.4), card_h - Inches(0.3))
                tf_c = tb_c.text_frame
                tf_c.word_wrap = True
                p_ch = tf_c.paragraphs[0]
                p_ch.text = b.prefix or f"Takeaway {idx + 1:02d}"
                p_ch.font.name = self.layout.FONT_HEADING
                p_ch.font.size = Pt(12)
                p_ch.font.bold = True
                p_ch.font.color.rgb = self.colors.primary if not self.colors.is_dark else self.colors.accent

                p_cb = tf_c.add_paragraph()
                p_cb.text = b.text
                p_cb.font.name = self.layout.FONT_BODY
                p_cb.font.size = Pt(11)
                p_cb.font.color.rgb = self.colors.text_secondary
                p_cb.space_before = Pt(6)
                p_cb.line_spacing = 1.2

            # Bottom Callout
            bot_y = Inches(5.9)
            callout = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), bot_y, Inches(11.733), Inches(0.7)
            )
            callout.fill.solid()
            callout.fill.fore_color.rgb = self.colors.card_tint
            callout.line.color.rgb = self.colors.secondary
            callout.line.width = Pt(1)

            tb_c = callout.text_frame
            tb_c.word_wrap = True
            p_co = tb_c.paragraphs[0]
            p_co.text = "Next Steps: Complete curriculum seminar problem sets and prepare for subsequent module discussions."
            p_co.font.name = self.layout.FONT_BODY
            p_co.font.size = Pt(10.5)
            p_co.font.bold = True
            p_co.font.color.rgb = self.colors.primary

        else:
            # Classic balanced synthesis container
            main_height = Inches(3.6)
            self.layout.create_card_container(
                slide,
                left=Inches(0.8),
                top=Inches(1.7),
                width=Inches(11.733),
                height=main_height,
                accent_color=self.colors.accent,
            )

            self._render_bullets_in_box(
                slide,
                left=Inches(1.2),
                top=Inches(1.95),
                width=Inches(11.0),
                height=main_height - Inches(0.5),
                bullets=sec.bullets[:4],
            )

            callout = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.8), Inches(5.5), Inches(11.733), Inches(1.0)
            )
            callout.fill.solid()
            callout.fill.fore_color.rgb = self.colors.card_tint
            callout.line.color.rgb = self.colors.secondary
            callout.line.width = Pt(1)

            tf_c = callout.text_frame
            tf_c.word_wrap = True
            tf_c.margin_left = Inches(0.3)
            tf_c.margin_top = Inches(0.18)

            p_c1 = tf_c.paragraphs[0]
            p_c1.text = "Key Takeaway & Next Actions:"
            p_c1.font.name = self.layout.FONT_HEADING
            p_c1.font.size = Pt(11)
            p_c1.font.bold = True
            p_c1.font.color.rgb = self.colors.primary

            p_c2 = tf_c.add_paragraph()
            p_c2.text = "Integrate these concepts into practical assignments. Review subsequent modules for advanced application."
            p_c2.font.name = self.layout.FONT_BODY
            p_c2.font.size = Pt(10.5)
            p_c2.font.color.rgb = self.colors.text_secondary

        return slide

    def _render_visual_slide(
        self,
        prs: Presentation,
        data: PresentationData,
        plan: Dict[str, Any],
        total_slides: int,
    ) -> Slide:
        """Render a native visual diagram or image slide with graceful text fallback."""
        visual: Optional[SlideVisual] = plan.get("visual")
        if not visual or visual.visual_type == VisualType.NONE:
            return self._render_content_slide(prs, data, plan, total_slides)

        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)
        self.layout.apply_canvas_background(slide)

        title = plan.get("title", "Lecture Topic")
        category = plan.get("category", "CURRICULUM VISUAL")

        self.layout.add_header(slide, title=title, category=category)
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        bounds = RectBounds(
            left=Inches(0.8),
            top=Inches(1.8),
            width=Inches(11.733),
            height=Inches(4.8),
        )

        try:
            rendered = False
            vtype = visual.visual_type

            if vtype == VisualType.PROCESS_FLOW and visual.process_flow:
                rendered = self.diagram_renderer.render_process_flow(slide, visual.process_flow, bounds)
            elif vtype == VisualType.TIMELINE and visual.timeline:
                rendered = self.diagram_renderer.render_timeline(slide, visual.timeline, bounds)
            elif vtype == VisualType.COMPARISON and visual.comparison:
                rendered = self.diagram_renderer.render_comparison(slide, visual.comparison, bounds)
            elif vtype == VisualType.CONCEPT_MAP and visual.concept_map:
                rendered = self.diagram_renderer.render_concept_map(slide, visual.concept_map, bounds)
            elif vtype == VisualType.LOCAL_IMAGE and visual.image:
                rendered = self.diagram_renderer.render_image(slide, visual.image, bounds)

            if rendered:
                return slide

        except Exception as exc:
            print(f"[NOTE] Visual rendering fallback for '{title}': {exc}")

        # Fallback cleanly to text card layout
        for shape in list(slide.shapes):
            sp = shape._element
            sp.getparent().remove(sp)

        self.layout.apply_canvas_background(slide)
        self.layout.add_header(slide, title=title, category="MODULE CONTENT")
        self.layout.add_footer(slide, plan["slide_num"], total_slides, data.title, data.institution)

        current_top = Inches(1.7)
        available_height = Inches(4.8)
        bullets = plan.get("bullets", [])
        paragraphs = plan.get("paragraphs", [])

        if paragraphs:
            para_text = " ".join(paragraphs)
            para_card_height = Inches(1.1)
            self.layout.create_card_container(
                slide,
                left=Inches(0.8),
                top=current_top,
                width=Inches(11.733),
                height=para_card_height,
                accent_color=self.colors.accent,
                bg_color=self.colors.card_tint,
            )
            p_box = slide.shapes.add_textbox(
                Inches(1.1), current_top + Inches(0.12), Inches(11.2), Inches(0.9)
            )
            tf_p = p_box.text_frame
            tf_p.word_wrap = True
            p_elem = tf_p.paragraphs[0]
            p_elem.text = para_text
            p_elem.font.name = self.layout.FONT_BODY
            p_elem.font.size = Pt(13.5)
            p_elem.font.color.rgb = self.colors.text_primary
            current_top += para_card_height + Inches(0.2)
            available_height -= (para_card_height + Inches(0.2))

        self.layout.create_card_container(
            slide,
            left=Inches(0.8),
            top=current_top,
            width=Inches(11.733),
            height=available_height,
            accent_color=self.colors.secondary,
        )
        self._render_bullets_in_box(
            slide,
            left=Inches(1.2),
            top=current_top + Inches(0.25),
            width=Inches(11.0),
            height=available_height - Inches(0.5),
            bullets=bullets,
        )
        return slide

    def _is_visual_valid(self, visual: SlideVisual) -> bool:
        """Validate that visual data is sufficiently populated before planning."""
        if not visual or visual.visual_type == VisualType.NONE:
            return False
        vtype = visual.visual_type
        if vtype == VisualType.PROCESS_FLOW:
            return bool(visual.process_flow and len(visual.process_flow.steps) >= 2)
        elif vtype == VisualType.TIMELINE:
            return bool(visual.timeline and len(visual.timeline.milestones) >= 2)
        elif vtype == VisualType.COMPARISON:
            return bool(
                visual.comparison and
                visual.comparison.left_column and
                visual.comparison.right_column and
                visual.comparison.left_column.points and
                visual.comparison.right_column.points
            )
        elif vtype == VisualType.CONCEPT_MAP:
            return bool(visual.concept_map and len(visual.concept_map.branches) >= 2)
        elif vtype == VisualType.LOCAL_IMAGE:
            return bool(visual.image and visual.image.file_path and Path(visual.image.file_path).exists())
        return False

    @staticmethod
    def _get_visual_category(visual_type: VisualType) -> str:
        """Return clean category header label for visual slide."""
        categories = {
            VisualType.PROCESS_FLOW: "PROCESS FLOWCHART",
            VisualType.TIMELINE: "HISTORICAL TIMELINE",
            VisualType.COMPARISON: "PARADIGM COMPARISON",
            VisualType.CONCEPT_MAP: "CONCEPT ARCHITECTURE",
            VisualType.LOCAL_IMAGE: "ILLUSTRATIVE MEDIA",
        }
        return categories.get(visual_type, "CURRICULUM VISUAL")

    def _render_bullets_in_box(
        self,
        slide: Slide,
        left: Inches,
        top: Inches,
        width: Inches,
        height: Inches,
        bullets: List[BulletPoint],
    ) -> None:
        """Render formatted bullet points with dynamic content-fitting, scaling font and spacing to fit."""
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Calculate container dimensions in points
        avail_w_pt = float(width) / 12700.0
        avail_h_pt = float(height) / 12700.0

        chosen_font_pt = 15.0
        chosen_spacing = 1.2
        chosen_space_after = 12.0

        for candidate_pt in [15.0, 14.0, 13.0, 12.0, 11.0, 10.5, 10.0]:
            cand_spacing = 1.2 if candidate_pt >= 14.0 else (1.15 if candidate_pt >= 12.0 else 1.1)
            cand_space_after = 12.0 if candidate_pt >= 14.0 else (8.0 if candidate_pt >= 12.0 else 4.0)

            # Estimate characters per line (approx 0.52 char width factor, 24pt bullet indent)
            usable_w = max(100.0, avail_w_pt - 24.0)
            chars_per_line = max(20, int(usable_w / (candidate_pt * 0.52)))

            total_lines = 0
            for b in bullets:
                raw_text = f"{b.prefix}: {b.text}" if b.prefix else b.text
                for subline in raw_text.splitlines():
                    sub_len = len(subline.strip())
                    if sub_len == 0:
                        total_lines += 1
                    else:
                        total_lines += max(1, int(math.ceil(sub_len / float(chars_per_line))))

            req_h = (total_lines * candidate_pt * cand_spacing) + (max(0, len(bullets) - 1) * cand_space_after)
            if req_h <= avail_h_pt:
                chosen_font_pt = candidate_pt
                chosen_spacing = cand_spacing
                chosen_space_after = cand_space_after
                break
        else:
            chosen_font_pt = 10.0
            chosen_spacing = 1.05
            chosen_space_after = 2.0

        # Filter only valid non-empty bullets
        valid_bullets = [b for b in bullets if (b.text and b.text.strip()) or (b.prefix and b.prefix.strip())]
        if not valid_bullets:
            return

        for i, b in enumerate(valid_bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(chosen_space_after)
            p.line_spacing = chosen_spacing

            # Decorative bullet symbol
            bullet_sym_run = p.add_run()
            bullet_sym_run.text = "▪  "
            bullet_sym_run.font.name = self.layout.FONT_HEADING
            bullet_sym_run.font.size = Pt(max(10.0, chosen_font_pt - 1.0))
            bullet_sym_run.font.bold = True
            bullet_sym_run.font.color.rgb = self.colors.secondary

            prefix_val = b.prefix.strip() if b.prefix else ""
            body_val = b.text.strip() if b.text else ""

            # Prevent prefix duplication if body_val already begins with prefix
            if prefix_val:
                p_clean = prefix_val.rstrip(":")
                if body_val.lower().startswith((p_clean + ":").lower()):
                    body_val = body_val[len(p_clean) + 1:].strip()
                elif body_val.lower().startswith(p_clean.lower()):
                    body_val = body_val[len(p_clean):].strip().lstrip(":- ")

            # Bold Prefix if available (e.g. "Term:")
            if prefix_val:
                prefix_run = p.add_run()
                prefix_run.text = f"{prefix_val.rstrip(':')}: "
                prefix_run.font.name = self.layout.FONT_HEADING
                prefix_run.font.size = Pt(chosen_font_pt)
                prefix_run.font.bold = True
                prefix_run.font.color.rgb = self.colors.primary

            # Body Text
            if body_val:
                body_run = p.add_run()
                body_run.text = body_val
                body_run.font.name = self.layout.FONT_BODY
                body_run.font.size = Pt(chosen_font_pt)
                body_run.font.bold = False
                body_run.font.color.rgb = self.colors.text_secondary

    # =========================================================================
    # Programmatic Verification Utility
    # =========================================================================

    @staticmethod
    def verify_presentation(file_path: str | Path) -> Dict[str, Any]:
        """Reopen a generated presentation to verify its structural integrity."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Presentation file does not exist: {path}")

        file_size = path.stat().st_size
        if file_size == 0:
            raise PPTGeneratorError("Generated file is 0 bytes.")

        prs = Presentation(str(path))
        slide_count = len(prs.slides)
        slide_titles: List[str] = []
        total_shapes = 0

        for slide in prs.slides:
            total_shapes += len(slide.shapes)
            # Find candidate title text
            title_text = ""
            # Prioritize shapes with larger font sizes or title boxes over badges
            candidates = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text and not text.startswith("Slide ") and text not in {
                        "ACADION", "CURRICULUM SERIES", "OVERVIEW & AGENDA",
                        "MODULE CONTENT", "COURSE CONCLUSION", "ACADEMIC LECTURE NOTES",
                        "CURRICULUM VISUAL", "HISTORICAL TIMELINE", "PROCESS FLOWCHART",
                        "PARADIGM COMPARISON", "CONCEPT ARCHITECTURE", "ILLUSTRATIVE MEDIA",
                        "VS", "KEY DISTINCTION & TAKEAWAY:", "CARDS & GRID",
                        "TWO-COLUMN COMPARISON", "KEY METRICS & HIGHLIGHTS",
                        "IMAGE & TEXT SPLIT", "CURRICULUM MODULE", "CORE ARCHITECTURE"
                    }:
                        # Check font size of first paragraph if available
                        first_para = shape.text_frame.paragraphs[0]
                        size = first_para.font.size.pt if first_para.font and first_para.font.size else 12
                        candidates.append((size, text.splitlines()[0]))

            if candidates:
                # Pick the candidate with the largest font size
                candidates.sort(key=lambda x: x[0], reverse=True)
                title_text = candidates[0][1]

            slide_titles.append(title_text or "Slide")

        return {
            "file_path": str(path.resolve()),
            "file_size_bytes": file_size,
            "slide_count": slide_count,
            "slide_titles": slide_titles,
            "total_shapes": total_shapes,
            "widescreen_16_9": (
                round(prs.slide_width / Inches(1), 2) == 13.33 and
                round(prs.slide_height / Inches(1), 2) == 7.5
            ),
        }
