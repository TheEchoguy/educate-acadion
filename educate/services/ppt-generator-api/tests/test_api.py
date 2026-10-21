"""
Automated backend API tests for Phase 5: FastAPI service endpoints.
Tests health check, text/file generation, error mapping, path traversal safety,
editing/regeneration, and presentation downloads.
"""

from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch
import pytest

from fastapi.testclient import TestClient
import pymupdf
import docx
from pptx import Presentation

from app.api import app, STORAGE_DIR
from app.models import BulletPoint, PresentationData, Section


@pytest.fixture
def client():
    """Test client for FastAPI app."""
    return TestClient(app)


# =====================================================================
# 1. Health Endpoint Tests
# =====================================================================

def test_health_check_endpoints(client: TestClient):
    """Test both /health and /api/health report healthy status and safe config."""
    for path in ("/health", "/api/health"):
        res = client.get(path)
        assert res.status_code == 200
        data = res.json()
        assert data["status"] == "healthy"
        assert data["version"] == "0.5.0"
        assert "deterministic" in data["modes"]
        assert "ai" in data["modes"]
        assert isinstance(data["ai_configured"], bool)
        assert "model" in data
        # Ensure API key is NEVER exposed in response
        assert "api_key" not in data
        assert "GEMINI_API_KEY" not in str(data)


# =====================================================================
# 2. Text-Based Generation Tests
# =====================================================================

def test_generate_from_pasted_text(client: TestClient):
    """Test deterministic presentation generation from pasted text."""
    notes = (
        "# Operating Systems Principles\n\n"
        "## Concurrency & Deadlocks\n"
        "- Mutual Exclusion: Only one process holds resource\n"
        "- Hold and Wait: Process retains allocated resources\n\n"
        "## Summary\n"
        "- Prevention: Resource ordering prevents cyclic wait\n"
    )

    res = client.post(
        "/api/generate",
        data={
            "notes_text": notes,
            "mode": "deterministic",
            "title_override": "Custom OS Lecture",
        },
    )

    assert res.status_code == 200
    data = res.json()
    assert data["title"] == "Custom OS Lecture"
    assert data["mode"] == "deterministic"
    assert data["slide_count"] >= 3
    assert len(data["slides"]) == data["slide_count"]
    assert "file_id" in data
    assert data["download_url"].startswith("/api/download/")

    # Test downloading the presentation
    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    assert "presentationml" in dl_res.headers.get("content-type", "")
    assert len(dl_res.content) > 10_000

    # Programmatically verify downloaded PPTX
    prs = Presentation(BytesIO(dl_res.content))
    assert len(prs.slides) == data["slide_count"]
    slide_texts = [
        shape.text_frame.text
        for s in prs.slides
        for shape in s.shapes
        if shape.has_text_frame
    ]
    assert any("Custom OS Lecture" in t for t in slide_texts)
    assert any("Mutual Exclusion" in t for t in slide_texts)


# =====================================================================
# 3. File Upload Generation Tests (TXT, PDF, DOCX)
# =====================================================================

def test_generate_from_uploaded_txt(client: TestClient):
    """Test generation via uploaded text file."""
    content = b"# Artificial Intelligence\n\n## Overview\n- Logic: Symbolic reasoning systems"
    files = {"file": ("lecture.txt", content, "text/plain")}

    res = client.post("/api/generate", data={"mode": "deterministic"}, files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["title"] == "Artificial Intelligence"
    assert data["source_metadata"]["file_type"] == "txt"


def test_generate_from_uploaded_pdf(client: TestClient, tmp_path: Path):
    """Test generation via uploaded PDF file."""
    pdf_path = tmp_path / "temp.pdf"
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_text(
        (50, 72),
        "# Robotics Architecture\n\n## Kinematics\n- Forward Kinematics: Computing tool pose",
    )
    doc.save(str(pdf_path))
    doc.close()

    with open(pdf_path, "rb") as f:
        files = {"file": ("robotics.pdf", f.read(), "application/pdf")}

    res = client.post("/api/generate", data={"mode": "deterministic"}, files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["title"] == "Robotics Architecture"
    assert data["source_metadata"]["file_type"] == "pdf"


def test_generate_from_uploaded_docx(client: TestClient, tmp_path: Path):
    """Test generation via uploaded Word document."""
    docx_path = tmp_path / "temp.docx"
    doc = docx.Document()
    doc.add_heading("Quantum Mechanics", level=1)
    doc.add_heading("Wavefunction", level=2)
    doc.add_paragraph("Schrodinger Equation: Dynamic wave equation", style="List Bullet")
    doc.save(str(docx_path))

    with open(docx_path, "rb") as f:
        files = {
            "file": (
                "quantum.docx",
                f.read(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }

    res = client.post("/api/generate", data={"mode": "deterministic"}, files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["title"] == "Quantum Mechanics"
    assert data["source_metadata"]["file_type"] == "docx"


# =====================================================================
# 4. Error Handling Tests
# =====================================================================

def test_generate_empty_input_returns_400(client: TestClient):
    """Test providing neither text nor file produces 400 Bad Request."""
    res = client.post("/api/generate", data={"mode": "deterministic"})
    assert res.status_code == 400
    assert "EmptyContent" in res.json().get("error", "")


def test_generate_invalid_mode_returns_400(client: TestClient):
    """Test providing an unsupported mode produces 400 Bad Request."""
    res = client.post(
        "/api/generate",
        data={"notes_text": "# Title\n- Bullet", "mode": "invalid_mode"},
    )
    assert res.status_code == 400
    assert "Invalid generation mode" in res.json().get("detail", "")


def test_generate_unsupported_file_extension(client: TestClient):
    """Test uploading an unsupported extension produces 400 Bad Request."""
    files = {"file": ("script.py", b"print('hello')", "text/plain")}
    res = client.post("/api/generate", data={"mode": "deterministic"}, files=files)
    assert res.status_code == 400
    assert "UnsupportedFileType" in res.json().get("error", "")


def test_generate_scanned_pdf_returns_422(client: TestClient, tmp_path: Path):
    """Test uploading a scanned/empty PDF produces 422 with OCR required hint."""
    pdf_path = tmp_path / "scanned.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(str(pdf_path))
    doc.close()

    with open(pdf_path, "rb") as f:
        files = {"file": ("scanned.pdf", f.read(), "application/pdf")}

    res = client.post("/api/generate", data={"mode": "deterministic"}, files=files)
    assert res.status_code == 422
    data = res.json()
    assert data.get("error") == "ScannedDocumentError"
    assert "OCR" in data.get("hint", "")


# =====================================================================
# 5. AI Mode & Key Protection Tests
# =====================================================================

def test_ai_mode_missing_api_key_returns_401(client: TestClient):
    """Test AI mode without configured API key returns 401 Unauthorized with instructional hint."""
    with patch.dict("os.environ", {}, clear=True):
        res = client.post(
            "/api/generate",
            data={
                "notes_text": "# Title\n\n## Section\n- Bullet point here",
                "mode": "ai",
            },
        )
        assert res.status_code == 401
        data = res.json()
        assert data.get("error") == "MissingAPIKey"
        assert "Deterministic Mode" in data.get("hint", "")


def test_ai_mode_mocked_success(client: TestClient):
    """Test AI mode with mocked Gemini service produces valid presentation response."""
    mock_plan_data = PresentationData(
        title="AI Planned Lecture",
        content_sections=[
            Section(
                title="AI Section",
                bullets=[BulletPoint(prefix="Concept", text="Verified educational point")],
                speaker_notes="Instructional talking point.",
            )
        ],
    )

    with patch(
        "app.gemini_service.GeminiContentPlanner.plan_from_text",
        return_value=mock_plan_data,
    ):
        res = client.post(
            "/api/generate",
            data={
                "notes_text": "# Source Notes\n- Raw text",
                "mode": "ai",
            },
        )
        assert res.status_code == 200
        data = res.json()
        assert data["title"] == "AI Planned Lecture"
        assert data["mode"] == "ai"
        # Slide 0: Title, Slide 1: Overview, Slide 2: AI Section
        assert data["slides"][2]["speaker_notes"] == "Instructional talking point."


# =====================================================================
# 6. Plan Editing and Regeneration Tests
# =====================================================================

def test_regenerate_from_edited_plan(client: TestClient):
    """Test regenerating a presentation from teacher's edited slide plan."""
    payload = {
        "title": "Teacher Edited Curriculum",
        "subtitle": "Revised Section Content",
        "slides": [
            {
                "title": "Deep Learning Fundamentals",
                "bullets": [
                    {"prefix": "Perceptron", "text": "Linear binary classification model"},
                    {"prefix": "Backprop", "text": "Reverse-mode automatic differentiation"},
                ],
                "speaker_notes": "Emphasize gradient calculation via chain rule.",
            }
        ],
    }

    res = client.post("/api/generate-from-plan", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["title"] == "Teacher Edited Curriculum"
    assert data["slide_count"] == 4  # Title + Overview + Content + Summary
    assert len(data["slides"]) == 4

    # Download regenerated file
    dl_res = client.get(data["download_url"])
    assert dl_res.status_code == 200
    prs = Presentation(BytesIO(dl_res.content))
    assert len(prs.slides) == 4


# =====================================================================
# 7. Path Traversal & Download Security Tests
# =====================================================================

def test_download_path_traversal_protection(client: TestClient):
    """Test that invalid file IDs (including path traversal attempts) are rejected with 400 or 404."""
    invalid_ids = [
        "../../etc/passwd",
        "..\\..\\windows\\win.ini",
        "invalid-id",
        "12345",
        "'; DROP TABLE slides;--",
    ]
    for bad_id in invalid_ids:
        res = client.get(f"/api/download/{bad_id}")
        assert res.status_code in (400, 404)


def test_download_nonexistent_file_returns_404(client: TestClient):
    """Test that a valid UUID pointing to a non-existent file returns 404."""
    random_uuid = "00000000-0000-0000-0000-000000000000"
    res = client.get(f"/api/download/{random_uuid}")
    assert res.status_code == 404
    assert "not found" in res.json().get("detail", "")


# =====================================================================
# 8. Slide Preview Endpoint & Rendering Tests
# =====================================================================

def test_preview_slide_endpoint_success(client: TestClient):
    """Test that generated presentations provide preview URLs and serve valid PNG images."""
    notes = (
        "# Preview Test Presentation\n\n"
        "## Preview Slide One\n"
        "- Layout: Clean academic theme\n"
        "- Quality: Crisp high-resolution preview\n"
    )
    res = client.post("/api/generate", data={"notes_text": notes, "mode": "deterministic"})
    assert res.status_code == 200
    data = res.json()
    assert "preview_urls" in data
    assert len(data["preview_urls"]) > 0

    # Request first slide preview
    preview_url = data["preview_urls"][0]
    img_res = client.get(preview_url)
    assert img_res.status_code == 200
    assert "image/png" in img_res.headers.get("content-type", "")
    # Verify PNG magic header bytes
    assert img_res.content[:8] == b"\x89PNG\r\n\x1a\n"
    assert len(img_res.content) > 1000


def test_preview_slide_path_traversal_protection(client: TestClient):
    """Test that path traversal attempts on preview endpoint are safely rejected."""
    bad_urls = [
        "/api/preview/../../etc/passwd/1",
        "/api/preview/invalid-uuid/1",
        "/api/preview/00000000-0000-0000-0000-000000000000/0",
        "/api/preview/00000000-0000-0000-0000-000000000000/9999",
    ]
    for url in bad_urls:
        res = client.get(url)
        assert res.status_code in (400, 404)


def test_preview_slide_nonexistent_returns_404(client: TestClient):
    """Test that non-existent slide preview image returns 404."""
    random_uuid = "00000000-0000-0000-0000-000000000000"
    res = client.get(f"/api/preview/{random_uuid}/99")
    assert res.status_code == 404

