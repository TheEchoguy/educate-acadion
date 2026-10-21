"""
Automated regression tests for Themes and Layouts Gallery Application & Serialization.
Addresses requirements from Section 8:
- Theme application to single slide and all slides.
- Layout transformation and geometry recalculation.
- Content, image, and speaker notes preservation.
- Obsolete native shape pruning (no duplicate/ghost shapes).
- Source PPTX immutability.
- Verification by reopening exported PPTX with python-pptx.
"""

import copy
import hashlib
import io
from pathlib import Path
import uuid
import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from app.api import STORAGE_DIR, EXPORTS_DIR
from app.layouts import THEMES, list_themes, list_layouts, get_theme
from app.models import (
    BorderDTO,
    FillDTO,
    PresentationEditorModel,
    SlideDetailDTO,
    SlideObjectDTO,
    TextContentDTO,
    TextParagraphDTO,
    TextRunDTO,
)
from app.pptx_parser import parse_presentation
from app.pptx_serializer import serialize_presentation


@pytest.fixture
def sample_multislide_deck(tmp_path: Path):
    """Create a real native PPTX presentation with 3 slides, background, cards, text, and notes."""
    fid = "test-gallery-regression-deck"
    pptx_path = STORAGE_DIR / f"{fid}.pptx"

    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 1920 virtual
    prs.slide_height = Inches(7.5)    # 1080 virtual

    # Slide 1: Cover
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.name = "Background"
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(15, 44, 89)

    title_box = slide1.shapes.add_textbox(Inches(1), Inches(2), Inches(11.333), Inches(2))
    title_box.name = "Title Box"
    p = title_box.text_frame.paragraphs[0]
    p.text = "Introduction to Computer Vision"
    p.runs[0].font.name = "Calibri"
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
    slide1.notes_slide.notes_text_frame.text = "Speaker notes for slide 1"

    # Slide 2: Content slide with card shapes and bullet points
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg2.name = "Background"
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = RGBColor(248, 250, 252)

    s2_title = slide2.shapes.add_textbox(Inches(1), Inches(0.8), Inches(11.333), Inches(1))
    s2_title.name = "Slide 2 Title"
    p2 = s2_title.text_frame.paragraphs[0]
    p2.text = "Convolutional Neural Networks"
    p2.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    card1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2.2), Inches(5), Inches(4))
    card1.name = "Card Shape 1"
    card1.fill.solid()
    card1.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card1.line.color.rgb = RGBColor(226, 232, 240)

    card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.5), Inches(2.2), Inches(5), Inches(4))
    card2.name = "Card Shape 2"
    card2.fill.solid()
    card2.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card2.line.color.rgb = RGBColor(226, 232, 240)

    t_bullet1 = slide2.shapes.add_textbox(Inches(1.2), Inches(2.4), Inches(4.6), Inches(3.5))
    t_bullet1.name = "Bullet Box 1"
    t_bullet1.text_frame.paragraphs[0].text = "Spatial hierarchy and convolution filters."

    t_bullet2 = slide2.shapes.add_textbox(Inches(6.7), Inches(2.4), Inches(4.6), Inches(3.5))
    t_bullet2.name = "Bullet Box 2"
    t_bullet2.text_frame.paragraphs[0].text = "Pooling layers and translation invariance."

    # Slide 3: Third slide
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    bg3 = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg3.name = "Background"
    bg3.fill.solid()
    bg3.fill.fore_color.rgb = RGBColor(248, 250, 252)

    s3_title = slide3.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(1))
    s3_title.name = "Slide 3 Title"
    s3_title.text_frame.paragraphs[0].text = "Conclusion and Next Steps"

    prs.save(str(pptx_path))
    model = parse_presentation(pptx_path, fid, STORAGE_DIR)

    yield fid, pptx_path, model

    # Teardown
    if pptx_path.exists():
        pptx_path.unlink()


def test_theme_apply_to_single_slide_updates_styles(sample_multislide_deck):
    """Test 1 & 2: Applying theme to single slide updates its style properties without affecting other slides."""
    fid, pptx_path, model = sample_multislide_deck

    # Save original source hash
    orig_bytes = pptx_path.read_bytes()
    orig_hash = hashlib.sha256(orig_bytes).hexdigest()

    # Modify Slide 2 to Emerald Professional (#F0FDF4 background, #064E3B primary)
    working_model = copy.deepcopy(model)
    slide2 = working_model.slides[1]
    slide2.theme = "emerald_pro"
    slide2.background_color = "#F0FDF4"
    for obj in slide2.objects:
        if obj.name == "Background" or (obj.x == 0 and obj.y == 0 and obj.width >= 1800):
            obj.fill = FillDTO(fill_type="solid", color="#F0FDF4", alpha=1.0)
        elif "Card" in obj.name or obj.object_type == "shape":
            obj.fill = FillDTO(fill_type="solid", color="#FFFFFF", alpha=1.0)
            obj.border = BorderDTO(color="#D1FAE5", width=2, style="solid")
        elif obj.text_content and "Title" in obj.name:
            for p in obj.text_content.paragraphs:
                for r in p.runs:
                    r.font_color = "#064E3B"

    # Export
    export_path, filename, warnings = serialize_presentation(
        source_pptx_path=pptx_path,
        model=working_model,
        local_assets=[],
        export_id=str(uuid.uuid4()),
        exports_dir=EXPORTS_DIR,
    )

    # Verify source PPTX remained untouched (Immutability check - Test 10)
    assert hashlib.sha256(pptx_path.read_bytes()).hexdigest() == orig_hash

    # Reopen exported PPTX
    out_prs = Presentation(str(export_path))
    assert len(out_prs.slides) == 3

    # Check Slide 2 styles
    out_s2 = out_prs.slides[1]
    s2_bg = next((s for s in out_s2.shapes if s.name == "Background"), None)
    assert s2_bg is not None
    assert s2_bg.fill.fore_color.rgb == RGBColor(0xF0, 0xFD, 0xF4)

    # Check Slide 1 remained untouched (Academic Navy #0F2C59)
    out_s1 = out_prs.slides[0]
    s1_bg = next((s for s in out_s1.shapes if s.name == "Background"), None)
    assert s1_bg is not None
    assert s1_bg.fill.fore_color.rgb == RGBColor(15, 44, 89)

    # Check Slide 3 remained untouched (#F8FAFC)
    out_s3 = out_prs.slides[2]
    s3_bg = next((s for s in out_s3.shapes if s.name == "Background"), None)
    assert s3_bg is not None
    assert s3_bg.fill.fore_color.rgb == RGBColor(0xF8, 0xFA, 0xFC)

    # Cleanup export
    if export_path.exists():
        export_path.unlink()


def test_theme_apply_to_all_slides(sample_multislide_deck):
    """Test 3: Applying theme to all slides updates every slide in the presentation."""
    fid, pptx_path, model = sample_multislide_deck
    working_model = copy.deepcopy(model)
    working_model.theme = "dark_tech"

    for slide in working_model.slides:
        slide.theme = "dark_tech"
        slide.background_color = "#0B132B"
        for obj in slide.objects:
            if obj.name == "Background" or (obj.x == 0 and obj.y == 0 and obj.width >= 1800):
                obj.fill = FillDTO(fill_type="solid", color="#0B132B", alpha=1.0)
            elif obj.object_type == "shape":
                obj.fill = FillDTO(fill_type="solid", color="#1A2639", alpha=1.0)
                obj.border = BorderDTO(color="#334155", width=1, style="solid")
            elif obj.text_content:
                for p in obj.text_content.paragraphs:
                    for r in p.runs:
                        r.font_color = "#06B6D4"

    export_path, filename, warnings = serialize_presentation(
        source_pptx_path=pptx_path,
        model=working_model,
        local_assets=[],
        export_id=str(uuid.uuid4()),
        exports_dir=EXPORTS_DIR,
    )

    out_prs = Presentation(str(export_path))
    for s_idx, slide in enumerate(out_prs.slides):
        bg = next((s for s in slide.shapes if s.name == "Background"), None)
        assert bg is not None, f"Slide {s_idx+1} missing Background shape"
        assert bg.fill.fore_color.rgb == RGBColor(0x0B, 0x13, 0x2B), f"Slide {s_idx+1} background not updated"

    if export_path.exists():
        export_path.unlink()


def test_layout_transformation_repositions_geometry_and_prunes_obsolete(sample_multislide_deck):
    """Test 11, 13, 18, 19: Applying a layout restructures coordinates, preserves content, and prunes obsolete shapes."""
    fid, pptx_path, model = sample_multislide_deck
    working_model = copy.deepcopy(model)

    slide2 = working_model.slides[1]
    # Original shape count on Slide 2 was 6 shapes (bg, title, card1, card2, bullet1, bullet2)
    assert len(slide2.objects) == 6

    # Restructure Slide 2 into a clean two-column comparison layout with 5 shapes:
    # 1. Background (reuses bg shape)
    # 2. Title (reuses title shape, repositioned)
    # 3. Column 1 Card (reuses card1 shape, repositioned to left: 80, top: 220, w: 840, h: 720)
    # 4. Column 2 Card (reuses card2 shape, repositioned to left: 1000, top: 220, w: 840, h: 720)
    # 5. Combined Text (reuses bullet1 shape, repositioned)
    # Bullet2 is omitted so it must be pruned cleanly by the serializer.
    reused_objects = [
        # BG
        slide2.objects[0],
        # Title repositioned
        SlideObjectDTO(
            object_id=slide2.objects[1].object_id,
            name="Slide Title",
            object_type="text_box",
            x=80, y=70, width=1760, height=100,
            x_emu=508000, y_emu=444500, width_emu=11176000, height_emu=635000,
            rotation=0, z_index=1,
            text_content=slide2.objects[1].text_content,
        ),
        # Column 1 Card
        SlideObjectDTO(
            object_id=slide2.objects[2].object_id,
            name="TwoCol_Card_1",
            object_type="shape",
            shape_type="ROUNDED_RECTANGLE",
            x=80, y=220, width=840, height=720,
            x_emu=508000, y_emu=1397000, width_emu=5334000, height_emu=4572000,
            rotation=0, z_index=2,
            fill=FillDTO(fill_type="solid", color="#FFFFFF", alpha=1.0),
            border=BorderDTO(color="#CBD5E1", width=1, style="solid"),
        ),
        # Column 2 Card
        SlideObjectDTO(
            object_id=slide2.objects[3].object_id,
            name="TwoCol_Card_2",
            object_type="shape",
            shape_type="ROUNDED_RECTANGLE",
            x=1000, y=220, width=840, height=720,
            x_emu=6350000, y_emu=1397000, width_emu=5334000, height_emu=4572000,
            rotation=0, z_index=3,
            fill=FillDTO(fill_type="solid", color="#FFFFFF", alpha=1.0),
            border=BorderDTO(color="#CBD5E1", width=1, style="solid"),
        ),
        # Text in Column 1
        SlideObjectDTO(
            object_id=slide2.objects[4].object_id,
            name="TwoCol_BodyText",
            object_type="text_box",
            x=110, y=250, width=780, height=660,
            x_emu=698500, y_emu=1587500, width_emu=4953000, height_emu=4191000,
            rotation=0, z_index=4,
            text_content=TextContentDTO(
                paragraphs=[
                    TextParagraphDTO(
                        paragraph_index=0, alignment="left", level=0, is_bullet=True,
                        runs=[TextRunDTO(run_index=0, text="Spatial hierarchy preserved in two-column comparison.")]
                    )
                ]
            ),
        ),
    ]
    slide2.objects = reused_objects
    slide2.layout_name = "two_column"

    export_path, filename, warnings = serialize_presentation(
        source_pptx_path=pptx_path,
        model=working_model,
        local_assets=[],
        export_id=str(uuid.uuid4()),
        exports_dir=EXPORTS_DIR,
    )

    out_prs = Presentation(str(export_path))
    out_s2 = out_prs.slides[1]

    # Verify obsolete shape was cleanly removed (count reduced from 6 to 5)
    assert len(out_s2.shapes) == 5, f"Expected 5 shapes, found {len(out_s2.shapes)}"

    # Verify Column 1 and Column 2 positions
    col1 = next((s for s in out_s2.shapes if "Card Shape 1" in s.name), None)
    col2 = next((s for s in out_s2.shapes if "Card Shape 2" in s.name), None)
    assert col1 is not None and col2 is not None
    assert col1.left < col2.left
    assert col1.width == col2.width

    # Verify preserved content
    body = next((s for s in out_s2.shapes if "Bullet Box 1" in s.name or "Spatial hierarchy" in getattr(s, "text", "")), None)
    assert body is not None
    assert "Spatial hierarchy" in body.text_frame.text

    if export_path.exists():
        export_path.unlink()


def test_blank_layout_preserves_background_and_clears_content(sample_multislide_deck):
    """Test 20: Applying blank layout retains slide background and speaker notes while clearing content."""
    fid, pptx_path, model = sample_multislide_deck
    working_model = copy.deepcopy(model)

    slide1 = working_model.slides[0]
    # Retain only the background shape
    bg_shape = slide1.objects[0]
    slide1.objects = [bg_shape]
    slide1.layout_name = "blank"

    export_path, filename, warnings = serialize_presentation(
        source_pptx_path=pptx_path,
        model=working_model,
        local_assets=[],
        export_id=str(uuid.uuid4()),
        exports_dir=EXPORTS_DIR,
    )

    out_prs = Presentation(str(export_path))
    out_s1 = out_prs.slides[0]

    # Only 1 shape should remain on the slide
    assert len(out_s1.shapes) == 1
    assert out_s1.shapes[0].name == "Background"

    # Speaker notes preserved
    assert "Speaker notes for slide 1" in out_s1.notes_slide.notes_text_frame.text

    if export_path.exists():
        export_path.unlink()


def test_all_14_themes_and_layouts_catalog_available():
    """Test 23: Verify all 14 registered themes and 6 layout families are available."""
    themes = list_themes()
    assert len(themes) >= 14
    theme_ids = {t["id"] for t in themes}
    for expected in [
        "academic_navy", "modern_blue", "emerald_pro", "purple_gradient",
        "corporate_minimal", "dark_tech", "warm_creative", "clean_light",
        "midnight_blue", "cyber_neon", "obsidian_gold", "dark_purple",
        "graphite_pro", "deep_ocean",
    ]:
        assert expected in theme_ids

    layouts = list_layouts()
    assert len(layouts) >= 6
    layout_ids = {l["id"] for l in layouts}
    for expected_layout in [
        "smart_mixed", "cards_grid", "two_column", "visual_flow",
        "stats_metrics", "classic_bullets",
    ]:
        assert expected_layout in layout_ids
