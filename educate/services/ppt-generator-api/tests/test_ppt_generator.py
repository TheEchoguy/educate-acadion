"""
Unit and integration tests for PPTGenerator.
"""

from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.content_parser import ContentParser
from app.models import BulletPoint, PresentationData, Section, SectionType
from app.ppt_generator import PPTGenerator


def test_generate_from_sample_notes(tmp_path: Path):
    """Test generating a presentation from realistic sample notes and verifying structure."""
    sample_file = Path("sample_data/sample_notes.txt")
    assert sample_file.exists(), "Sample notes file must exist in sample_data/"

    parser = ContentParser()
    data = parser.parse_file(sample_file)

    generator = PPTGenerator()
    out_file = tmp_path / "output_presentation.pptx"
    result_path = generator.generate(data, out_file)

    assert result_path.exists()
    assert result_path.stat().st_size > 10_000

    # Reopen programmatically
    prs = Presentation(str(result_path))
    assert len(prs.slides) >= 10, "Expected at least 10 slides due to pagination"

    # Verify widescreen 16:9 dimensions
    assert round(prs.slide_width / Inches(1), 2) == 13.33
    assert round(prs.slide_height / Inches(1), 2) == 7.5

    # Run verification utility
    ver = PPTGenerator.verify_presentation(result_path)
    assert ver["slide_count"] == len(prs.slides)
    assert ver["widescreen_16_9"] is True
    assert ver["total_shapes"] > 50


def test_content_pagination_avoids_overflow(tmp_path: Path):
    """Verify that a section with > 4 bullets is split across multiple slides."""
    data = PresentationData(
        title="Pagination Testing",
        content_sections=[
            Section(
                title="Extensive List of Topics",
                bullets=[
                    BulletPoint(text=f"Detailed bullet item number {i}")
                    for i in range(1, 9)
                ],
            )
        ]
    )

    generator = PPTGenerator()
    out_file = tmp_path / "pagination_test.pptx"
    generator.generate(data, out_file)

    prs = Presentation(str(out_file))
    # Expected: Title slide (1) + Overview (1) + Extensive List (2 parts) + Summary (1) = 5 slides
    assert len(prs.slides) == 5

    # Check slide titles contain Part 1/2 and Part 2/2
    ver = PPTGenerator.verify_presentation(out_file)
    titles_str = " ".join(ver["slide_titles"])
    assert "Part 1/2" in titles_str
    assert "Part 2/2" in titles_str


def test_essential_text_and_slide_numbers_in_presentation(tmp_path: Path):
    """Verify that essential text and slide numbers are present in slides."""
    data = PresentationData(
        title="Robotics & Autonomous Systems",
        content_sections=[
            Section(
                title="Sensors and Actuators",
                bullets=[
                    BulletPoint(text="LiDAR sensors provide dense 3D spatial point clouds.", prefix="LiDAR"),
                ],
            )
        ]
    )

    generator = PPTGenerator()
    out_file = tmp_path / "text_verification.pptx"
    generator.generate(data, out_file)

    prs = Presentation(str(out_file))
    all_slide_texts = []
    for slide in prs.slides:
        slide_text = " ".join(
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame
        )
        all_slide_texts.append(slide_text)

    combined_text = " ".join(all_slide_texts)
    assert "Robotics & Autonomous Systems" in combined_text
    assert "Sensors and Actuators" in combined_text
    assert "LiDAR" in combined_text
    assert "Slide 3 of 4" in combined_text  # Footer verification


@pytest.mark.parametrize("title_comp", ["editorial", "cinematic", "split"])
@pytest.mark.parametrize("theme_id,expected_is_dark", [
    ("clean_light", False),
    ("academic_navy", False),
    ("dark_tech", True),
    ("cyber_neon", True),
])
def test_title_slide_theme_consistency(tmp_path: Path, title_comp: str, theme_id: str, expected_is_dark: bool):
    """Verify that title slide canvas and containers strictly reflect the active theme."""
    from app.layouts import get_theme

    theme = get_theme(theme_id)
    assert theme.is_dark == expected_is_dark

    data = PresentationData(
        title="Title Slide Theme Verification",
        title_composition=title_comp,
        content_sections=[
            Section(
                title="Foundational Section",
                bullets=[BulletPoint(text="Testing theme consistency across title compositions.")],
            )
        ]
    )

    generator = PPTGenerator(theme=theme_id)
    out_file = tmp_path / f"title_test_{theme_id}_{title_comp}.pptx"
    generator.generate(data, out_file)

    prs = Presentation(str(out_file))
    cover_slide = prs.slides[0]

    # Find the Background shape
    bg_shape = None
    for shape in cover_slide.shapes:
        if shape.name == "Background":
            bg_shape = shape
            break

    assert bg_shape is not None, f"Cover slide missing 'Background' shape for {theme_id} ({title_comp})"
    assert bg_shape.fill.type is not None

    bg_rgb = (bg_shape.fill.fore_color.rgb[0], bg_shape.fill.fore_color.rgb[1], bg_shape.fill.fore_color.rgb[2])
    expected_bg = (theme.bg_light[0], theme.bg_light[1], theme.bg_light[2])
    assert bg_rgb == expected_bg, (
        f"Cover slide background mismatch for theme {theme_id} ({title_comp}): "
        f"got {bg_rgb}, expected {expected_bg}"
    )

    # In cinematic mode, verify the container frame uses the theme's card_bg
    if title_comp == "cinematic":
        frame_shape = next((s for s in cover_slide.shapes if s.name == "HeroFrame"), None)
        assert frame_shape is not None, f"Cover slide missing 'HeroFrame' for {theme_id}"
        frame_rgb = (frame_shape.fill.fore_color.rgb[0], frame_shape.fill.fore_color.rgb[1], frame_shape.fill.fore_color.rgb[2])
        expected_frame = (theme.card_bg[0], theme.card_bg[1], theme.card_bg[2])
        assert frame_rgb == expected_frame, (
            f"Cinematic frame fill mismatch for theme {theme_id}: got {frame_rgb}, expected {expected_frame}"
        )

    # In split mode, verify the hero column uses card_tint (light) or card_bg (dark)
    if title_comp == "split":
        col_shape = next((s for s in cover_slide.shapes if s.name == "HeroColumn"), None)
        assert col_shape is not None, f"Cover slide missing 'HeroColumn' for {theme_id}"
        col_rgb = (col_shape.fill.fore_color.rgb[0], col_shape.fill.fore_color.rgb[1], col_shape.fill.fore_color.rgb[2])
        expected_col = (theme.card_bg[0], theme.card_bg[1], theme.card_bg[2]) if theme.is_dark else (theme.card_tint[0], theme.card_tint[1], theme.card_tint[2])
        assert col_rgb == expected_col, (
            f"Split column fill mismatch for theme {theme_id}: got {col_rgb}, expected {expected_col}"
        )


@pytest.mark.parametrize("theme_id", ["clean_light", "academic_navy", "dark_tech", "cyber_neon"])
def test_section_divider_theme_propagation(theme_id, tmp_path):
    """Verify that section divider slides consistently apply the active theme canvas and card fill."""
    data = PresentationData(
        title="Divider Slide Theme Verification",
        content_sections=[
            Section(
                title="Foundational Section",
                bullets=[BulletPoint(text="Testing theme consistency across section divider.")],
            )
        ]
    )
    generator = PPTGenerator(theme=theme_id)
    theme = generator.colors

    prs = Presentation()
    prs.slide_width = generator.layout.SLIDE_WIDTH
    prs.slide_height = generator.layout.SLIDE_HEIGHT

    divider_plan = {
        "type": "section_divider",
        "slide_num": 2,
        "title": "Module Deep Dive",
        "category": "ADVANCED TOPICS",
        "subtitle": "Theoretical Foundations & Implementations",
    }
    slide = generator._render_section_divider_slide(prs, data, divider_plan, 5)

    # 1. Background must match theme bg_light
    canvas_shape = next((s for s in slide.shapes if s.name == "Background"), None)
    assert canvas_shape is not None, f"Divider slide missing 'Background' for {theme_id}"
    canvas_rgb = (canvas_shape.fill.fore_color.rgb[0], canvas_shape.fill.fore_color.rgb[1], canvas_shape.fill.fore_color.rgb[2])
    expected_bg = (theme.bg_light[0], theme.bg_light[1], theme.bg_light[2])
    assert canvas_rgb == expected_bg

    # 2. DividerCard must match theme card_bg
    card_shape = next((s for s in slide.shapes if s.name == "DividerCard"), None)
    assert card_shape is not None, f"Divider slide missing 'DividerCard' for {theme_id}"
    card_rgb = (card_shape.fill.fore_color.rgb[0], card_shape.fill.fore_color.rgb[1], card_shape.fill.fore_color.rgb[2])
    expected_card = (theme.card_bg[0], theme.card_bg[1], theme.card_bg[2])
    assert card_rgb == expected_card


def test_pillow_fallback_preview_consistency(tmp_path):
    """Verify that Pillow fallback preview renderer operates consistently without crashing."""
    from app.preview_renderer import render_pptx_slides_pillow_fallback

    data = PresentationData(
        title="Preview Test Presentation",
        content_sections=[
            Section(
                title="Foundational Section",
                bullets=[BulletPoint(text="Testing Pillow fallback preview consistency.")],
            )
        ]
    )

    out_file = tmp_path / "test_preview_gen.pptx"
    generator = PPTGenerator(theme="clean_light")
    generator.generate(data, output_path=out_file)

    preview_dir = tmp_path / "previews"
    images = render_pptx_slides_pillow_fallback(out_file, preview_dir)
    assert len(images) > 0, "Pillow fallback should render at least 1 preview image"
    for img_path in images:
        assert img_path.exists() and img_path.stat().st_size > 0


