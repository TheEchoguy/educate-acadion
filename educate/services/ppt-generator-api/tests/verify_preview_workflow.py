"""
End-to-end verification script for Acadion AI PPT Generator PowerPoint Preview.
Tests the complete workflow:
  1. Generate presentation (deterministic mode with notes)
  2. Verify preview URLs and fetch slide PNG images
  3. Validate downloaded original PPTX
  4. Edit slide plan and regenerate presentation
  5. Verify updated slide preview PNG images
  6. Validate downloaded regenerated PPTX with updated slide content
"""

from io import BytesIO
import json
import sys
import urllib.parse
import urllib.request
from pptx import Presentation

BASE_URL = "http://127.0.0.1:8000"


def main():
    print("=" * 65)
    print("STARTING E2E PREVIEW WORKFLOW VERIFICATION")
    print("=" * 65)

    # -------------------------------------------------------------
    # STEP 1: Generate presentation from curriculum notes
    # -------------------------------------------------------------
    print("\n[STEP 1] Generating presentation via POST /api/generate...")
    notes = (
        "# Distributed Systems Principles\n"
        "Curriculum Notes — Scalability & Consensus\n\n"
        "## Consensus Algorithms\n"
        "- Paxos: Asynchronous consensus protocol with proven safety guarantees\n"
        "- Raft: Understandable leader-based consensus algorithm designed for ease of implementation\n"
        "- Byzantine Fault Tolerance: Tolerates arbitrary node failures and malicious actors\n\n"
        "## Scalability Patterns\n"
        "- Horizontal Partitioning: Sharding datasets across autonomous database nodes\n"
        "- Eventual Consistency: Sacrifices immediate consistency for high availability in CAP theorem\n"
    )

    data = urllib.parse.urlencode({
        "notes_text": notes,
        "mode": "deterministic",
        "title_override": "Distributed Systems & Scalability",
    }).encode("utf-8")

    req = urllib.request.Request(f"{BASE_URL}/api/generate", data=data, method="POST")
    with urllib.request.urlopen(req) as resp:
        res = json.loads(resp.read().decode("utf-8"))

    file_id = res["file_id"]
    slide_count = res["slide_count"]
    preview_urls = res.get("preview_urls", [])

    print(f"  -> Generated: '{res['title']}'")
    print(f"  -> File ID: {file_id}")
    print(f"  -> Slide count: {slide_count}")
    print(f"  -> Preview URLs generated: {len(preview_urls)}")

    assert slide_count >= 4, f"Expected at least 4 slides, got {slide_count}"
    assert len(preview_urls) == slide_count, f"Mismatch in preview URLs: {len(preview_urls)} vs {slide_count}"

    # -------------------------------------------------------------
    # STEP 2: Verify high-fidelity slide preview images
    # -------------------------------------------------------------
    print("\n[STEP 2] Verifying slide preview images via GET /api/preview/{file_id}/{slide_number}...")
    for idx, url in enumerate(preview_urls, start=1):
        img_req = urllib.request.Request(f"{BASE_URL}{url}")
        with urllib.request.urlopen(img_req) as img_resp:
            assert img_resp.status == 200, f"Failed to fetch preview {url}"
            content_type = img_resp.headers.get("content-type", "")
            assert "image/png" in content_type, f"Expected image/png, got {content_type}"
            img_bytes = img_resp.read()
            # Verify PNG signature (\x89PNG\r\n\x1a\n)
            assert img_bytes[:8] == b"\x89PNG\r\n\x1a\n", f"Invalid PNG magic bytes for slide {idx}"
            assert len(img_bytes) > 2000, f"Slide {idx} image too small ({len(img_bytes)} bytes)"
            print(f"  -> Slide {idx} preview verified: {len(img_bytes):,} bytes, Content-Type: {content_type}")

    # -------------------------------------------------------------
    # STEP 3: Validate downloaded original PowerPoint (.pptx)
    # -------------------------------------------------------------
    print("\n[STEP 3] Downloading and verifying original PPTX via GET /api/download/{file_id}...")
    dl_url = res["download_url"]
    dl_req = urllib.request.Request(f"{BASE_URL}{dl_url}")
    with urllib.request.urlopen(dl_req) as dl_resp:
        assert dl_resp.status == 200
        pptx_bytes = dl_resp.read()
        assert len(pptx_bytes) > 10_000, "PPTX file size too small"
        prs = Presentation(BytesIO(pptx_bytes))
        print(f"  -> Downloaded original presentation: {len(pptx_bytes):,} bytes")
        print(f"  -> Presentation reopened successfully! Slide count: {len(prs.slides)}")
        assert len(prs.slides) == slide_count

    # -------------------------------------------------------------
    # STEP 4: Edit slide plan and regenerate presentation
    # -------------------------------------------------------------
    print("\n[STEP 4] Editing slide plan and regenerating via POST /api/generate-from-plan...")
    edited_payload = {
        "title": "Distributed Systems & Scalability (Updated Edition)",
        "subtitle": "Educational Curriculum Courseware",
        "slides": [
            {
                "title": "Advanced Consensus Algorithms & Raft",
                "bullets": [
                    {"prefix": "Leader Election", "text": "Heartbeat-driven randomized timeout election protocol"},
                    {"prefix": "Log Replication", "text": "AppendEntries RPC commits monotonic entries to majority of cluster quorum"},
                ],
                "speaker_notes": "Emphasize that Raft decomposes consensus into leader election, log replication, and safety.",
            },
            {
                "title": "High-Throughput Partitioning",
                "bullets": [
                    {"prefix": "Consistent Hashing", "text": "Minimizes key remapping during cluster node resizing"},
                    {"prefix": "Vector Clocks", "text": "Captures causal relationships between concurrent distributed writes"},
                ],
                "speaker_notes": "Contrast Dynamo-style vector clocks with Google Spanner TrueTime API.",
            },
        ],
    }

    regen_req = urllib.request.Request(
        f"{BASE_URL}/api/generate-from-plan",
        data=json.dumps(edited_payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )

    with urllib.request.urlopen(regen_req) as regen_resp:
        regen_res = json.loads(regen_resp.read().decode("utf-8"))

    new_file_id = regen_res["file_id"]
    new_slide_count = regen_res["slide_count"]
    new_preview_urls = regen_res.get("preview_urls", [])

    print(f"  -> Regenerated: '{regen_res['title']}'")
    print(f"  -> New File ID: {new_file_id}")
    print(f"  -> Regenerated slide count: {new_slide_count}")
    print(f"  -> New Preview URLs: {len(new_preview_urls)}")

    assert new_file_id != file_id, "Regenerated file_id should be distinct from original"
    assert len(new_preview_urls) == new_slide_count

    # -------------------------------------------------------------
    # STEP 5: Verify updated slide preview images
    # -------------------------------------------------------------
    print("\n[STEP 5] Verifying regenerated slide preview images...")
    for idx, url in enumerate(new_preview_urls, start=1):
        img_req = urllib.request.Request(f"{BASE_URL}{url}")
        with urllib.request.urlopen(img_req) as img_resp:
            assert img_resp.status == 200
            img_bytes = img_resp.read()
            assert img_bytes[:8] == b"\x89PNG\r\n\x1a\n"
            print(f"  -> Regenerated Slide {idx} preview verified: {len(img_bytes):,} bytes")

    # -------------------------------------------------------------
    # STEP 6: Validate downloaded regenerated PowerPoint (.pptx)
    # -------------------------------------------------------------
    print("\n[STEP 6] Downloading and inspecting regenerated PPTX...")
    new_dl_req = urllib.request.Request(f"{BASE_URL}{regen_res['download_url']}")
    with urllib.request.urlopen(new_dl_req) as new_dl_resp:
        assert new_dl_resp.status == 200
        new_pptx_bytes = new_dl_resp.read()
        new_prs = Presentation(BytesIO(new_pptx_bytes))
        print(f"  -> Downloaded regenerated PPTX: {len(new_pptx_bytes):,} bytes")
        print(f"  -> Slides verified: {len(new_prs.slides)}")

        # Verify edited content appears in the presentation shapes
        all_texts = [
            shape.text_frame.text
            for s in new_prs.slides
            for shape in s.shapes
            if shape.has_text_frame
        ]

        assert any("Advanced Consensus Algorithms & Raft" in t for t in all_texts), (
            "Edited slide title 'Advanced Consensus Algorithms & Raft' not found in PPTX"
        )
        assert any("Leader Election" in t for t in all_texts), (
            "Edited bullet prefix 'Leader Election' not found in PPTX"
        )
        print("  -> Edited slide titles and bullets verified in PowerPoint presentation!")

    print("\n" + "=" * 65)
    print("ALL 6 STEPS COMPLETED & VERIFIED SUCCESSFULLY!")
    print("=" * 65)


if __name__ == "__main__":
    main()
