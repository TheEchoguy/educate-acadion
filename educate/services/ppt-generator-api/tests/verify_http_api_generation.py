"""
Programmatic verification script that generates a presentation through the FastAPI HTTP API,
saves the downloaded .pptx file, reopens it with python-pptx, and asserts structure & content.
"""

from pathlib import Path
import sys

# Ensure workspace root is in sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient
from pptx import Presentation

from app.api import app


def test_and_verify_live_http_generation():
    client = TestClient(app)

    # 1. Test Health
    health_res = client.get("/health")
    assert health_res.status_code == 200, f"Health check failed: {health_res.text}"
    health_data = health_res.json()
    print(f"[+] HTTP API Health Check Passed: {health_data}")

    # 2. Generate Presentation via HTTP POST /api/generate
    curriculum_notes = (
        "# Distributed Systems & Microservices\n"
        "Advanced Engineering Masterclass — Acadion Learning\n\n"
        "## Core Principles\n"
        "- Decoupled Services: Independent lifecycle and deployment boundaries\n"
        "- Resilience: Circuit breaker patterns and bounded retry policies\n"
        "- Observability: Distributed tracing across service mesh boundaries\n\n"
        "## Monolith vs Microservices\n"
        "- Monolith: Unified codebase with centralized database\n"
        "- Monolith: Simpler initial scaffolding but harder long-term scaling\n"
        "- Microservices: Autonomous domain boundaries with polyglot persistence\n"
        "- Microservices: High availability through fault-isolated micro-components\n\n"
        "## Summary & Architecture Review\n"
        "- Event-Driven Architecture: Asynchronous decoupling via message brokers\n"
        "- Zero Trust Security: Continuous authentication and mutual TLS\n"
    )

    print("[*] Sending POST /api/generate request...")
    res = client.post(
        "/api/generate",
        data={
            "notes_text": curriculum_notes,
            "mode": "deterministic",
            "title_override": "Distributed Systems & Microservices",
        },
    )

    assert res.status_code == 200, f"Generate request failed: {res.text}"
    result = res.json()
    print(f"[+] HTTP Generate Successful:")
    print(f"    - File ID    : {result['file_id']}")
    print(f"    - Title      : {result['title']}")
    print(f"    - Slide count: {result['slide_count']}")
    print(f"    - Download   : {result['download_url']}")

    # 3. Download PPTX via GET /api/download/{file_id}
    print(f"[*] Downloading generated PPTX via {result['download_url']}...")
    dl_res = client.get(result["download_url"])
    assert dl_res.status_code == 200, f"Download failed: {dl_res.status_code}"

    out_path = Path("output/http_api_verified.pptx")
    out_path.write_bytes(dl_res.content)
    print(f"[+] Saved downloaded file to: {out_path.resolve()} ({len(dl_res.content):,} bytes)")

    # 4. Reopen programmatically with python-pptx and verify
    prs = Presentation(str(out_path))
    slide_count = len(prs.slides)
    assert slide_count == result["slide_count"], f"Mismatch in slide count: {slide_count} vs {result['slide_count']}"

    # Check 16:9 widescreen dimensions
    width_in = prs.slide_width / 914400
    height_in = prs.slide_height / 914400
    assert abs(width_in - 13.333) < 0.1, f"Expected 13.333\" width, got {width_in}\""
    assert abs(height_in - 7.5) < 0.1, f"Expected 7.5\" height, got {height_in}\""

    # Check all slide text
    all_texts = []
    for s in prs.slides:
        for shape in s.shapes:
            if shape.has_text_frame:
                all_texts.append(shape.text_frame.text)
    full_text = " ".join(all_texts)

    expected_keywords = [
        "Distributed Systems & Microservices",
        "Decoupled Services",
        "Monolith vs Microservices",
        "Circuit breaker",
        "Event-Driven Architecture",
    ]

    print("[*] Verifying slide text contents...")
    for kw in expected_keywords:
        assert kw.lower() in full_text.lower(), f"Missing expected keyword '{kw}' in generated presentation!"
        print(f"    [OK] Verified keyword: '{kw}'")

    print(f"\n[SUCCESS] Successfully verified presentation generated via HTTP API ({slide_count} slides, 16:9 widescreen)!\n")


if __name__ == "__main__":
    test_and_verify_live_http_generation()
