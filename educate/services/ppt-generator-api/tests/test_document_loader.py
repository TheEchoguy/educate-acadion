"""
Automated tests for Phase 4: Teacher Input and Document Ingestion.
Tests TXT, PDF, DOCX, and pasted text loading, resource limits,
scanned document detection, error handling, and pipeline integration.
"""

from pathlib import Path
from unittest.mock import MagicMock, patch
import pytest

import docx
import pymupdf

from app.content_parser import ContentParser
from app.document_loader import DocumentLoader
from app.gemini_service import GeminiContentPlanner
from app.main import run
from app.models import (
    BulletPoint,
    CorruptedDocumentError,
    DocumentIngestionError,
    DocumentMetadata,
    EmptyContentError,
    FileSizeLimitExceededError,
    InputFileNotFoundError,
    PresentationData,
    ScannedDocumentError,
    Section,
    TextLengthLimitExceededError,
    UnsupportedFileTypeError,
)
from app.ppt_generator import PPTGenerator


# =====================================================================
# Fixture Helpers
# =====================================================================

def make_test_pdf(path: Path, content_pages: list[str]) -> Path:
    """Create a multi-page PDF fixture containing educational text."""
    doc = pymupdf.open()
    for text in content_pages:
        page = doc.new_page()
        page.insert_text((50, 72), text)
    doc.save(str(path))
    doc.close()
    return path


def make_test_docx(
    path: Path,
    title: str,
    sections: list[tuple[str, list[str]]],
    include_table: bool = False
) -> Path:
    """Create a structured Word document fixture with headings and bullet points."""
    doc = docx.Document()
    doc.add_heading(title, level=1)
    for sec_title, bullets in sections:
        doc.add_heading(sec_title, level=2)
        for bullet in bullets:
            doc.add_paragraph(bullet, style="List Bullet")

    if include_table:
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Term"
        table.rows[0].cells[1].text = "Definition"
        table.rows[1].cells[0].text = "Supervised Learning"
        table.rows[1].cells[1].text = "Learning from labeled training data"

    doc.save(str(path))
    return path


# =====================================================================
# 1. TXT Ingestion Tests
# =====================================================================

def test_load_txt_utf8(tmp_path: Path):
    """Test standard UTF-8 text file extraction."""
    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("# Deep Learning Lecture\n\n## Overview\n- Neural Networks: Foundation of DL", encoding="utf-8")

    loader = DocumentLoader()
    text, meta = loader.load(txt_file)

    assert "Deep Learning Lecture" in text
    assert "Neural Networks: Foundation of DL" in text
    assert meta.file_type == "txt"
    assert meta.page_count == 1
    assert meta.word_count > 0
    assert meta.character_count == len(text)


def test_load_txt_utf8_bom(tmp_path: Path):
    """Test UTF-8 text with BOM marker is cleanly stripped and decoded."""
    txt_file = tmp_path / "bom_notes.txt"
    bom_content = b"\xef\xbb\xbf# Quantum Computing\n\n## Principles\n- Superposition: Qubits in linear states"
    txt_file.write_bytes(bom_content)

    loader = DocumentLoader()
    text, meta = loader.load(txt_file)

    assert text.startswith("# Quantum Computing")
    assert "\ufeff" not in text
    assert meta.file_type == "txt"


def test_load_txt_latin1_fallback(tmp_path: Path):
    """Test fallback decoding for non-UTF8 legacy encoded files."""
    txt_file = tmp_path / "latin1_notes.txt"
    # Contains accented characters encoded in Latin-1
    content_latin1 = "# Modèle d'Apprentissage\n\n## Introduction\n- Théorie: Analyse des données".encode("latin-1")
    txt_file.write_bytes(content_latin1)

    loader = DocumentLoader()
    text, meta = loader.load(txt_file)

    assert "Modèle d'Apprentissage" in text
    assert "Théorie" in text


# =====================================================================
# 2. PDF Ingestion Tests
# =====================================================================

def test_load_pdf_multipage(tmp_path: Path):
    """Test page-by-page PDF extraction with metadata."""
    pdf_path = tmp_path / "lecture.pdf"
    pages = [
        "# Introduction to Robotics\n\n## Kinematics\n- Forward Kinematics: Computing end-effector pose",
        "## Dynamics and Control\n- Inverse Kinematics: Determining joint configurations\n- PID Control: Closed-loop regulation"
    ]
    make_test_pdf(pdf_path, pages)

    loader = DocumentLoader()
    text, meta = loader.load(pdf_path)

    assert "Introduction to Robotics" in text
    assert "Kinematics" in text
    assert "Dynamics and Control" in text
    assert meta.file_type == "pdf"
    assert meta.page_count == 2
    assert meta.word_count > 10


def test_load_pdf_scanned_detection(tmp_path: Path):
    """Test that a PDF with 0 extractable characters raises ScannedDocumentError requesting OCR."""
    pdf_path = tmp_path / "scanned_document.pdf"
    # Create empty page (simulating scanned image PDF without OCR text layer)
    doc = pymupdf.open()
    doc.new_page()
    doc.save(str(pdf_path))
    doc.close()

    loader = DocumentLoader()
    with pytest.raises(ScannedDocumentError) as exc_info:
        loader.load(pdf_path)

    err_msg = str(exc_info.value)
    assert "scanned document" in err_msg.lower()
    assert "ocr" in err_msg.lower()


# =====================================================================
# 3. DOCX Ingestion Tests
# =====================================================================

def test_load_docx_structure(tmp_path: Path):
    """Test Word document ingestion preserving headings, bullets, and tables."""
    docx_path = tmp_path / "curriculum.docx"
    sections = [
        ("Architecture Overview", [
            "Transformer Architecture: Self-attention mechanisms",
            "Multi-Head Attention: Parallel projection representations"
        ]),
        ("Training Dynamics", [
            "AdamW Optimizer: Weight decay decoupling",
            "Learning Rate Warmup: Initial stability phase"
        ])
    ]
    make_test_docx(docx_path, "Natural Language Processing", sections, include_table=True)

    loader = DocumentLoader()
    text, meta = loader.load(docx_path)

    assert "# Natural Language Processing" in text
    assert "## Architecture Overview" in text
    assert "- Transformer Architecture: Self-attention mechanisms" in text
    assert "## Training Dynamics" in text
    assert "Supervised Learning | Learning from labeled training data" in text
    assert meta.file_type == "docx"


def test_load_docx_empty(tmp_path: Path):
    """Test empty Word document raises EmptyContentError."""
    docx_path = tmp_path / "empty.docx"
    doc = docx.Document()
    doc.save(str(docx_path))

    loader = DocumentLoader()
    with pytest.raises(EmptyContentError):
        loader.load(docx_path)


# =====================================================================
# 4. Pasted Text Ingestion Tests
# =====================================================================

def test_load_pasted_text_success():
    """Test ingesting direct pasted lecture notes."""
    raw = "   # Neural Networks 101\n\n- Perceptron: Binary classifier\n- MLP: Multilayer Perceptron   \n\n"
    loader = DocumentLoader()
    text, meta = loader.load_text(raw)

    assert text.startswith("# Neural Networks 101")
    assert meta.file_type == "raw_text"
    assert meta.page_count == 1
    assert meta.word_count == 12


def test_load_pasted_text_empty():
    """Test empty or whitespace pasted text raises EmptyContentError."""
    loader = DocumentLoader()
    with pytest.raises(EmptyContentError):
        loader.load_text("   \n\t  ")


# =====================================================================
# 5. Resource Limits and Security Validation
# =====================================================================

def test_missing_file_raises_error(tmp_path: Path):
    """Test missing file raises InputFileNotFoundError."""
    loader = DocumentLoader()
    with pytest.raises(InputFileNotFoundError):
        loader.load(tmp_path / "does_not_exist.pdf")


def test_unsupported_file_extension(tmp_path: Path):
    """Test unsupported file extensions raise UnsupportedFileTypeError."""
    bad_file = tmp_path / "payload.exe"
    bad_file.write_text("dummy", encoding="utf-8")

    loader = DocumentLoader()
    with pytest.raises(UnsupportedFileTypeError) as exc_info:
        loader.load(bad_file)
    assert ".exe" in str(exc_info.value)


def test_file_size_limit_exceeded(tmp_path: Path):
    """Test file size limit check rejects files exceeding max_file_size_bytes."""
    large_file = tmp_path / "large.txt"
    # Configure tiny 100-byte limit
    loader = DocumentLoader(max_file_size_bytes=100)
    large_file.write_text("A" * 200, encoding="utf-8")

    with pytest.raises(FileSizeLimitExceededError) as exc_info:
        loader.load(large_file)
    assert "exceeds maximum allowed limit" in str(exc_info.value).lower()


def test_extracted_text_length_limit_exceeded(tmp_path: Path):
    """Test extracted text limit check rejects oversized content."""
    txt_file = tmp_path / "oversized.txt"
    # Max text limit of 50 chars
    loader = DocumentLoader(max_file_size_bytes=10_000, max_text_length=50)
    txt_file.write_text("# Educational Lecture\n" + "Word " * 20, encoding="utf-8")

    with pytest.raises(TextLengthLimitExceededError) as exc_info:
        loader.load(txt_file)
    assert "maximum allowed limit" in str(exc_info.value).lower()


def test_corrupted_pdf_magic_bytes(tmp_path: Path):
    """Test file with .pdf extension but invalid header bytes raises CorruptedDocumentError."""
    corrupted_pdf = tmp_path / "fake.pdf"
    corrupted_pdf.write_bytes(b"NOT A REAL PDF FILE HEADER")

    loader = DocumentLoader()
    with pytest.raises(CorruptedDocumentError) as exc_info:
        loader.load(corrupted_pdf)
    assert "signature" in str(exc_info.value).lower() or "header" in str(exc_info.value).lower()


def test_corrupted_docx_magic_bytes(tmp_path: Path):
    """Test file with .docx extension but invalid zip header raises CorruptedDocumentError."""
    corrupted_docx = tmp_path / "fake.docx"
    corrupted_docx.write_bytes(b"THIS IS NOT A ZIP ARCHIVE")

    loader = DocumentLoader()
    with pytest.raises(CorruptedDocumentError) as exc_info:
        loader.load(corrupted_docx)
    assert "signature" in str(exc_info.value).lower() or "archive" in str(exc_info.value).lower()


# =====================================================================
# 6. Pipeline Integration Tests (Deterministic & PPT Generation)
# =====================================================================

def test_content_parser_integration_with_pdf(tmp_path: Path):
    """Test ContentParser.parse_file directly parses PDF into PresentationData."""
    pdf_path = tmp_path / "lecture.pdf"
    content = (
        "# Artificial Intelligence Fundamentals\n\n"
        "## Overview & History\n"
        "- 1950: Alan Turing proposes the Turing Test\n"
        "- 2012: AlexNet demonstrates deep learning breakthrough\n\n"
        "## Machine Learning Process\n"
        "- Data Collection: Gathering domain dataset\n"
        "- Feature Engineering: Extracting informative signals\n"
        "- Model Training: Optimizing parameter weights\n"
    )
    make_test_pdf(pdf_path, [content])

    parser = ContentParser()
    data = parser.parse_file(pdf_path)

    assert data.title == "Artificial Intelligence Fundamentals"
    assert len(data.all_sections) >= 2

    # Verify presentation generation from parsed PDF
    output_pptx = tmp_path / "pdf_presentation.pptx"
    gen = PPTGenerator()
    resolved_path = gen.generate(data, str(output_pptx))
    assert Path(resolved_path).exists()

    ver = PPTGenerator.verify_presentation(resolved_path)
    assert ver["slide_count"] >= 3
    assert any("Artificial Intelligence Fundamentals" in t for t in ver["slide_titles"])


def test_content_parser_integration_with_docx(tmp_path: Path):
    """Test ContentParser.parse_file directly parses DOCX into PresentationData."""
    docx_path = tmp_path / "lecture.docx"
    sections = [
        ("Cloud Architecture Models", [
            "IaaS: Infrastructure as a Service",
            "PaaS: Platform as a Service",
            "SaaS: Software as a Service"
        ]),
        ("Summary & Takeaways", [
            "Scalability: Dynamic elastic resource allocation",
            "Reliability: Multi-region failover protocols"
        ])
    ]
    make_test_docx(docx_path, "Cloud Computing Principles", sections)

    parser = ContentParser()
    data = parser.parse_file(docx_path)

    assert data.title == "Cloud Computing Principles"
    assert len(data.all_sections) >= 2

    # Verify presentation generation from parsed DOCX
    output_pptx = tmp_path / "docx_presentation.pptx"
    gen = PPTGenerator()
    resolved_path = gen.generate(data, str(output_pptx))
    assert Path(resolved_path).exists()

    ver = PPTGenerator.verify_presentation(resolved_path)
    assert ver["slide_count"] >= 3


# =====================================================================
# 7. Gemini Service Integration Test (Mocked Client)
# =====================================================================

def test_gemini_service_integration_with_docx(tmp_path: Path):
    """Test GeminiContentPlanner.plan_from_file loads DOCX and invokes client with prompt."""
    docx_path = tmp_path / "ai_lecture.docx"
    make_test_docx(docx_path, "Robotics Curriculum", [("Actuators", ["DC Motors: Rotational motion"])])

    mock_response = MagicMock()
    mock_response.text = (
        '{\n'
        '  "title": "Robotics Curriculum",\n'
        '  "subtitle": "Engineering Lecture Series",\n'
        '  "recommended_slide_count": 6,\n'
        '  "overview": {\n'
        '    "title": "Overview",\n'
        '    "bullets": [{"prefix": "Scope", "text": "Hardware and software design"}]\n'
        '  },\n'
        '  "sections": [\n'
        '    {\n'
        '      "title": "Actuators & Sensors",\n'
        '      "bullets": [{"prefix": "Motors", "text": "Continuous drive mechanics"}],\n'
        '      "speaker_notes": "Discuss electric motor selection criteria."\n'
        '    }\n'
        '  ],\n'
        '  "conclusion": {\n'
        '    "title": "Synthesis",\n'
        '    "bullets": [{"prefix": "Future", "text": "Autonomous systems integration"}]\n'
        '  }\n'
        '}'
    )

    mock_client = MagicMock()
    mock_client.models.generate_content.return_value = mock_response

    planner = GeminiContentPlanner(client=mock_client)
    data = planner.plan_from_file(docx_path)

    assert data.title == "Robotics Curriculum"
    assert len(data.content_sections) == 1
    assert data.content_sections[0].title == "Actuators & Sensors"
    assert data.content_sections[0].speaker_notes == "Discuss electric motor selection criteria."


# =====================================================================
# 8. CLI Integration Tests
# =====================================================================

def test_cli_pasted_text_option(tmp_path: Path):
    """Test CLI execution using --text for pasted educational notes."""
    output_pptx = tmp_path / "pasted_output.pptx"
    pasted_content = (
        "# Microservices Architecture\n\n"
        "## Core Principles\n"
        "- Loose Coupling: Independent deployment services\n"
        "- High Cohesion: Single responsibility boundary\n\n"
        "## Summary\n"
        "- Resiliency: Circuit breakers and retry patterns\n"
    )

    exit_code = run(
        input_path=None,
        output_path=str(output_pptx),
        mode="deterministic",
        verify=True,
        text_input=pasted_content,
    )

    assert exit_code == 0
    assert output_pptx.exists()
    ver = PPTGenerator.verify_presentation(str(output_pptx))
    assert ver["slide_count"] >= 3


def test_cli_pdf_input(tmp_path: Path):
    """Test CLI execution with a PDF file input."""
    pdf_path = tmp_path / "cli_notes.pdf"
    content = (
        "# Operating Systems\n\n"
        "## Process Management\n"
        "- PCB: Process Control Block state tracking\n"
        "- Context Switch: CPU state save and restore\n\n"
        "## Summary\n"
        "- Multitasking: Preemptive scheduling algorithms\n"
    )
    make_test_pdf(pdf_path, [content])
    output_pptx = tmp_path / "pdf_cli_out.pptx"

    exit_code = run(
        input_path=str(pdf_path),
        output_path=str(output_pptx),
        mode="deterministic",
        verify=True,
    )

    assert exit_code == 0
    assert output_pptx.exists()


def test_cli_docx_input(tmp_path: Path):
    """Test CLI execution with a DOCX file input."""
    docx_path = tmp_path / "cli_notes.docx"
    make_test_docx(docx_path, "Data Structures", [
        ("Trees & Graphs", ["Binary Search Tree: O(log N) lookup", "B-Tree: Balanced multi-way disk storage"]),
        ("Summary", ["Trade-offs: Time versus space complexity"])
    ])
    output_pptx = tmp_path / "docx_cli_out.pptx"

    exit_code = run(
        input_path=str(docx_path),
        output_path=str(output_pptx),
        mode="deterministic",
        verify=True,
    )

    assert exit_code == 0
    assert output_pptx.exists()


def test_cli_unsupported_format_error(tmp_path: Path, capsys: pytest.CaptureFixture):
    """Test CLI displays error and exits with code 1 for unsupported file extension."""
    invalid_file = tmp_path / "notes.rtf"
    invalid_file.write_text("Sample", encoding="utf-8")
    output_pptx = tmp_path / "out.pptx"

    exit_code = run(input_path=str(invalid_file), output_path=str(output_pptx))
    assert exit_code == 1
    captured = capsys.readouterr()
    assert "Unsupported File Format" in captured.err


def test_cli_scanned_pdf_error(tmp_path: Path, capsys: pytest.CaptureFixture):
    """Test CLI displays OCR requirement error and exits with code 1 for scanned PDF."""
    scanned_pdf = tmp_path / "scanned.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(str(scanned_pdf))
    doc.close()
    output_pptx = tmp_path / "out.pptx"

    exit_code = run(input_path=str(scanned_pdf), output_path=str(output_pptx))
    assert exit_code == 1
    captured = capsys.readouterr()
    assert "Scanned Document" in captured.err or "OCR" in captured.err
