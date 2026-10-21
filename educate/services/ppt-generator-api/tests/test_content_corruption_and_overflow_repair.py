"""
Regression tests reproducing and verifying the repair of:
1. Header/footer and watermark leakage (e.g. VTUSYNC, BCS701 running headers).
2. AI topic hallucination and lack of grounding (6LoWPAN, RPL).
3. Incorrect source-to-topic-to-slide mapping.
4. Text overflow and container collision on dense content.
5. Strict hard slide limits (<= 20 slides per part).
"""

from pathlib import Path
import re
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.models import BulletPoint, ContentBlockType, PresentationData, Section, SectionType
from app.notes_extractor import NotesExtractor
from app.topic_discovery import TopicDiscoveryEngine, DiscoveredTopicDTO
from app.multi_part_planner import MultiPartPlanner
from app.ppt_generator import PPTGenerator
from app.readability_validator import ReadabilityValidator


SAMPLE_PDF_PATH = Path("sample_data/BCS701_Module_3.pdf")


def test_header_and_watermark_removal_on_source_pdf():
    """Verify that recurring 129-character headers and watermarks are stripped."""
    assert SAMPLE_PDF_PATH.exists(), f"Source fixture missing at {SAMPLE_PDF_PATH}"

    extractor = NotesExtractor()
    blocks, manifest = extractor.extract_from_file(SAMPLE_PDF_PATH)

    assert manifest.total_pages == 57
    assert manifest.word_count > 3000

    # Running header should not appear as any block's content
    for b in blocks:
        text = b.normalized_text.lower()
        # Verify running header stripped
        assert "bcs701]-2024-25" not in text, f"Running header leaked into block {b.block_id}: {b.normalized_text[:50]}"
        # Verify watermark stripped
        assert "vtusync.in" not in text, f"Watermark leaked into block {b.block_id}"
        assert "studied smart, not hard" not in text, f"Footer watermark leaked into block {b.block_id}"

    # Verify genuine curriculum headings were preserved
    headings = [b.normalized_text.strip() for b in blocks if b.block_type == ContentBlockType.HEADING]
    assert len(headings) >= 15

    heading_corpus = " ".join(headings).lower()
    assert "iot design methodology" in heading_corpus
    assert "weather monitoring" in heading_corpus
    assert "python data types" in heading_corpus
    assert "file handling" in heading_corpus
    assert "classes" in heading_corpus
    assert "question bank" in heading_corpus


def test_topic_discovery_has_zero_hallucinated_rpl_or_6lowpan():
    """Verify that topic discovery does not output RPL or 6LoWPAN since they do not exist in the source."""
    extractor = NotesExtractor()
    blocks, manifest = extractor.extract_from_file(SAMPLE_PDF_PATH)

    engine = TopicDiscoveryEngine()
    topics = engine.discover_topics(blocks, manifest, mode="deterministic")

    assert len(topics) >= 5

    topic_titles = [t.title.lower() for t in topics]
    for title in topic_titles:
        assert "rpl" not in title, f"Hallucinated RPL topic found: {title}"
        assert "6lowpan" not in title, f"Hallucinated 6LoWPAN topic found: {title}"

    # Verify real topics exist
    full_titles = " ".join(topic_titles)
    assert "iot design methodology" in full_titles
    assert "case study" in full_titles or "weather" in full_titles
    assert "python" in full_titles or "control flow" in full_titles
    assert "classes" in full_titles or "packages" in full_titles
    assert "question bank" in full_titles


def test_strict_grounding_validation_rejects_hallucinations():
    """Verify that TopicDiscoveryEngine._validate_grounding rejects ungrounded topics."""
    extractor = NotesExtractor()
    blocks, _ = extractor.extract_from_file(SAMPLE_PDF_PATH)
    engine = TopicDiscoveryEngine()

    fake_hallucinated_topics = [
        DiscoveredTopicDTO(
            topic_id="topic_001",
            title="Routing in Low Power and Lossy Networks (RPL)",
            is_inferred=True,
            page_start=39,
            page_end=57,
        ),
        DiscoveredTopicDTO(
            topic_id="topic_002",
            title="6LoWPAN Adaptation Layer Protocols",
            is_inferred=True,
            page_start=19,
            page_end=38,
        ),
    ]

    is_grounded = engine._validate_grounding(fake_hallucinated_topics, blocks)
    assert is_grounded is False, "Ungrounded topics should be rejected by _validate_grounding"


def test_multi_part_plan_source_grounding_and_slide_limits():
    """Verify that the multi-part plan contains only grounded topics, respects slide limits, and isolates review questions."""
    extractor = NotesExtractor()
    blocks, manifest = extractor.extract_from_file(SAMPLE_PDF_PATH)

    engine = TopicDiscoveryEngine()
    topics = engine.discover_topics(blocks, manifest, mode="deterministic")

    planner = MultiPartPlanner()
    plan = planner.build_plan(blocks, topics, manifest, slides_per_part=20)

    assert plan.total_parts >= 5

    for part in plan.parts:
        # Every part must have planned slides <= 20
        assert part.planned_slides <= 20, f"Part {part.part_number} exceeds limit: {part.planned_slides} slides"
        # No part should be titled RPL
        assert "rpl" not in part.title.lower(), f"Part {part.part_number} has RPL title: {part.title}"
        assert "6lowpan" not in part.title.lower(), f"Part {part.part_number} has 6LoWPAN title: {part.title}"

    # Verify the question bank units are assigned to a Question Bank topic/part
    qb_units = [u for u in plan.content_units if u.unit_type == "review_question"]
    assert len(qb_units) >= 1
    for u in qb_units:
        assert 50 <= u.source_pages[0] <= 57
        # Assigned part must not be named RPL
        assigned_part = next(p for p in plan.parts if p.part_number == u.assigned_part)
        assert "rpl" not in assigned_part.title.lower()


def test_ppt_generator_dense_content_text_fitting(tmp_path: Path):
    """Verify that dense text scales font smoothly and does not overflow container or collide."""
    dense_bullets = [
        BulletPoint(
            text="The IoT design methodology establishes an exhaustive engineering process covering Purpose & Requirements, Process Specification, Domain Model, Information Model, Service Specification, and Operational Views to guarantee robust physical-digital integration across edge sensors and cloud controllers.",
            prefix="Systemic Methodology",
        ),
        BulletPoint(
            text="Operational view specifications define the deployment configurations, communication topology, cloud middleware connections, and telemetry pipelines required to ingest real-time environmental metrics reliably without packet loss.",
            prefix="Operational Topology",
        ),
        BulletPoint(
            text="State service abstractions in REST-ful endpoints allow native edge devices to report temperature, humidity, and atmospheric barometric readings while decoupling actuation from ingestion.",
            prefix="Service Decoupling",
        ),
        BulletPoint(
            text="Component integration ties heterogeneous microcontrollers, SPI/I2C communication buses, and external transceivers to central application logic implemented in Python.",
            prefix="Component Synthesis",
        ),
    ]

    sec = Section(
        title="Comprehensive System Integration Framework",
        section_type=SectionType.CONTENT,
        bullets=dense_bullets,
    )

    data = PresentationData(
        title="IoT Architectural Engineering",
        subtitle="Advanced Pedagogical Notes",
        institution="University Engineering Series",
        content_sections=[sec],
    )

    out_file = tmp_path / "dense_test.pptx"
    gen = PPTGenerator(theme="academic_navy")
    gen.generate(data, out_file, max_slides=20)

    assert out_file.exists()

    # Post-serialization validation
    result = ReadabilityValidator.validate_pptx(out_file, max_allowed_slides=20)
    assert result.hard_limit_passed is True
    assert result.actual_exported_slides <= 20
    assert result.serious_layout_issues_detected == 0, f"Readability issues detected: {result.warnings}"


def test_serialized_pptx_full_pipeline_verification(tmp_path: Path):
    """Verify that an exported multi-part presentation series generates clean, valid PPTX files within slide limits."""
    extractor = NotesExtractor()
    blocks, manifest = extractor.extract_from_file(SAMPLE_PDF_PATH)

    engine = TopicDiscoveryEngine()
    topics = engine.discover_topics(blocks, manifest, mode="deterministic")

    planner = MultiPartPlanner()
    plan = planner.build_plan(blocks, topics, manifest, slides_per_part=20)

    # Pick Part 1 and Part 10 (Classes) and Part 12 (Question Bank)
    generator = PPTGenerator(theme="academic_navy")

    for part in [plan.parts[0], plan.parts[min(9, len(plan.parts) - 1)], plan.parts[-1]]:
        part_units = [u for u in plan.content_units if u.unit_id in part.unit_ids]
        sections = []
        for u in part_units:
            bullets = [BulletPoint(text=pt) for pt in u.key_points]
            sections.append(
                Section(
                    title=u.title,
                    section_type=SectionType.CONTENT,
                    bullets=bullets,
                )
            )

        part_data = PresentationData(
            title=part.title,
            subtitle=part.subtitle,
            institution="VTU Academic Series",
            content_sections=sections,
        )

        out_pptx = tmp_path / f"part_{part.part_number:02d}.pptx"
        generator.generate(part_data, out_pptx, max_slides=20)

        assert out_pptx.exists()
        prs = Presentation(str(out_pptx))
        assert len(prs.slides) <= 20, f"Part {part.part_number} exported {len(prs.slides)} slides > 20"

        # Validate readability
        val = ReadabilityValidator.validate_pptx(out_pptx, max_allowed_slides=20)
        assert val.hard_limit_passed is True
        assert val.serious_layout_issues_detected == 0


def test_no_course_name_substituted_as_topic_titles():
    """Verify topic discovery does not populate topics with generic course names or bare module tags."""
    extractor = NotesExtractor()
    blocks, manifest = extractor.extract_from_file(SAMPLE_PDF_PATH)
    engine = TopicDiscoveryEngine()
    topics = engine.discover_topics(blocks, manifest, mode="deterministic")

    assert len(topics) >= 10
    for t in topics:
        clean = t.title.strip().lower()
        assert not re.match(r"^internet of things(\s+\[bcs701\])?$", clean), f"Generic course title used as topic: {t.title}"
        assert not re.match(r"^module\s+[0-9ivxlcdm]+$", clean), f"Bare module label used as topic: {t.title}"
        assert not re.match(r"^bcs\d+", clean), f"Course code used as topic: {t.title}"

    # Verify key steps are distinct topics
    step_topics = [t.title for t in topics if "step" in t.title.lower()]
    assert len(step_topics) >= 5, f"Expected distinct step topics, found: {step_topics}"


def test_agenda_cards_contain_distinct_topics_and_no_body_overflow(tmp_path: Path):
    """Verify that the generated agenda slide has distinct topics across cards without repeated course names."""
    extractor = NotesExtractor()
    blocks, manifest = extractor.extract_from_file(SAMPLE_PDF_PATH)
    engine = TopicDiscoveryEngine()
    topics = engine.discover_topics(blocks, manifest, mode="deterministic")

    planner = MultiPartPlanner()
    plan = planner.build_plan(blocks, topics, manifest, slides_per_part=20)

    from app.multi_part_generator import MultiPartGenerator
    from app.models import MultiPartJobDTO

    job = MultiPartJobDTO(
        job_id="test_job_agenda",
        document_name="BCS701_Module_3.pdf",
        document_hash=manifest.file_hash,
        slides_per_part=20,
        plan=plan,
        theme="academic_navy",
    )

    mpg = MultiPartGenerator(output_dir=tmp_path)
    part_result, p_topics, p_units = mpg._generate_single_part(
        job=job,
        part_plan=plan.parts[0],
        units_by_id={u.unit_id: u for u in plan.content_units},
        all_topics=plan.topics,
    )

    assert part_result.status == "completed"
    assert part_result.slide_count <= 20

    # Inspect the agenda slide (slide 2)
    prs = Presentation(str(tmp_path / f"{part_result.file_id}.pptx"))
    slide_2 = prs.slides[1]

    # Find the module title boxes inside the roadmap cards
    card_titles = []
    for sh in slide_2.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            # If shape is a card title (located between Y 1.4" and 2.8")
            if Inches(1.3) <= sh.top <= Inches(2.9) and Inches(0.8) <= sh.left <= Inches(12.0) and len(t) > 3:
                if not t.startswith("STAGE"):
                    card_titles.append(t)

    assert len(card_titles) >= 2, f"Expected at least 2 agenda card titles, found: {card_titles}"
    # Verify titles are distinct!
    unique_card_titles = set(card_titles)
    assert len(unique_card_titles) == len(card_titles), f"Agenda cards contain repeated titles: {card_titles}"
    for ct in card_titles:
        assert "bcs701" not in ct.lower(), f"Leaked course code in card title: {ct}"


def test_embedded_images_retained_and_linked_to_split_image_slides(tmp_path: Path):
    """Verify that embedded figures are extracted and assigned to diagram content units with valid image paths."""
    extractor = NotesExtractor()
    blocks, manifest = extractor.extract_from_file(SAMPLE_PDF_PATH)
    engine = TopicDiscoveryEngine()
    topics = engine.discover_topics(blocks, manifest, mode="deterministic")

    planner = MultiPartPlanner()
    plan = planner.build_plan(blocks, topics, manifest, slides_per_part=20)

    img_units = [u for u in plan.content_units if u.image_path]
    assert len(img_units) >= 10, f"Expected at least 10 image-bearing units, found {len(img_units)}"

    # Check that image files physically exist
    for u in img_units[:5]:
        assert Path(u.image_path).exists(), f"Extracted image missing at: {u.image_path}"
        assert u.unit_type == "diagram"


def test_no_empty_bullet_slides_in_full_exported_deck(tmp_path: Path):
    """Verify that no slide in the generated presentation contains only empty bullet markers."""
    extractor = NotesExtractor()
    blocks, manifest = extractor.extract_from_file(SAMPLE_PDF_PATH)
    engine = TopicDiscoveryEngine()
    topics = engine.discover_topics(blocks, manifest, mode="deterministic")
    planner = MultiPartPlanner()
    plan = planner.build_plan(blocks, topics, manifest, slides_per_part=20)

    from app.multi_part_generator import MultiPartGenerator
    from app.models import MultiPartJobDTO

    job = MultiPartJobDTO(
        job_id="test_job_empty_bullets",
        document_name="BCS701_Module_3.pdf",
        document_hash=manifest.file_hash,
        slides_per_part=20,
        plan=plan,
        theme="academic_navy",
    )

    mpg = MultiPartGenerator(output_dir=tmp_path)
    part_result, _, _ = mpg._generate_single_part(
        job=job,
        part_plan=plan.parts[0],
        units_by_id={u.unit_id: u for u in plan.content_units},
        all_topics=plan.topics,
    )

    pptx_path = tmp_path / f"{part_result.file_id}.pptx"
    prs = Presentation(str(pptx_path))

    for s_idx, slide in enumerate(prs.slides, 1):
        if s_idx == 1:
            continue
        # Check all body textboxes
        for sh in slide.shapes:
            if sh.has_text_frame and Inches(1.5) <= sh.top <= Inches(6.8):
                txt = sh.text_frame.text.strip()
                if txt:
                    # Must not be only bullet symbols like ▪ or •
                    cleaned = re.sub(r"[\s▪•\-\*\+➤➢■\u27a2\u2022\uf0b7]+", "", txt)
                    assert len(cleaned) > 0, f"Slide {s_idx} contains shape with only bullet markers: {repr(txt)}"


@pytest.mark.parametrize("limit", [10, 15, 20, 25])
def test_hard_slide_limits_enforced_strictly(limit: int, tmp_path: Path):
    """Verify that user-configured slide budget is enforced as a strict hard maximum across 10, 15, 20, 25."""
    extractor = NotesExtractor()
    blocks, manifest = extractor.extract_from_file(SAMPLE_PDF_PATH)
    engine = TopicDiscoveryEngine()
    topics = engine.discover_topics(blocks, manifest, mode="deterministic")
    planner = MultiPartPlanner()
    plan = planner.build_plan(blocks, topics, manifest, slides_per_part=limit)

    for p in plan.parts:
        assert p.planned_slides <= limit, f"Planned slides ({p.planned_slides}) exceeds limit ({limit})"

    from app.multi_part_generator import MultiPartGenerator
    from app.models import MultiPartJobDTO

    job = MultiPartJobDTO(
        job_id=f"test_limit_{limit}",
        document_name="BCS701_Module_3.pdf",
        document_hash=manifest.file_hash,
        slides_per_part=limit,
        plan=plan,
        theme="academic_navy",
    )

    mpg = MultiPartGenerator(output_dir=tmp_path)
    part_result, _, _ = mpg._generate_single_part(
        job=job,
        part_plan=plan.parts[0],
        units_by_id={u.unit_id: u for u in plan.content_units},
        all_topics=plan.topics,
    )

    assert part_result.hard_limit_passed is True
    assert part_result.actual_exported_slides <= limit
    prs = Presentation(str(tmp_path / f"{part_result.file_id}.pptx"))
    assert len(prs.slides) <= limit

