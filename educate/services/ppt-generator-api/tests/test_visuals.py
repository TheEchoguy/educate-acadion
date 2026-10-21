"""
Unit and integration tests for Phase 3 Visual Intelligence.
Tests native diagram renderers, bounds validation, image placement,
error handling, and graceful text fallback.
"""

from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.diagrams import DiagramRenderer, RectBounds
from app.layouts import SlideLayoutEngine
from app.models import (
    BulletPoint,
    ComparisonColumn,
    ComparisonData,
    ConceptMapData,
    ConceptMapNode,
    ImageData,
    InvalidVisualDataError,
    PresentationData,
    ProcessFlowData,
    ProcessStep,
    Section,
    SlideVisual,
    TimelineData,
    TimelineMilestone,
    VisualRenderError,
    VisualType,
)
from app.ppt_generator import PPTGenerator


@pytest.fixture
def test_slide():
    """Provides a fresh blank slide with 16:9 widescreen layout."""
    layout_engine = SlideLayoutEngine()
    prs = Presentation()
    layout_engine.set_presentation_dimensions(prs)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return prs, slide, layout_engine


@pytest.fixture
def default_bounds():
    """Standard 16:9 content area bounding box."""
    return RectBounds(
        left=Inches(0.8),
        top=Inches(1.8),
        width=Inches(11.733),
        height=Inches(4.8),
    )


# =========================================================================
# Diagram Renderer Tests
# =========================================================================

def test_render_process_flow_success(test_slide, default_bounds):
    """Test rendering a native process flowchart with 4 steps."""
    prs, slide, layout = test_slide
    renderer = DiagramRenderer(layout)

    data = ProcessFlowData(steps=[
        ProcessStep(step_number=1, title="Ingestion", description="Collect training data"),
        ProcessStep(step_number=2, title="Preprocessing", description="Tokenize and normalize text"),
        ProcessStep(step_number=3, title="Training", description="Optimize neural loss function"),
        ProcessStep(step_number=4, title="Evaluation", description="Validate on holdout set"),
    ])

    success = renderer.render_process_flow(slide, data, default_bounds)
    assert success is True
    # 4 step cards + 4 top bars + 4 badges + 4 titles + 4 descriptions + 3 arrows = 23 shapes
    assert len(slide.shapes) >= 15


def test_render_process_flow_invalid_count(test_slide, default_bounds):
    """Test that fewer than 2 steps raises InvalidVisualDataError."""
    prs, slide, layout = test_slide
    renderer = DiagramRenderer(layout)

    data = ProcessFlowData(steps=[
        ProcessStep(step_number=1, title="Single Step"),
    ])
    with pytest.raises(InvalidVisualDataError):
        renderer.render_process_flow(slide, data, default_bounds)


def test_render_timeline_success(test_slide, default_bounds):
    """Test rendering a native timeline diagram with milestones."""
    prs, slide, layout = test_slide
    renderer = DiagramRenderer(layout)

    data = TimelineData(milestones=[
        TimelineMilestone(date_or_period="1950", title="Turing Test", description="Alan Turing landmark paper"),
        TimelineMilestone(date_or_period="1956", title="Dartmouth Workshop", description="McCarthy coins AI"),
        TimelineMilestone(date_or_period="1997", title="Deep Blue", description="Kasparov chess victory"),
        TimelineMilestone(date_or_period="2012", title="AlexNet", description="Deep learning breakthrough"),
    ])

    success = renderer.render_timeline(slide, data, default_bounds)
    assert success is True
    # Verify presence of axis line, milestone markers, badges, and cards
    assert len(slide.shapes) >= 16


def test_render_timeline_invalid_count(test_slide, default_bounds):
    """Test that empty or 1-milestone timeline raises InvalidVisualDataError."""
    prs, slide, layout = test_slide
    renderer = DiagramRenderer(layout)

    data = TimelineData(milestones=[
        TimelineMilestone(date_or_period="1950", title="Lone Milestone"),
    ])
    with pytest.raises(InvalidVisualDataError):
        renderer.render_timeline(slide, data, default_bounds)


def test_render_comparison_success(test_slide, default_bounds):
    """Test rendering a native two-column comparison diagram with VS badge and takeaway."""
    prs, slide, layout = test_slide
    renderer = DiagramRenderer(layout)

    data = ComparisonData(
        left_column=ComparisonColumn(
            name="Classical Machine Learning",
            tag="Feature Engineering",
            points=["Requires handcrafted feature extractors", "Effective on tabular structured datasets", "High interpretability"],
        ),
        right_column=ComparisonColumn(
            name="Deep Learning",
            tag="Representation Learning",
            points=["Learns hierarchical features end-to-end", "Requires massive labeled datasets", "Lower transparency"],
        ),
        takeaway="Select classical ML for structured business data and DL for unstructured media.",
    )

    success = renderer.render_comparison(slide, data, default_bounds)
    assert success is True
    # Left card + right card + VS badge + takeaway card
    assert len(slide.shapes) >= 6

    # Verify "VS" text is present in one of the shapes
    shape_texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    assert any("VS" in t for t in shape_texts)
    assert any("Classical Machine Learning" in t for t in shape_texts)


def test_render_concept_map_success(test_slide, default_bounds):
    """Test rendering a native concept map with center hub and 4 branches."""
    prs, slide, layout = test_slide
    renderer = DiagramRenderer(layout)

    data = ConceptMapData(
        center_concept="Artificial Intelligence Core",
        branches=[
            ConceptMapNode(title="Computer Vision", description="Scene analysis and object classification"),
            ConceptMapNode(title="Natural Language", description="Speech comprehension and synthesis"),
            ConceptMapNode(title="Robotics", description="Actuation and spatial perception"),
            ConceptMapNode(title="Knowledge Graphs", description="Ontologies and reasoning"),
        ],
    )

    success = renderer.render_concept_map(slide, data, default_bounds)
    assert success is True
    # Center hub + 4 branch cards + 4 accent strips + 4 connectors + 4 text boxes
    assert len(slide.shapes) >= 10


def test_render_image_placement(test_slide, default_bounds, tmp_path: Path):
    """Test placing a local image, ensuring aspect ratio preservation and alt text."""
    from PIL import Image

    # Create a test 600x400 image
    img_file = tmp_path / "diagram.png"
    img = Image.new("RGB", (600, 400), color=(15, 44, 89))
    img.save(img_file)

    prs, slide, layout = test_slide
    renderer = DiagramRenderer(layout)

    data = ImageData(
        file_path=str(img_file),
        caption="Figure 1: Deep Neural Network Architecture",
        alt_text="Schematic of neural network layers",
    )

    success = renderer.render_image(slide, data, default_bounds)
    assert success is True

    # Find picture shape
    pic = [s for s in slide.shapes if s.shape_type == 13]  # 13 is MSO_SHAPE_TYPE.PICTURE
    assert len(pic) == 1
    # Check that image aspect ratio 600:400 (1.5) is preserved
    aspect = pic[0].width / pic[0].height
    assert round(aspect, 2) == 1.50


def test_render_image_missing_file_raises_error(test_slide, default_bounds):
    """Test that missing image file raises VisualRenderError."""
    prs, slide, layout = test_slide
    renderer = DiagramRenderer(layout)

    data = ImageData(file_path="non_existent_image_123.png")
    with pytest.raises(VisualRenderError):
        renderer.render_image(slide, data, default_bounds)


def test_render_image_invalid_extension_raises_error(test_slide, default_bounds, tmp_path: Path):
    """Test that unsupported image extension raises InvalidVisualDataError."""
    fake_exe = tmp_path / "danger.exe"
    fake_exe.write_text("not an image", encoding="utf-8")

    prs, slide, layout = test_slide
    renderer = DiagramRenderer(layout)

    data = ImageData(file_path=str(fake_exe))
    with pytest.raises(InvalidVisualDataError):
        renderer.render_image(slide, data, default_bounds)


# =========================================================================
# PPT Generator Visual Integration & Fallback Tests
# =========================================================================

def test_generator_visual_slide_integration(tmp_path: Path):
    """Test PPTGenerator building presentation containing visual diagram slides."""
    data = PresentationData(
        title="Modern Machine Learning Architectures",
        content_sections=[
            Section(
                title="Historical Evolution",
                speaker_notes="Walk students through foundational AI milestones.",
                visual=SlideVisual(
                    visual_type=VisualType.TIMELINE,
                    timeline=TimelineData(milestones=[
                        TimelineMilestone(date_or_period="1956", title="Dartmouth", description="AI founded"),
                        TimelineMilestone(date_or_period="2012", title="AlexNet", description="CNN era"),
                    ]),
                ),
            ),
            Section(
                title="Model Paradigms Comparison",
                speaker_notes="Highlight differences between Supervised and Unsupervised learning.",
                visual=SlideVisual(
                    visual_type=VisualType.COMPARISON,
                    comparison=ComparisonData(
                        left_column=ComparisonColumn(name="Supervised", points=["Uses ground-truth labels"]),
                        right_column=ComparisonColumn(name="Unsupervised", points=["Discovers latent structure"]),
                    ),
                ),
            ),
        ]
    )

    generator = PPTGenerator()
    out_file = tmp_path / "visual_integration.pptx"
    result_path = generator.generate(data, out_file)

    assert result_path.exists()
    prs = Presentation(str(result_path))

    # Verify widescreen dimensions
    assert round(prs.slide_width / Inches(1), 2) == 13.33
    assert round(prs.slide_height / Inches(1), 2) == 7.5

    # Check that speaker notes are preserved on the visual slides
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
    combined_notes = " ".join(notes)
    assert "foundational AI milestones" in combined_notes
    assert "Supervised and Unsupervised" in combined_notes


def test_generator_visual_fallback_on_invalid_data(tmp_path: Path):
    """Verify that PPTGenerator falls back cleanly to text card layout when visual data is invalid."""
    data = PresentationData(
        title="Robust Fallback Test",
        content_sections=[
            Section(
                title="Broken Visual Section",
                bullets=[BulletPoint(text="Fallback bullet point should still be rendered.")],
                visual=SlideVisual(
                    visual_type=VisualType.LOCAL_IMAGE,
                    image=ImageData(file_path="non_existent_fallback_file.png"),
                ),
            )
        ]
    )

    generator = PPTGenerator()
    out_file = tmp_path / "fallback_test.pptx"
    # Should not raise exception; must fall back cleanly
    result_path = generator.generate(data, out_file)

    assert result_path.exists()
    prs = Presentation(str(result_path))
    # Check that the fallback text bullet was preserved
    all_text = " ".join(
        shape.text_frame.text
        for s in prs.slides
        for shape in s.shapes
        if shape.has_text_frame
    )
    assert "Fallback bullet point should still be rendered." in all_text
