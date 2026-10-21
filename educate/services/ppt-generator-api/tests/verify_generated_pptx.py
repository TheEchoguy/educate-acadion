"""
Programmatic verification script to verify that generated presentations from PDF and DOCX
contain the expected educational source text across their slide content.
"""

from pathlib import Path
from pptx import Presentation


def verify_presentation_content(pptx_path: Path, expected_keywords: list[str]) -> bool:
    """Open presentation and verify presence of expected source keywords."""
    prs = Presentation(str(pptx_path))
    all_text = []

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    all_text.append(para.text)

    full_presentation_text = " ".join(all_text)
    print(f"\n[+] Verifying {pptx_path.name} (Slides: {len(prs.slides)}):")

    missing = []
    for kw in expected_keywords:
        if kw.lower() in full_presentation_text.lower():
            print(f"    [OK] Found expected keyword: '{kw}'")
        else:
            missing.append(kw)
            print(f"    [FAIL] Missing expected keyword: '{kw}'")

    if missing:
        raise AssertionError(f"Missing expected keywords in {pptx_path.name}: {missing}")
    return True


if __name__ == "__main__":
    # 1. Verify PDF-generated presentation
    pdf_pptx = Path("output/pdf_robotics_presentation.pptx")
    pdf_keywords = [
        "Autonomous Robotics",
        "Sensor Fusion",
        "Unimate",
        "Sojourner",
        "Autonomous Navigation Pipeline",
        "State Estimation",
        "Model Predictive Control",
        "Human-Robot Collaboration",
    ]
    verify_presentation_content(pdf_pptx, pdf_keywords)

    # 2. Verify DOCX-generated presentation
    docx_pptx = Path("output/docx_cloud_presentation.pptx")
    docx_keywords = [
        "Cloud Computing",
        "Elastic Scaling",
        "Monolithic vs Microservices",
        "Distributed Consensus Workflow",
        "Log Replication",
        "Zero Trust Security",
    ]
    verify_presentation_content(docx_pptx, docx_keywords)

    print("\n[SUCCESS] Both PDF and DOCX generated presentations verified programmatically!")
