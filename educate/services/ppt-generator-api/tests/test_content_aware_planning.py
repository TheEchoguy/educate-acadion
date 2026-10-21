"""
Unit and integration tests for content-aware slide planning and varied compositions.
Verifies auto design selection, composition assignment, and valid PPTX rendering.
"""

from pathlib import Path
import tempfile
import pytest
from pptx import Presentation

from app.models import (
    BulletPoint,
    PresentationData,
    Section,
    SectionType,
    SlideVisual,
    VisualType,
    ProcessFlowData,
    ComparisonData,
    ComparisonColumn,
)
from app.layouts import auto_select_theme, get_theme
from app.content_parser import ContentParser
from app.ppt_generator import PPTGenerator


def test_auto_select_theme_domain_heuristics():
    """Verify domain classification heuristic selects suitable themes."""
    assert auto_select_theme("Deep Neural Networks and AI", "Machine learning transformers.") in ("modern_blue", "dark_tech")
    assert auto_select_theme("Cellular Biology and Oncology", "Genomic sequences and cancer cell mitosis.") == "emerald_pro"
    assert auto_select_theme("Corporate Finance and Quarterly Revenue", "Fiscal quarterly performance and balance sheets.") == "corporate_minimal"
    assert auto_select_theme("Contemporary Graphic Design and Aesthetics", "Creative typography and art gallery.") == "purple_gradient"
    assert auto_select_theme("Ancient Greco-Roman Civilizations", "Historical philosophical texts.") == "warm_creative"
    assert auto_select_theme("Introduction to General Studies", "Standard textbook material.") == "academic_navy"


def test_rule_based_composition_detection():
    """Verify rule-based parser automatically assigns varied compositions based on content semantics."""
    parser = ContentParser()
    
    # Discussion section
    sec_disc = parser._detect_composition(
        "Discussion Questions & Seminar Debate",
        [BulletPoint(text="What are the implications?")],
        None
    )
    assert sec_disc == "discussion"

    # Case study section
    sec_case = parser._detect_composition(
        "AlphaFold Case Study: Breakthrough in Biology",
        [BulletPoint(text="50-year grand challenge solved.")],
        None
    )
    assert sec_case == "case_study"

    # Comparison section
    sec_comp = parser._detect_composition(
        "Classical Machine Learning vs Deep Learning",
        [BulletPoint(text="Manual features vs automated representations.")],
        None
    )
    assert sec_comp == "two_column"

    # Core architecture section
    sec_arch = parser._detect_composition(
        "Transformer Neural Architecture Mechanics",
        [BulletPoint(text="Self-attention dot-product.")],
        None
    )
    assert sec_arch == "asymmetric"

    # 3-pillar section
    sec_pillars = parser._detect_composition(
        "Core Foundational Pillars",
        [BulletPoint(text="Pillar 1"), BulletPoint(text="Pillar 2"), BulletPoint(text="Pillar 3")],
        None
    )
    assert sec_pillars == "multi_column"

    # KPI / Metric section
    sec_metrics = parser._detect_composition(
        "Global Scaling Impact",
        [BulletPoint(text="Achieved 99.4% accuracy rate.")],
        None
    )
    assert sec_metrics == "kpi_metrics"


def test_varied_slide_renderers_generate_valid_pptx():
    """Verify that decks with varied composition types render cleanly into valid widescreen PPTX files."""
    data = PresentationData(
        title="Modern AI Architecture & Ethical Horizons",
        subtitle="Advanced Faculty Curriculum",
        institution="Acadion University",
        recommended_theme="modern_blue",
        title_composition="cinematic",
        overview_composition="agenda_flow",
        conclusion_composition="discussion_roadmap",
        overview_section=Section(
            title="Curriculum Overview & Agenda",
            section_type=SectionType.INTRO_OVERVIEW,
            paragraphs=["Comprehensive course exploring foundational mechanics and ethical challenges."],
            bullets=[
                BulletPoint(prefix="Module 1", text="Neural Architecture Fundamentals"),
                BulletPoint(prefix="Module 2", text="Empirical Case Study Analysis"),
                BulletPoint(prefix="Module 3", text="Socratic Seminar Discussion"),
            ],
        ),
        content_sections=[
            Section(
                title="Deep Transformer Architecture",
                composition_type="asymmetric",
                teaching_objective="Understand self-attention mechanisms",
                key_message="Attention computes pairwise token relevance matrices without recurrence.",
                bullets=[
                    BulletPoint(prefix="Self-Attention", text="Q, K, V matrix projections compute pairwise token relevance."),
                    BulletPoint(prefix="Multi-Head", text="Extracts concurrent representational features across diverse subspaces."),
                ],
                speaker_notes="Walk students through the attention formula on slide.",
            ),
            Section(
                title="AlphaFold Case Study & Breakthrough",
                composition_type="case_study",
                paragraphs=["Solving the 50-year protein folding challenge in molecular biology."],
                bullets=[
                    BulletPoint(prefix="Challenge", text="Predicting 3D conformations from linear amino acid sequences."),
                    BulletPoint(prefix="Architecture", text="Evoformer neural network leveraging spatial geometric constraints."),
                    BulletPoint(prefix="Impact", text="Over 200 million predicted protein structures catalogued globally."),
                    BulletPoint(prefix="Takeaway", text="Demonstrates empirical AI driving scientific discovery."),
                ],
                speaker_notes="Highlight the immense impact on biomedical therapeutics.",
            ),
            Section(
                title="Foundational Pillars of Intelligence",
                composition_type="multi_column",
                bullets=[
                    BulletPoint(prefix="Representation", text="Hierarchical latent vector embeddings capturing semantic nuances."),
                    BulletPoint(prefix="Optimization", text="Stochastic gradient descent minimizing empirical loss over parameters."),
                    BulletPoint(prefix="Generalization", text="Out-of-distribution robustness on unseen real-world tasks."),
                ],
            ),
            Section(
                title="Seminar Inquiry & Socratic Debate",
                composition_type="discussion",
                bullets=[
                    BulletPoint(prefix="Governance", text="How should frontier AI deployment be regulated internationally?"),
                    BulletPoint(prefix="Safety", text="What mitigations exist against severe algorithmic bias in healthcare?"),
                ],
                speaker_notes="Facilitate active student debate between both cohorts.",
            ),
            Section(
                title="Editorial Insights on Future Alignment",
                composition_type="editorial",
                key_message="Responsible scaling requires multi-stakeholder oversight and rigorous safety audits.",
                bullets=[
                    BulletPoint(prefix="Safety Audits", text="Pre-deployment red teaming and automated vulnerability probing."),
                    BulletPoint(prefix="Transparency", text="Mechanistic interpretability tools inspecting latent neuron activations."),
                ],
            ),
        ],
        conclusion_section=Section(
            title="Curriculum Synthesis & Horizon Outlook",
            section_type=SectionType.SUMMARY_CONCLUSION,
            bullets=[
                BulletPoint(prefix="Theoretical Mastery", text="Consolidate understanding of attention mechanisms and gradient updates."),
                BulletPoint(prefix="Practical Governance", text="Apply ethical review frameworks in capstone project deployments."),
                BulletPoint(prefix="Future Inquiry", text="Investigate neurosymbolic integration as the next computing frontier."),
            ],
            speaker_notes="Conclude presentation and direct students to homework problem sets.",
        ),
    )

    with tempfile.TemporaryDirectory() as tmp_dir:
        out_path = Path(tmp_dir) / "test_varied_deck.pptx"
        generator = PPTGenerator(theme="auto", layout_style="auto")
        resolved = generator.generate(data, str(out_path), layout_style="auto")

        assert Path(resolved).exists()
        assert Path(resolved).stat().st_size > 5000

        # Verify PPTX structure
        verification = PPTGenerator.verify_presentation(resolved)
        assert verification["slide_count"] == 8
        assert verification["widescreen_16_9"] is True
        assert verification["total_shapes"] > 40

        # Inspect slides with python-pptx
        prs = Presentation(str(resolved))
        assert len(prs.slides) == 8

        # Slide 1: Title
        assert prs.slides[0].notes_slide.notes_text_frame.text != ""
        # Slide 2: Overview / Agenda Flow
        assert len(prs.slides[1].shapes) >= 4
        # Slide 3: Asymmetric
        assert any("TEACHING OBJECTIVE" in s.text_frame.text for s in prs.slides[2].shapes if s.has_text_frame)
        # Slide 4: Case Study
        assert any("EMPIRICAL CONTEXT" in s.text_frame.text for s in prs.slides[3].shapes if s.has_text_frame)
        # Slide 5: Multi-column pillars
        assert any("PILLAR 01" in s.text_frame.text for s in prs.slides[4].shapes if s.has_text_frame)
        # Slide 6: Discussion
        assert any("SOCRATIC SEMINAR" in s.text_frame.text for s in prs.slides[5].shapes if s.has_text_frame)
        # Slide 7: Editorial
        assert any("CORE TAKEAWAY" in s.text_frame.text for s in prs.slides[6].shapes if s.has_text_frame)
        # Slide 8: Summary / Discussion Roadmap
        assert any("CURRICULUM SYNTHESIS" in s.text_frame.text for s in prs.slides[7].shapes if s.has_text_frame)
