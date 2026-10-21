"""
Unit tests for Gemini AI Content Planning Service (Phase 2).
Uses mocked SDK responses so tests run reliably offline without consuming Gemini API quota.
"""

from pathlib import Path
from unittest.mock import MagicMock, patch
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.gemini_service import (
    AIPresentationBullet,
    AIPresentationPlan,
    AIPresentationSection,
    GeminiContentPlanner,
)
from app.models import (
    EmptyContentError,
    GeminiAuthenticationError,
    GeminiOutputValidationError,
    GeminiQuotaExceededError,
    MissingAPIKeyError,
    ModelUnavailableError,
    PresentationData,
)
from app.ppt_generator import PPTGenerator


SAMPLE_VALID_PLAN_JSON = """{
  "title": "Quantum Computing Essentials",
  "subtitle": "Physics and Computational Theory",
  "institution": "Acadion Learning",
  "overview": {
    "title": "Lecture Overview & Syllabus",
    "subtitle": "Course Introduction",
    "narrative": "An introductory exploration of qubits, superposition, and quantum algorithms.",
    "bullets": [
      {"prefix": "Qubits", "text": "Two-state quantum mechanical systems functioning as basic information units."},
      {"prefix": "Superposition", "text": "Simultaneous linear combination of orthogonal basis states."}
    ],
    "speaker_notes": "Welcome students to the quantum lecture series."
  },
  "sections": [
    {
      "title": "Core Quantum Principles",
      "subtitle": "Theoretical Foundations",
      "narrative": "Fundamental mechanical properties governing quantum processing.",
      "bullets": [
        {"prefix": "Entanglement", "text": "Non-local correlations between interconnected qubit pairs."},
        {"prefix": "Interference", "text": "Constructive amplification of probability amplitudes."},
        {"prefix": "Decoherence", "text": "Loss of quantum state fidelity due to environmental interaction."}
      ],
      "speaker_notes": "Emphasize decoherence as the primary physical engineering bottleneck."
    }
  ],
  "conclusion": {
    "title": "Summary & Future Outlook",
    "subtitle": "Concluding Remarks",
    "narrative": "Quantum computing is transitioning from physics experimentation to fault-tolerant scale.",
    "bullets": [
      {"prefix": "Current Status", "text": "Noisy Intermediate-Scale Quantum (NISQ) era."},
      {"prefix": "Milestones", "text": "Development of quantum error correction codes."}
    ],
    "speaker_notes": "Direct students to lab problem set 1."
  },
  "recommended_slide_count": 5
}"""


def test_ai_presentation_plan_conversion():
    """Verify that AIPresentationPlan cleanly converts to domain PresentationData."""
    plan = AIPresentationPlan.model_validate_json(SAMPLE_VALID_PLAN_JSON)
    data = plan.to_presentation_data()

    assert isinstance(data, PresentationData)
    assert data.title == "Quantum Computing Essentials"
    assert data.subtitle == "Physics and Computational Theory"
    assert data.institution == "Acadion Learning"
    assert data.recommended_slide_count == 5

    # Overview verification
    assert data.overview_section is not None
    assert data.overview_section.title == "Lecture Overview & Syllabus"
    assert len(data.overview_section.bullets) == 2
    assert data.overview_section.bullets[0].prefix == "Qubits"
    assert data.overview_section.speaker_notes == "Welcome students to the quantum lecture series."

    # Content section verification
    assert len(data.content_sections) == 1
    sec = data.content_sections[0]
    assert sec.title == "Core Quantum Principles"
    assert len(sec.bullets) == 3
    assert sec.bullets[0].prefix == "Entanglement"
    assert sec.speaker_notes == "Emphasize decoherence as the primary physical engineering bottleneck."

    # Conclusion verification
    assert data.conclusion_section is not None
    assert data.conclusion_section.title == "Summary & Future Outlook"


def test_plan_from_text_with_mocked_sdk():
    """Test GeminiContentPlanner.plan_from_text using a mocked GenAI client."""
    mock_response = MagicMock()
    mock_response.text = SAMPLE_VALID_PLAN_JSON

    mock_client = MagicMock()
    mock_client.models.generate_content.return_value = mock_response

    planner = GeminiContentPlanner(api_key="mock_key", client=mock_client)
    data = planner.plan_from_text("Notes on quantum computing.")

    assert data.title == "Quantum Computing Essentials"
    assert len(data.content_sections) == 1
    mock_client.models.generate_content.assert_called_once()


def test_empty_notes_raises_error():
    """Verify that empty or whitespace notes raise EmptyContentError without calling API."""
    planner = GeminiContentPlanner(api_key="mock_key")
    with pytest.raises(EmptyContentError):
        planner.plan_from_text("")

    with pytest.raises(EmptyContentError):
        planner.plan_from_text("   \n\t  ")


def test_missing_api_key_raises_error():
    """Verify that MissingAPIKeyError is raised when API key is unset."""
    with patch.dict("os.environ", {}, clear=True):
        planner = GeminiContentPlanner(api_key="")
        with pytest.raises(MissingAPIKeyError) as exc_info:
            planner.plan_from_text("Valid notes content")
        assert "GEMINI_API_KEY is not configured" in str(exc_info.value)


def test_authentication_error_handling():
    """Verify that 401/403 or invalid API key errors map to GeminiAuthenticationError."""
    from google.genai.errors import ClientError

    mock_client = MagicMock()
    mock_error = ClientError(401, {"error": {"message": "API key not valid", "code": 401}}, None)
    mock_client.models.generate_content.side_effect = mock_error

    planner = GeminiContentPlanner(api_key="mock_invalid_key", client=mock_client)
    with pytest.raises(GeminiAuthenticationError) as exc_info:
        planner.plan_from_text("Some notes")
    assert "invalid" in str(exc_info.value).lower()


def test_quota_exceeded_error_handling():
    """Verify that 429 quota exhaustion errors map to GeminiQuotaExceededError."""
    from google.genai.errors import ClientError

    mock_client = MagicMock()
    mock_error = ClientError(429, {"error": {"message": "ResourceExhausted quota exceeded", "code": 429}}, None)
    mock_client.models.generate_content.side_effect = mock_error

    planner = GeminiContentPlanner(api_key="mock_key", client=mock_client)
    with pytest.raises(GeminiQuotaExceededError) as exc_info:
        planner.plan_from_text("Some notes")
    assert "quota or rate limit" in str(exc_info.value)


def test_model_unavailable_error_handling():
    """Verify that 404 model errors map to ModelUnavailableError."""
    from google.genai.errors import ClientError

    mock_client = MagicMock()
    mock_error = ClientError(404, {"error": {"message": "models/foo-bar not found", "code": 404}}, None)
    mock_client.models.generate_content.side_effect = mock_error

    planner = GeminiContentPlanner(api_key="mock_key", model="foo-bar", client=mock_client)
    with pytest.raises(ModelUnavailableError) as exc_info:
        planner.plan_from_text("Some notes")
    assert "foo-bar" in str(exc_info.value)


def test_malformed_json_raises_validation_error():
    """Verify that invalid JSON or incomplete schema raises GeminiOutputValidationError."""
    mock_response = MagicMock()
    mock_response.text = '{"malformed": "not a presentation plan"}'

    mock_client = MagicMock()
    mock_client.models.generate_content.return_value = mock_response

    planner = GeminiContentPlanner(api_key="mock_key", client=mock_client)
    with pytest.raises(GeminiOutputValidationError):
        planner.plan_from_text("Notes")


def test_ai_presentation_integration_with_ppt_engine(tmp_path: Path):
    """Verify that PresentationData from AI plan generates a valid 16:9 .pptx with speaker notes."""
    plan = AIPresentationPlan.model_validate_json(SAMPLE_VALID_PLAN_JSON)
    data = plan.to_presentation_data()

    generator = PPTGenerator()
    out_file = tmp_path / "ai_integration_test.pptx"
    result_path = generator.generate(data, out_file)

    assert result_path.exists()
    assert result_path.stat().st_size > 10_000

    # Reopen presentation and verify
    prs = Presentation(str(result_path))
    assert round(prs.slide_width / Inches(1), 2) == 13.33
    assert round(prs.slide_height / Inches(1), 2) == 7.5

    # Verify speaker notes are stored on the slides
    notes_texts = [
        s.notes_slide.notes_text_frame.text
        for s in prs.slides
        if s.has_notes_slide
    ]
    combined_notes = " ".join(notes_texts)
    assert "quantum lecture series" in combined_notes
    assert "decoherence as the primary physical engineering bottleneck" in combined_notes
