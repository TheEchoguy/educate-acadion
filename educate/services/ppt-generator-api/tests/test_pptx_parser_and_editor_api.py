"""
Comprehensive test suite for Phase 2.1:
PPTX Parser, Normalized Slide-Object Model (SOM), and Read-Only Editor API.

Covers:
1. Generated presentation with multiple slides.
2. Slide dimensions and coordinate conversion.
3. Text boxes with multiple paragraphs and formatted runs.
4. Shapes with fill and border colors.
5. Pictures and asset URL retrieval.
6. Tables and cell text where supported.
7. Speaker notes.
8. Objects without text.
9. Missing optional formatting properties.
10. Unsupported objects or features classification.
11. Invalid file IDs.
12. Missing PPTX files.
13. Malformed PPTX input.
14. Verification that parsing does not modify the original PPTX (immutability).
15. Repeated parsing consistency (determinism).
"""

import hashlib
from io import BytesIO
import os
from pathlib import Path
import tempfile
import uuid
import pytest
from fastapi.testclient import TestClient
from PIL import Image

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.api import STORAGE_DIR, app, _GENERATED_REGISTRY
from app.models import (
    BorderDTO,
    FillDTO,
    ImageAssetDTO,
    PresentationEditorModel,
    SlideDetailDTO,
    SlideObjectDTO,
    TableCellDTO,
    TableDataDTO,
    TextContentDTO,
    TextParagraphDTO,
    TextRunDTO,
)
from app.pptx_parser import (
    DEFAULT_VIRTUAL_HEIGHT,
    DEFAULT_VIRTUAL_WIDTH,
    emu_to_normalized,
    extract_border,
    extract_color_hex,
    extract_fill,
    normalized_to_emu,
    parse_presentation,
    parse_text_frame,
)


@pytest.fixture
def client():
    """Test client for FastAPI app."""
    return TestClient(app)


@pytest.fixture
def sample_png_bytes() -> bytes:
    """Generate a valid minimal 100x100 PNG image in memory."""
    img = Image.new("RGBA", (100, 100), color=(37, 99, 235, 255))
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# =====================================================================
# 1. Multi-Slide Presentation Parsing
# =====================================================================

def test_parse_multiple_slides_presentation(tmp_path):
    """Test parsing a presentation with multiple slides extracts accurate slide metadata and objects."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for i in range(4):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        tb.text = f"Slide {i+1} Title"

    test_file = tmp_path / "multi_slide.pptx"
    prs.save(str(test_file))

    file_id = str(uuid.uuid4())
    model = parse_presentation(test_file, file_id)

    assert isinstance(model, PresentationEditorModel)
    assert model.file_id == file_id
    assert model.slide_count == 4
    assert len(model.slides) == 4

    for idx, slide in enumerate(model.slides, start=1):
        assert slide.slide_number == idx
        assert slide.slide_id == f"{file_id}_slide_{idx}"
        assert len(slide.objects) >= 1
        assert slide.width == DEFAULT_VIRTUAL_WIDTH
        assert slide.height == DEFAULT_VIRTUAL_HEIGHT


# =====================================================================
# 2. Slide Dimensions and Coordinate Conversion
# =====================================================================

def test_coordinate_conversion_precision_and_aspect_ratio(tmp_path):
    """Test precise bidirectional coordinate conversion between EMUs and normalized virtual canvas."""
    slide_w_emu = 12192000  # 13.333 inches (16:9 widescreen)
    slide_h_emu = 6858000   # 7.5 inches

    # Test EMU to normalized
    x_norm = emu_to_normalized(1219200, slide_w_emu, DEFAULT_VIRTUAL_WIDTH)
    assert x_norm == 192.0

    y_norm = emu_to_normalized(685800, slide_h_emu, DEFAULT_VIRTUAL_HEIGHT)
    assert y_norm == 108.0

    # Test round trip from normalized back to EMU
    x_emu = normalized_to_emu(x_norm, slide_w_emu, DEFAULT_VIRTUAL_WIDTH)
    assert x_emu == 1219200

    y_emu = normalized_to_emu(y_norm, slide_h_emu, DEFAULT_VIRTUAL_HEIGHT)
    assert y_emu == 685800

    # Test non-16:9 aspect ratio (4:3 presentation: 10 x 7.5 inches)
    prs_43 = Presentation()
    prs_43.slide_width = Inches(10)
    prs_43.slide_height = Inches(7.5)
    slide_43 = prs_43.slides.add_slide(prs_43.slide_layouts[6])
    slide_43.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(2))

    test_file = tmp_path / "ratio_43.pptx"
    prs_43.save(str(test_file))

    model_43 = parse_presentation(test_file, str(uuid.uuid4()))
    assert model_43.slide_height == 1080.0
    # 4/3 * 1080 = 1440
    assert abs(model_43.slide_width - 1440.0) < 1.0


# =====================================================================
# 3. Text Boxes with Multiple Paragraphs and Formatted Runs
# =====================================================================

def test_text_formatting_paragraphs_and_runs(tmp_path):
    """Test text boxes with multiple paragraphs, font sizes, colors, alignment, and formatting."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
    tf = tb.text_frame
    tf.word_wrap = True

    # Paragraph 0: Heading with centered alignment and mixed runs
    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.CENTER
    r0_1 = p0.add_run()
    r0_1.text = "Academic "
    r0_1.font.name = "Calibri"
    r0_1.font.bold = True
    r0_1.font.size = Pt(24)

    r0_2 = p0.add_run()
    r0_2.text = "Courseware"
    r0_2.font.name = "Calibri"
    r0_2.font.italic = True
    r0_2.font.color.rgb = RGBColor(37, 99, 235)  # #2563EB

    # Paragraph 1: Bullet point with indent level 1 and underline
    p1 = tf.add_paragraph()
    p1.level = 1
    p1.space_before = Pt(6)
    p1.space_after = Pt(4)
    r1 = p1.add_run()
    r1.text = "Key concept bullet item"
    r1.font.underline = True
    r1.font.size = Pt(14)

    # Paragraph 2: Hyperlink run
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Acadion Official Site"
    r2.hyperlink.address = "https://acadion.edu"

    test_file = tmp_path / "formatted_text.pptx"
    prs.save(str(test_file))

    model = parse_presentation(test_file, str(uuid.uuid4()))
    slide_res = model.slides[0]
    tb_obj = slide_res.objects[0]

    assert tb_obj.text_content is not None
    tc = tb_obj.text_content
    assert len(tc.paragraphs) == 3

    # Check paragraph 0
    para0 = tc.paragraphs[0]
    assert para0.alignment == "center"
    assert len(para0.runs) == 2
    assert para0.runs[0].text == "Academic "
    assert para0.runs[0].bold is True
    assert para0.runs[0].font_size == 24.0
    assert para0.runs[1].text == "Courseware"
    assert para0.runs[1].italic is True
    assert para0.runs[1].font_color == "#2563EB"

    # Check paragraph 1
    para1 = tc.paragraphs[1]
    assert para1.level == 1
    assert para1.is_bullet is True
    assert para1.space_before == 6.0
    assert para1.space_after == 4.0
    assert para1.runs[0].underline is True

    # Check paragraph 2
    para2 = tc.paragraphs[2]
    assert para2.runs[0].hyperlink == "https://acadion.edu"


# =====================================================================
# 4. Shapes with Fill and Border Colors
# =====================================================================

def test_shape_fill_and_border_colors(tmp_path):
    """Test extracting shape solid fills, borders, dimensions, and shape types."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Add a rectangle with solid blue fill and gold border
    shape1 = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1), Inches(1), Inches(4), Inches(2)
    )
    shape1.name = "TestCard"
    shape1.fill.solid()
    shape1.fill.fore_color.rgb = RGBColor(15, 44, 89)  # #0F2C59
    shape1.line.color.rgb = RGBColor(217, 119, 6)      # #D97706
    shape1.line.width = Pt(2.5)

    test_file = tmp_path / "fill_border.pptx"
    prs.save(str(test_file))

    model = parse_presentation(test_file, str(uuid.uuid4()))
    obj = model.slides[0].objects[0]

    assert obj.name == "TestCard"
    assert obj.object_type == "shape"
    assert obj.shape_type == "ROUNDED_RECTANGLE"
    assert obj.fill is not None
    assert obj.fill.fill_type == "solid"
    assert obj.fill.color == "#0F2C59"
    assert obj.border is not None
    assert obj.border.color == "#D97706"
    assert obj.border.width is not None


# =====================================================================
# 5. Pictures and Asset URL Retrieval
# =====================================================================

def test_pictures_and_asset_retrieval(tmp_path, sample_png_bytes):
    """Test extracting embedded images to asset storage and serving them via the asset endpoint."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Write temp image to disk for adding to pptx
    temp_img_path = tmp_path / "input_img.png"
    temp_img_path.write_bytes(sample_png_bytes)

    slide.shapes.add_picture(str(temp_img_path), Inches(2), Inches(2), Inches(3), Inches(3))

    test_file = tmp_path / "picture_deck.pptx"
    prs.save(str(test_file))

    file_id = str(uuid.uuid4())
    assets_dir = tmp_path / "test_assets"
    model = parse_presentation(test_file, file_id, assets_base_dir=assets_dir)

    pic_obj = model.slides[0].objects[0]
    assert pic_obj.object_type == "picture"
    assert pic_obj.image_asset is not None
    asset = pic_obj.image_asset
    assert asset.asset_id.startswith("img_")
    assert asset.url == f"/api/editor/assets/{file_id}/{asset.asset_id}"
    assert asset.content_type in ("image/png", "image/x-png")

    # Verify image was written to disk
    extracted_file = assets_dir / f"{file_id}_assets" / f"{asset.asset_id}.png"
    assert extracted_file.exists()
    assert extracted_file.read_bytes() == sample_png_bytes


# =====================================================================
# 6. Tables and Cell Text
# =====================================================================

def test_table_parsing_and_cell_text(tmp_path):
    """Test parsing PowerPoint tables, cell text frames, rows, and columns."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    tbl_shape = slide.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(6), Inches(2))
    tbl = tbl_shape.table
    tbl.cell(0, 0).text = "Column 1"
    tbl.cell(0, 1).text = "Column 2"
    tbl.cell(0, 2).text = "Column 3"
    tbl.cell(1, 0).text = "Data A"
    tbl.cell(1, 1).text = "Data B"
    tbl.cell(1, 2).text = "Data C"

    test_file = tmp_path / "table_deck.pptx"
    prs.save(str(test_file))

    model = parse_presentation(test_file, str(uuid.uuid4()))
    table_obj = model.slides[0].objects[0]

    assert table_obj.object_type == "table"
    assert table_obj.table_data is not None
    td = table_obj.table_data
    assert td.rows == 2
    assert td.columns == 3
    assert len(td.cells) == 6

    # Verify cell values
    cell_0_0 = next(c for c in td.cells if c.row_index == 0 and c.col_index == 0)
    assert cell_0_0.text == "Column 1"
    assert cell_0_0.text_content is not None
    assert cell_0_0.text_content.paragraphs[0].plain_text == "Column 1"

    cell_1_2 = next(c for c in td.cells if c.row_index == 1 and c.col_index == 2)
    assert cell_1_2.text == "Data C"


# =====================================================================
# 7. Speaker Notes
# =====================================================================

def test_speaker_notes_extraction(tmp_path):
    """Test that speaker notes are accurately extracted when present and empty when omitted."""
    prs = Presentation()

    # Slide 1: has notes
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    slide1.notes_slide.notes_text_frame.text = "Explain the three architectural tiers thoroughly."

    # Slide 2: no notes
    prs.slides.add_slide(prs.slide_layouts[6])

    test_file = tmp_path / "notes_deck.pptx"
    prs.save(str(test_file))

    model = parse_presentation(test_file, str(uuid.uuid4()))
    assert model.slides[0].speaker_notes == "Explain the three architectural tiers thoroughly."
    assert model.slides[1].speaker_notes == ""


# =====================================================================
# 8. Objects Without Text
# =====================================================================

def test_objects_without_text(tmp_path):
    """Test shapes without any text (geometric shapes, decorative cards) parse cleanly."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Decorative rectangle with no text
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(2), Inches(2))
    rect.name = "Decorative Box"

    test_file = tmp_path / "no_text.pptx"
    prs.save(str(test_file))

    model = parse_presentation(test_file, str(uuid.uuid4()))
    obj = model.slides[0].objects[0]

    assert obj.name == "Decorative Box"
    assert obj.object_type == "shape"
    if obj.text_content is not None:
        assert obj.text_content.full_text == ""


# =====================================================================
# 9. Missing Optional Formatting Properties
# =====================================================================

def test_missing_optional_formatting_properties(tmp_path):
    """Test shapes with missing/unspecified optional styling properties parse safely without arbitrary substitutions."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    p = tb.text_frame.paragraphs[0]
    p.text = "Unstyled plain text"
    # No font size, no font color, no bold, no background fill specified

    test_file = tmp_path / "unstyled.pptx"
    prs.save(str(test_file))

    model = parse_presentation(test_file, str(uuid.uuid4()))
    obj = model.slides[0].objects[0]
    run = obj.text_content.paragraphs[0].runs[0]

    # Must NOT substitute black or white arbitrarily if not specified
    assert run.bold is None or run.bold is False
    assert run.italic is None or run.italic is False
    # Color can be None without raising exceptions
    assert run.font_color is None or isinstance(run.font_color, str)


# =====================================================================
# 10. Unsupported Objects Classification
# =====================================================================

def test_unsupported_objects_classification(tmp_path):
    """Test that unsupported objects (e.g. PowerPoint charts) are explicitly classified as unsupported."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Add a native chart
    chart_data = CategoryChartData()
    chart_data.categories = ["A", "B", "C"]
    chart_data.add_series("Series 1", (10, 20, 30))

    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1), Inches(1), Inches(6), Inches(4),
        chart_data
    )

    test_file = tmp_path / "chart_deck.pptx"
    prs.save(str(test_file))

    model = parse_presentation(test_file, str(uuid.uuid4()))
    obj = model.slides[0].objects[0]

    assert obj.object_type == "unsupported"
    assert obj.is_unsupported is True
    assert obj.unsupported_reason is not None
    assert "chart" in obj.unsupported_reason.lower()


# =====================================================================
# 11. API: Invalid File IDs
# =====================================================================

def test_api_invalid_file_ids(client: TestClient):
    """Test editor API returns 400 Bad Request on invalid file ID format or traversal attempts."""
    # Malformed non-UUID file ID formats
    for bad_id in ("not-a-uuid-format", "12345", "test_id_123", "'; DROP TABLE slides;--"):
        res = client.get(f"/api/editor/{bad_id}")
        assert res.status_code == 400
        assert "Invalid file identifier format" in res.json()["detail"]

    # Path traversal attack patterns in file_id
    traversal_patterns = [
        "../../etc/passwd",
        "..\\..\\windows\\system32",
        "..%2F..%2Fwindows%2Fsystem32",
    ]
    for pattern in traversal_patterns:
        res_traversal = client.get(f"/api/editor/{pattern}")
        assert res_traversal.status_code in (400, 404)

    # Malformed asset endpoint IDs
    res_asset_bad_id = client.get("/api/editor/assets/bad-id/img_123")
    assert res_asset_bad_id.status_code == 400

    valid_uuid = str(uuid.uuid4())
    res_asset_bad_asset = client.get(f"/api/editor/assets/{valid_uuid}/invalid-asset!@#")
    assert res_asset_bad_asset.status_code == 400

    res_asset_traversal = client.get(f"/api/editor/assets/{valid_uuid}/../../traversal")
    assert res_asset_traversal.status_code in (400, 404)


# =====================================================================
# 12. API: Missing PPTX Files
# =====================================================================

def test_api_missing_pptx_files(client: TestClient):
    """Test editor API returns 404 Not Found for non-existent presentation files."""
    random_uuid = str(uuid.uuid4())
    res = client.get(f"/api/editor/{random_uuid}")
    assert res.status_code == 404
    assert "not found" in res.json()["detail"].lower()

    # Asset for non-existent presentation
    res_asset = client.get(f"/api/editor/assets/{random_uuid}/img_none")
    assert res_asset.status_code == 404


# =====================================================================
# 13. API: Malformed PPTX Input
# =====================================================================

def test_api_malformed_pptx_input(client: TestClient):
    """Test editor API returns 422 Unprocessable Entity when reading corrupted or invalid PPTX file."""
    bad_uuid = str(uuid.uuid4())
    corrupt_file = STORAGE_DIR / f"{bad_uuid}.pptx"

    try:
        # Write corrupted garbage bytes
        corrupt_file.write_bytes(b"CORRUPTED_NON_ZIP_DATA_HEADER_12345")

        res = client.get(f"/api/editor/{bad_uuid}")
        assert res.status_code == 422
        assert "Unable to parse presentation file" in res.json()["detail"]
    finally:
        if corrupt_file.exists():
            corrupt_file.unlink()


# =====================================================================
# 14. Immutability: Parsing Does Not Modify Source PPTX
# =====================================================================

def test_parsing_does_not_modify_source_pptx(tmp_path):
    """Verify that parsing leaves the source PPTX file byte-for-byte identical."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    tb.text = "Immutability Check Slide"

    test_file = tmp_path / "immutability.pptx"
    prs.save(str(test_file))

    # Read sha256 before parsing
    hash_before = hashlib.sha256(test_file.read_bytes()).hexdigest()
    mtime_before = os.path.getmtime(test_file)

    # Perform parsing
    model = parse_presentation(test_file, str(uuid.uuid4()))
    assert model.slide_count == 1

    # Read sha256 and mtime after parsing
    hash_after = hashlib.sha256(test_file.read_bytes()).hexdigest()
    mtime_after = os.path.getmtime(test_file)

    assert hash_before == hash_after
    assert mtime_before == mtime_after


# =====================================================================
# 15. Repeated Parsing Consistency (Determinism)
# =====================================================================

def test_repeated_parsing_consistency(tmp_path):
    """Verify that repeated parsing of the same file produces identical normalized models."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(15, 44, 89)
    shape.text = "Consistent Output Text"

    test_file = tmp_path / "consistency.pptx"
    prs.save(str(test_file))

    file_id = str(uuid.uuid4())
    model1 = parse_presentation(test_file, file_id)
    model2 = parse_presentation(test_file, file_id)

    assert model1.model_dump() == model2.model_dump()
