"""
Script to generate sample PDF and DOCX educational lecture notes for Phase 4 verification.
"""

from pathlib import Path
import docx
import pymupdf


def generate_sample_pdf(output_path: Path):
    """Generate a clean, structured educational PDF document."""
    doc = pymupdf.open()
    
    # Page 1: Overview and History
    p1 = doc.new_page()
    p1_text = (
        "# Autonomous Robotics & Intelligent Systems\n"
        "Advanced Engineering Curriculum — Acadion Learning\n\n"
        "## Introduction & Foundations\n"
        "- Autonomous Agent: Self-governing entity operating in dynamic environments\n"
        "- Sensor Fusion: Combining IMU, LiDAR, and camera telemetry\n"
        "- Actuation: Closed-loop motor control and torque regulation\n\n"
        "## Robotics Milestones\n"
        "- 1961: Unimate becomes the first industrial robot deployed at GM\n"
        "- 1997: Sojourner rover navigates the Martian surface autonomously\n"
        "- 2015: Deep reinforcement learning enables agile legged locomotion\n"
        "- 2024: Foundation robotics models generalize zero-shot manipulation\n"
    )
    p1.insert_text((50, 72), p1_text, fontsize=11)

    # Page 2: Perception, Planning, and Control
    p2 = doc.new_page()
    p2_text = (
        "## Autonomous Navigation Pipeline\n"
        "- Environmental Sensing: Capturing point clouds and RGB visual frames\n"
        "- State Estimation: Extended Kalman Filtering and SLAM mapping\n"
        "- Path Planning: Global Dijkstra search and local trajectory rollout\n"
        "- Motion Execution: Model Predictive Control trajectory tracking\n\n"
        "## Summary & Future Outlook\n"
        "- Physical Safety: Hardware interlocks and formal verification bounds\n"
        "- Human-Robot Collaboration: Compliant actuation and intent prediction\n"
        "- Scalability: Fleet learning across distributed robotic deployments\n"
    )
    p2.insert_text((50, 72), p2_text, fontsize=11)

    doc.save(str(output_path))
    doc.close()
    print(f"[+] Created sample PDF: {output_path}")


def generate_sample_docx(output_path: Path):
    """Generate a clean, structured educational Word document."""
    doc = docx.Document()

    # Title
    doc.add_heading("Cloud Computing & Distributed Systems", level=1)
    p_sub = doc.add_paragraph("Curriculum Masterclass — Acadion Engineering Series")
    p_sub.style = "Subtitle"

    # Section 1: Overview
    doc.add_heading("Introduction & Architectural Foundations", level=2)
    doc.add_paragraph("Cloud Computing: On-demand availability of system resources and data storage", style="List Bullet")
    doc.add_paragraph("Elastic Scaling: Dynamically provisioning compute capacity to meet workload demand", style="List Bullet")
    doc.add_paragraph("Fault Tolerance: Redundant distributed infrastructure preventing single-point failure", style="List Bullet")

    # Section 2: Comparison
    doc.add_heading("Monolithic vs Microservices Architecture", level=2)
    doc.add_paragraph("Monolithic: Single unified codebase and shared relational database", style="List Bullet")
    doc.add_paragraph("Monolithic: Simpler initial deployment but complex long-term maintenance", style="List Bullet")
    doc.add_paragraph("Microservices: Independently deployable bounded contextual services", style="List Bullet")
    doc.add_paragraph("Microservices: Resilient fault isolation with decentralized data storage", style="List Bullet")

    # Section 3: Distributed Consensus
    doc.add_heading("Distributed Consensus Workflow", level=2)
    doc.add_paragraph("Client Proposal: Client submits state mutation request to leader", style="List Bullet")
    doc.add_paragraph("Log Replication: Leader appends entry and broadcasts to follower nodes", style="List Bullet")
    doc.add_paragraph("Quorum Agreement: Majority of nodes validate and acknowledge the log entry", style="List Bullet")
    doc.add_paragraph("Commit & Response: Leader commits state machine update and responds to client", style="List Bullet")

    # Section 4: Summary
    doc.add_heading("Summary & Best Practices", level=2)
    doc.add_paragraph("Observability: Comprehensive metrics, structured tracing, and distributed logs", style="List Bullet")
    doc.add_paragraph("Zero Trust Security: Mutual TLS authentication and least-privilege IAM roles", style="List Bullet")
    doc.add_paragraph("Cost Optimization: Auto-scaling and spot instance scheduling", style="List Bullet")

    doc.save(str(output_path))
    print(f"[+] Created sample DOCX: {output_path}")


if __name__ == "__main__":
    sample_dir = Path("sample_data")
    generate_sample_pdf(sample_dir / "sample_robotics_lecture.pdf")
    generate_sample_docx(sample_dir / "sample_cloud_curriculum.docx")
